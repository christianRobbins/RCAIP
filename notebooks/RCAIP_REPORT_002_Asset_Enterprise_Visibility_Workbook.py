# Databricks notebook source
# MAGIC %md
# MAGIC # RCAIP REPORT 002 — Asset / Enterprise Visibility Workbook
# MAGIC
# MAGIC CURRENT CHECKOUT AUTHORITY:
# MAGIC D0064, current Config/producer contracts and committed DIAG through Cell 38.
# MAGIC
# MAGIC HISTORICAL IMPLEMENTATION REFERENCE:
# MAGIC Recovered REPORT 001 Library artifacts are not checkout authority.
# MAGIC Deterministic Excel escaping is adopted as presentation engineering.
# MAGIC
# MAGIC Only new execution parameter: RCAIP_REPORT002_MODE.
# MAGIC No Config, producer, DIAG or recovery notebook is executed here.
# MAGIC No external publication-scope artifact is required.
# MAGIC
# MAGIC Would this exact presentation logic run for Paris? YES.
# MAGIC NO NEW ANALYTICAL AUTHORITY INTRODUCED

# COMMAND ----------

import hashlib
import io
import json
import math
import re
import tempfile
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from py4j.protocol import Py4JJavaError
from pyspark.sql import DataFrame
from pyspark.sql import functions as F
from pyspark.sql.utils import AnalysisException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from databricks.sdk import WorkspaceClient
from databricks.sdk.service.workspace import ExportFormat, ImportFormat

RCAIP_REPORT002_READY = False
for _name in (
    "RCAIP_REPORT002_PUBLISHED_PATH",
    "RCAIP_REPORT002_SHA256",
    "RCAIP_REPORT002_CERTIFICATE",
):
    globals().pop(_name, None)

MODES = {"DURABLE_EVIDENCE_V1", "CELL38_FULL"}
NOT_EXECUTED = "NOT EXECUTED / NOT AVAILABLE IN THIS EDITION"
NO_SIGNAL = "DISC002_NO_USABLE_ATTRIBUTION_SIGNAL"
POST_CLASS = "IP005_POST_CMDB_REGISTRATION_VALIDATION"
CROSS_CLASS = "DISC002_CONSISTENT_MULTI_SOURCE_ATTRIBUTION_SEED"
UMBRELLA = "ENTERPRISE_ATTRIBUTION_LOOKUP_OR_BOUNDED_POST_CORROBORATION"

SHEET_NAMES = [
    "Overview", "Asset Review Detail", "Attribution Evidence Review",
    "Post Validation Context", "Enterprise Crosswalk Context",
    "Evidence Limitations", "Source Evidence", "Tool Coverage",
]
OPERATOR_COLUMNS = [
    "Asset", "Device Context", "Current Observation", "Lifecycle",
    "Active Sources", "Stale Sources", "Identity Evidence",
    "Attribution Evidence", "Evidence Limitation", "Analytical Branch",
    "Analytical Continuation", "Customer Determination",
    "Workload Authorization",
]

MAX_ASSETS = 50000
MAX_INPUT_ROWS = 200000
MAX_INPUT_CELLS = 12000000
MAX_INPUT_BYTES = 192 * 1024 * 1024
MAX_EXPORT_CELLS = 8000000
MAX_FILE_BYTES = 100 * 1024 * 1024

SCOPE_FIELDS = ["post_slug", "rcaip_run_id"]
SOURCE_FLAGS = {
    "CMDB": "cmdb_present",
    "NMAP": "nmap_present",
    "Tenable": "tenable_present",
    "Intune": "intune_present",
    "MDE": "mde_present",
    "Qualys": "qualys_present",
}
FEATURE_IDENTITY_FIELDS = """
canonical_name canonical_serial canonical_ip device_family
identity_basis identity_strength source_mask_raw source_mask_active
source_mask_stale current_observation_status attribution_confidence
attribution_conflict_flag conflict_count
""".split()
FEATURE_FIELDS = FEATURE_IDENTITY_FIELDS + """
device_scope network_scope operational_boundary operational_boundary_basis
operational_boundary_confidence management_model management_owner
visibility_evidence_count visibility_evidence_mask visibility_class
post_visibility_state overall_assurance_state overall_review_priority
primary_routing_scope applicable_control_count applicability_review_count
healthy_control_count missing_control_count not_reporting_control_count
stale_reporting_control_count misconfigured_control_count
noncompliant_control_count configuration_unknown_count source_unavailable_count
identity_reconciliation_count evidence_only_relationship_count
adverse_relationship_count active_unresolved_ip_candidate
control_assurance_contract_version control_assurance_build_version
control_assurance_build_mode
""".split()
DISC_FIELDS = """
candidate_scope attribution_evidence_status attribution_confidence
attribution_signal_count attribution_source_count
distinct_attribution_value_count direct_assignment_signal_count
proposed_crosswalk_value dominant_value_source_count dominant_value_signal_count
dominant_value_sources attribution_route recommended_next_action
validation_responsibility final_disposition disposition_status
automatic_assignment_allowed automatic_finding_allowed automatic_ticket_allowed
automatic_remediation_allowed
""".split()
ATTR_FIELDS = """
qualification_class qualification_next_gate qualification_evidence_basis
boundary_semantic_class discovery002_coverage_state
existing_validation_responsibility workload_promotion_state
""".split()
MDE_PROVENANCE = """
enterprise_mde_source_logical enterprise_mde_source_physical
enterprise_mde_source_delta_version enterprise_mde_source_delta_timestamp
enterprise_mde_source_schema_sha256 search_authority_sha256
""".split()
AUTHORITY_FIELDS = """
asset_key search_eligibility_class selected_short_name
short_name_lookup_eligible collision_safe_full_name_variants
selected_supporting_source_families selected_supporting_sources
original_mde_assurance_state search_input_contract_version
search_input_build_version
""".split()

# Exact central-sheet projections. Native-grain and certification products
# retain their complete published schema.
PROJECT = {
    "asset": ["asset_key", *SCOPE_FIELDS, *SOURCE_FLAGS.values()] + """
host_key observed_source_count_raw observed_source_count_active
observed_source_count_stale current_observation_status
requires_post_validation validation_reason reconciliation_confidence
lifecycle_reconciliation_status
""".split(),
    "mart": ["asset_key", *SCOPE_FIELDS, *FEATURE_IDENTITY_FIELDS,
             "feature_contract_version", "feature_build_version",
             "feature_build_mode"],
    "feature": ["asset_key", *SCOPE_FIELDS, *FEATURE_FIELDS],
    "lifecycle": ["asset_key", *SCOPE_FIELDS, *SOURCE_FLAGS.values()] + """
host_key current_observation_status evidence_state lifecycle_state
lifecycle_status most_recent_record_timestamp most_recent_active_timestamp
days_since_most_recent_record days_since_most_recent_active
record_age_bucket active_age_bucket validation_responsibility recommended_action
""".split(),
    "operational": ["asset_key", *SCOPE_FIELDS, *SOURCE_FLAGS.values()] + """
current_observation_status operational_population_class current_evidence_flag
stale_evidence_flag unknown_freshness_flag cmdb_present_flag
tool_expectation_count tool_validation_required_count
automatic_finding_allowed automatic_ticket_allowed automatic_remediation_allowed
""".split(),
    "investigation": ["asset_key", *SCOPE_FIELDS, *SOURCE_FLAGS.values()] + """
current_observation_status operational_population_class current_evidence_flag
stale_evidence_flag cmdb_present_flag tool_expectation_count
primary_case_type identity_priority recommended_owner_scope identity_reason
identity_recommended_action
""".split(),
    "discovery": ["asset_key", *SCOPE_FIELDS, *DISC_FIELDS],
    "mde": ["asset_key", *SCOPE_FIELDS, *MDE_PROVENANCE] + """
rediscovery_visibility_class lifecycle_case_state final_enrichment_state
enrichment_contract_version enrichment_build_version
original_classification_preserved operational_change_inferred
assurance_case_changed routing_changed priority_changed finding_created
ticket_created remediation_inferred rediscovery_delta_type
""".split(),
}

REQUIRED = {
    "asset": ["asset_key", *SCOPE_FIELDS, *SOURCE_FLAGS.values(),
              "current_observation_status", "requires_post_validation",
              "validation_reason"],
    "mart": ["asset_key", *SCOPE_FIELDS, *FEATURE_IDENTITY_FIELDS,
             "feature_contract_version", "feature_build_version",
             "feature_build_mode"],
    "feature": ["asset_key", *SCOPE_FIELDS, *FEATURE_IDENTITY_FIELDS,
                "device_scope", "network_scope", "operational_boundary",
                "operational_boundary_confidence", "management_model",
                "management_owner", "control_assurance_contract_version",
                "control_assurance_build_version", "control_assurance_build_mode"],
    "lifecycle": ["asset_key", *SCOPE_FIELDS, "evidence_state", "lifecycle_state"],
    "source": ["asset_key", *SCOPE_FIELDS, "observed_source", "source_role",
               "observed_ip", "observed_name", "serial_number",
               "last_observed_timestamp", "timestamp_basis",
               "source_active_flag", "source_stale_flag", "days_since_observed",
               "observation_age_bucket", "source_freshness_status"],
    "tool": ["asset_key", *SCOPE_FIELDS, "tool_name", "required_tool",
             "tool_record_present", "tool_active_reporting",
             "tool_stale_evidence", "last_seen_at", "expectation_basis",
             "expectation_rule_source", "coverage_status", "reporting_status"],
    "operational": ["asset_key", *SCOPE_FIELDS, "operational_population_class",
                    "tool_expectation_count"],
    "investigation": ["asset_key", *SCOPE_FIELDS, "primary_case_type"],
    "availability": [*SCOPE_FIELDS, "source_key", "source_label", "required",
                     "available", "table_name", "resolution_basis", "row_count",
                     "column_count", "read_error", "missing_source_interpretation"],
    "ip005_summary": [*SCOPE_FIELDS, "metric_name", "metric_value"],
    "post_context": ["asset_key", *SCOPE_FIELDS],
    "enterprise_context": ["asset_key", *SCOPE_FIELDS],
    "discovery": ["asset_key", *SCOPE_FIELDS, "attribution_evidence_status"],
    "disc_availability": [*SCOPE_FIELDS, "source_system", "logical_name",
                          "physical_table_name", "availability_status",
                          "required_for_discovery", "optional_source",
                          "discovery_mode", "queried_for_candidate_data",
                          "column_count", "metadata_error"],
    "disc_gaps": [*SCOPE_FIELDS, "gap_id", "gap_category", "source_system",
                  "gap_status", "affected_asset_count", "operational_impact",
                  "recommended_owner", "recommended_action", "optional_source"],
    "disc_closure": [*SCOPE_FIELDS, "status", "ready", "candidate_asset_count",
                     "disposition_asset_count", "no_usable_signal_count"],
    "mde_search": ["asset_key", *SCOPE_FIELDS, "search_eligible", *AUTHORITY_FIELDS[1:]],
    "mde": PROJECT["mde"],
}

# aliases, existing producer fallback, storage scope
# POST: producer replaces the Post output.
# SHARED: producer retains other Posts.
# HISTORY: producer retains other runs for this Post.
SPECS = {
    "asset": (["ASSET_RECONCILIATION_SCORED"], None, "POST"),
    "mart": (["IP006_ASSET_FEATURE_MART"], None, "POST"),
    "feature": (["IP006_CONTROL_ASSURANCE_FEATURE"], None, "POST"),
    "lifecycle": (["IP003_ASSET_LIFECYCLE_ROLLUP"], None, "POST"),
    "source": (["IP003_SOURCE_LIFECYCLE_OBSERVATIONS"], None, "POST"),
    "tool": (["IP002_TOOL_COVERAGE_REVIEW"], None, "POST"),
    "operational": (
        ["OPERATIONAL_ASSET_VIEW", "IP005_OPERATIONAL_ASSET_VIEW"],
        "operational_asset_view", "POST",
    ),
    "investigation": (
        ["OPERATIONAL_INVESTIGATION_UNIVERSE",
         "IP005_OPERATIONAL_INVESTIGATION_UNIVERSE"],
        "operational_investigation_universe", "POST",
    ),
    "availability": (
        ["IP005_SOURCE_AVAILABILITY"], "ip005_source_availability", "POST",
    ),
    "ip005_summary": (
        ["IP005_FINAL_SUMMARY"], "ip005_final_summary", "POST",
    ),
    "post_context": (
        ["IP003_POST_VALIDATION_TABLE"], "ip003_post_validation_table", "POST",
    ),
    "enterprise_context": (
        ["IP003_ENT_CLEANUP_TABLE"], "ip003_ent_cleanup_table", "POST",
    ),
    "discovery": (
        [], "rcaip_discovery_002_enterprise_attribution_disposition", "SHARED",
    ),
    "disc_availability": (
        [], "rcaip_discovery_002_source_availability", "SHARED",
    ),
    "disc_gaps": (
        [], "rcaip_discovery_002_enterprise_attribution_source_gaps", "SHARED",
    ),
    "disc_closure": (
        [], "rcaip_discovery_002_closure_certificate", "SHARED",
    ),
    "mde_search": (["IP006_MDE_REDISCOVERY_SEARCH_INPUT"], None, "HISTORY"),
    "mde": (["IP006_MDE_REDISCOVERY_ENRICHMENT"], None, "HISTORY"),
}
CORE = [
    "asset", "mart", "feature", "lifecycle", "source", "tool",
    "operational", "investigation", "availability", "ip005_summary",
]
DISC_GROUP = ["discovery", "disc_availability", "disc_gaps", "disc_closure"]

RUNTIME = {
    "closed_members": "RCAIP_DIAG001_UNRESOLVED_IP_CASE_MEMBER_DF",
    "branch": "RCAIP_DIAG001_BRANCH_RESUME_DF",
    "attribution": "RCAIP_DIAG001_ATTRIBUTION_QUALIFIED_DF",
    "parent": "RCAIP_DIAG001_ATTRIBUTION_NEXT_COHORT_DF",
    "intune": "RCAIP_DIAG001_CELL31_INTUNE_DISPOSITION_DF",
    "history": "RCAIP_DIAG001_CELL31_MDE_HISTORY_DEPENDENCY_DF",
    "treatment": "RCAIP_DIAG001_CELL38_TREATMENT_DF",
    "history_summary": "RCAIP_DIAG001_CELL38_HISTORY_DEPENDENCY_SUMMARY_DF",
}

LABELS = {}

def add_labels(producer, field, pairs):
    for raw, display in pairs.items():
        LABELS[(producer, field, raw)] = display

add_labels("source", "source_freshness_status", {
    "CURRENT": "Current source evidence",
    "CURRENT_BY_TIMESTAMP": "Current by recorded timestamp",
    "STALE": "Stale source evidence",
    "STALE_BY_TIMESTAMP": "Stale by recorded timestamp",
    "UNKNOWN": "Source freshness unknown",
})
add_labels("DIAG_CELL32", "c32_qualification_state", {
    "NETWORK_ATTRIBUTION_EVIDENCE_UNAVAILABLE": "Attribution evidence unavailable",
    "NO_CERTIFIED_OBSERVATION_LINK": "No certified observation link",
    "OBSERVATION_CONTEXT_INCOMPLETE": "Observation context incomplete",
    "STRUCTURAL_OBSERVATION_CONTEXT_ONLY": "Network observation only",
    "INDEPENDENT_NETWORK_ATTRIBUTION_CONTEXT": "Independent network support",
    "PARTIAL_INDEPENDENT_NETWORK_ATTRIBUTION_CONTEXT": "Partial independent support",
    "CERTIFIED_NETWORK_REVIEW_CONTEXT": "Network review context recorded",
})
for producer, prefix in (("DIAG_CELL34", "c34"), ("DIAG_CELL36", "c36")):
    add_labels(producer, prefix + "_existing_review_context", {
        "ATTRIBUTION_CONFLICT_RECORDED": "Attribution conflict recorded",
        "OTHER_CONFLICT_CONTEXT_RECORDED": "Other conflict context recorded",
        "CONFLICT_CONTEXT_INCOMPLETE": "Conflict context incomplete",
        "NO_CONFLICT_RECORDED_IN_CONSUMED_FIELDS":
            "No conflict recorded in these fields",
    })
add_labels("DIAG_CELL34", "c34_observation_context", {
    "MULTI_SOURCE_ACTIVE_OBSERVATION_CONTEXT": "Multiple active observation sources",
    "NMAP_OBSERVATION_WITH_CMDB_REGISTRATION_CONTEXT":
        "Nmap observation with CMDB context",
})
add_labels("DIAG_CELL36", "c36_observation_context", {
    "TENABLE_ACTIVE_SOURCE_OBSERVATION_ONLY": "Tenable active observation",
})
INTUNE_LABELS = {
    "INTUNE_IOS_IPADOS_MANAGED_MOBILE_CONTEXT": "Managed-mobile context",
    "WINDOWS_HISTORICAL_MDE_RECONCILIATION_CANDIDATE":
        "Windows history reconciliation candidate",
    "WINDOWS_EXISTING_MDE_CONTEXT": "Prior MDE context available",
    "INTUNE_SHARED_SERIAL_IDENTITY_REVIEW": "Shared-serial identity context",
    "PLATFORM_CONTEXT_REVIEW_REQUIRED": "Platform context review",
}
add_labels("intune", "qualification_class", INTUNE_LABELS)
add_labels("DIAG_CELL31", "c31_intune_qualification_class", INTUNE_LABELS)
GATE_LABELS = {
    UMBRELLA: "Attribution evidence continuation",
    "MOBILE_APPLICABILITY_AND_OPERATIONAL_ATTRIBUTION_CONTEXT":
        "Mobile applicability and attribution context",
    "GOVERNED_MDE_HISTORY_PREFLIGHT": "Governed MDE history preflight — parked",
    "REUSE_EXISTING_MDE_REDISCOVERY_EVIDENCE": "Reuse existing MDE evidence",
    "INTUNE_SOURCE_IDENTITY_RECONCILIATION": "Intune identity reconciliation",
    "PLATFORM_SEMANTIC_REVIEW": "Platform interpretation review",
    "POST_INDICATIVE_IDENTITY_AND_CMDB_VALIDATION":
        "Post-indicative identity/CMDB validation",
    "GOVERNED_CROSSWALK_VALIDATION": "Governed crosswalk validation",
}
add_labels("attribution", "qualification_next_gate", GATE_LABELS)
add_labels("DIAG_CELL38", "c38_qualified_continuation_gate", GATE_LABELS)
add_labels("DIAG_CELL38", "c38_customer_determination_state", {
    "NOT_YET_DETERMINED": "Customer not yet determined",
})
add_labels("attribution", "workload_promotion_state", {
    "NOT_READY_FOR_CASE_FORMATION": "Not ready for case formation",
    "NOT_READY_FOR_OPERATIONAL_WORKLOAD": "Not ready for operational workload",
})

# COMMAND ----------

class CompatibilityError(RuntimeError):
    pass


def input_read(name, stage, operation):
    """Translate only recognized read failures; unknown/fatal errors remain STOPs."""
    try:
        return operation()
    except (FileNotFoundError, PermissionError, ConnectionError, TimeoutError) as exc:
        raise CompatibilityError(
            f"{name}: {stage}: {type(exc).__name__}: {exc}"
        ) from exc
    except AnalysisException as exc:
        # Missing durable inputs are optional-group omissions; unresolved columns
        # and other analysis/programming errors must still stop publication.
        if exc.getErrorClass() not in {
            "TABLE_OR_VIEW_NOT_FOUND", "PATH_NOT_FOUND", "DELTA_TABLE_NOT_FOUND",
        }:
            raise
        raise CompatibilityError(
            f"{name}: {stage}: {type(exc).__name__}: {exc}"
        ) from exc
    except Py4JJavaError as exc:
        # Spark may wrap a source I/O failure. Require a fully recognized chain;
        # never treat JVM resource failures or arbitrary Spark errors as absence.
        wrappers = {
            "org.apache.spark.SparkException",
            "java.util.concurrent.ExecutionException",
            "java.util.concurrent.CompletionException",
            "java.io.IOException", "java.io.UncheckedIOException",
        }
        recoverable = {
            "java.io.FileNotFoundException", "java.nio.file.NoSuchFileException",
            "java.nio.file.AccessDeniedException",
            "java.net.ConnectException", "java.net.SocketTimeoutException",
            "java.net.UnknownHostException",
        }
        cause = exc.java_exception
        for _ in range(16):
            kind = cause.getClass().getName()
            if kind not in wrappers | recoverable:
                raise
            next_cause = cause.getCause()
            if next_cause is None:
                if kind not in recoverable:
                    raise
                raise CompatibilityError(
                    f"{name}: {stage}: {type(exc).__name__}: {exc}"
                ) from exc
            cause = next_cause
        raise  # Unbounded/unrecognized cause chain is not ordinary unavailability.


def selected_parent_keys(attribution, cohort_id):
    """Reproduce Cell 26's four-field signature and non-null inner-join keys."""
    fields = (
        "qualification_class", "qualification_next_gate",
        "qualification_evidence_basis", "existing_validation_responsibility",
    )
    keys = set()
    for key, row in attribution.items():
        values = [row[field] for field in fields]
        if any(value is not None and not isinstance(value, str) for value in values):
            raise CompatibilityError("Cell 26 cohort signature requires string fields")
        # concat_ws skips nulls; Cell 26's subsequent equality join excludes them.
        signature = " | ".join(value for value in values if value is not None)
        derived_id = "ATTRQ-" + hashlib.sha256(signature.encode("utf-8")).hexdigest()[:12]
        if derived_id == cohort_id and all(value is not None for value in values):
            keys.add(key)
    return keys


def build_and_publish_report002():
    # Every execution-cell retry invalidates the prior publication before reads.
    global RCAIP_REPORT002_READY
    RCAIP_REPORT002_READY = False
    for name in (
        "RCAIP_REPORT002_PUBLISHED_PATH",
        "RCAIP_REPORT002_SHA256",
        "RCAIP_REPORT002_CERTIFICATE",
    ):
        globals().pop(name, None)

    checks, manifest, omissions = [], {}, {}
    data, frames, schema_columns = {}, {}, {}
    input_cells = 0
    input_bytes = 0

    def require(condition, message):
        if not condition:
            raise CompatibilityError("REPORT 002: " + message)

    def check(name, condition):
        require(condition, name)
        checks.append({"check": name, "status": "PASS"})

    def need(name):
        require(name in globals(), "Missing existing runtime object: " + name)
        return globals()[name]

    def blank(value):
        return value is None or (isinstance(value, str) and not value.strip())

    def canonical(value):
        if hasattr(value, "asDict"):
            return canonical(value.asDict(recursive=True))
        if isinstance(value, dict):
            require(all(isinstance(k, str) for k in value),
                    "Unsupported non-string map key")
            return {k: canonical(value[k]) for k in sorted(value)}
        if isinstance(value, (list, tuple)):
            return [canonical(v) for v in value]
        if isinstance(value, datetime):
            return {"datetime": value.isoformat()}
        if isinstance(value, date):
            return {"date": value.isoformat()}
        if isinstance(value, Decimal):
            return {"decimal": str(value)}
        if isinstance(value, bytes):
            return {"bytes": value.hex()}
        if isinstance(value, float):
            require(math.isfinite(value), "Non-finite evidence number")
        require(value is None or isinstance(value, (str, bool, int, float)),
                "Unsupported evidence type: " + type(value).__name__)
        return value

    def encoded(value):
        return json.dumps(
            canonical(value), ensure_ascii=False, sort_keys=True,
            separators=(",", ":"), allow_nan=False,
        ).encode("utf-8")

    def sha(value):
        return hashlib.sha256(value).hexdigest()

    def fingerprint(rows):
        return sha("".join(sorted(sha(encoded(r)) for r in rows)).encode("ascii"))

    def key_fingerprint(keys):
        return sha(encoded(sorted(keys)))

    mode = need("RCAIP_REPORT002_MODE")
    require(mode in MODES, "Explicit supported publication mode required")
    for name in (
        "spark", "RCAIP_CONFIG_000_READY", "RCAIP_TABLE_REGISTRY_READY",
        "RCAIP_POST_NAME", "RCAIP_POST_SLUG", "RCAIP_RUN_ID",
        "RCAIP_PROFILE_HASH", "RCAIP_RESULTS_NAMESPACE",
        "RCAIP_RUN_REGISTRY_TABLE", "RCAIP_PROFILE_REGISTRY_TABLE",
        "RCAIP_TABLES", "get_rcaip_table",
    ):
        need(name)
    require(RCAIP_CONFIG_000_READY is True
            and RCAIP_TABLE_REGISTRY_READY is True,
            "Existing Config runtime is not ready")

    post, run = str(RCAIP_POST_SLUG), str(RCAIP_RUN_ID)
    namespace = str(RCAIP_RESULTS_NAMESPACE)
    require(re.fullmatch(r"[A-Za-z0-9_-]+", post) is not None,
            "Unsafe filename Post slug")
    require(re.fullmatch(r"[A-Za-z0-9_.-]+", run) is not None,
            "Unsafe filename run ID")
    require(len(post) + len(run) < 180, "Filename scope is too long")

    def quoted(table):
        parts = str(table).split(".")
        require(all(parts), "Invalid physical table name")
        return ".".join("`" + p.replace("`", "``") + "`" for p in parts)

    def pin(table):
        def read_snapshot():
            detail = spark.sql("DESCRIBE DETAIL " + quoted(table)).first().asDict()
            require(str(detail.get("format", "")).lower() == "delta",
                    table + ": expected existing Delta product")
            history = spark.sql(
                "DESCRIBE HISTORY " + quoted(table) + " LIMIT 1"
            ).first()
            require(history is not None, table + ": Delta history unavailable")
            version = int(history["version"])
            df = spark.read.option("versionAsOf", version).table(table)
            return df, {
                "physical_table": table, "delta_version": version,
                "delta_timestamp": str(history["timestamp"]),
                "delta_table_id": detail.get("id"),
            }
        return input_read(table, "pin Delta snapshot", read_snapshot)

    def bounded(name, df, required_fields=(), projection=None):
        nonlocal input_cells, input_bytes
        require(isinstance(df, DataFrame), name + ": not a DataFrame")
        require(len(df.columns) == len(set(df.columns)),
                name + ": duplicate column names")
        require(set(required_fields) <= set(df.columns),
                name + ": missing fields " +
                str(sorted(set(required_fields) - set(df.columns))))
        schema_columns[name] = list(df.columns)
        selected = (
            [c for c in dict.fromkeys(projection) if c in df.columns]
            if projection is not None else list(df.columns)
        )
        selected_df = df.select(*[
            F.col("`" + c.replace("`", "``") + "`") for c in selected
        ])
        count = input_read(name, "count projected rows", selected_df.count)
        require(count <= MAX_INPUT_ROWS, name + ": row safety limit exceeded")
        input_cells += count * len(selected)
        require(input_cells <= MAX_INPUT_CELLS, "Input cell budget exceeded")
        rows = []
        iterator = input_read(name, "open row iterator", selected_df.toLocalIterator)
        sentinel = object()
        while True:
            row = input_read(name, "collect next row", lambda: next(iterator, sentinel))
            if row is sentinel:
                break
            value = row.asDict(recursive=True)
            input_bytes += len(encoded(value))
            require(input_bytes <= MAX_INPUT_BYTES, "Input byte budget exceeded")
            rows.append(value)
        require(len(rows) == count, name + ": row count changed during collection")
        manifest.setdefault(name, {}).update({
            "row_count": count,
            "source_column_count": len(df.columns),
            "projected_columns": selected,
            "projected_schema_sha256": sha(selected_df.schema.json().encode()),
            "content_sha256": fingerprint(rows),
        })
        return rows

    def registry_rows():
        result = []
        for name, table in (
            ("run_registry", RCAIP_RUN_REGISTRY_TABLE),
            ("profile_registry", RCAIP_PROFILE_REGISTRY_TABLE),
        ):
            df, info = pin(table)
            manifest[name] = info
            field = "post_profile_key"
            require(field in df.columns, name + ": missing Post-profile key")
            df = df.filter(F.col(field) == F.lit(post))
            if name == "run_registry":
                df = df.filter(
                    (F.col("active") == F.lit(True))
                    & (F.col("run_status") == F.lit("CERTIFIED_INTAKE"))
                )
            else:
                df = df.filter(
                    (F.col("active") == F.lit(True))
                    & (F.upper(F.col("profile_status")) == F.lit("CERTIFIED"))
                )
            rows = bounded(name, df)
            require(len(rows) == 1, name + ": expected one active certified row")
            result.append(rows[0])
        return result

    run_record, profile_record = registry_rows()
    check("Existing certified run/profile binding",
          run_record["run_id"] == run
          and run_record["profile_hash"] == profile_record["profile_hash"]
          == RCAIP_PROFILE_HASH
          and run_record["results_namespace"] == namespace)

    def resolve(name):
        aliases, stem, storage = SPECS[name]
        for alias in aliases:
            if alias in RCAIP_TABLES:
                return get_rcaip_table(alias, require_exists=False)
        require(stem is not None, name + ": required registry binding absent")
        return (
            f"{namespace}.{stem}" if storage == "SHARED"
            else f"{namespace}.{post}_{stem}"
        )

    def load(name):
        table = resolve(name)
        df, info = pin(table)
        manifest[name] = dict(info, status="CHECKING", producer_contract=name)
        require(set(SCOPE_FIELDS) <= set(df.columns),
                name + ": missing row-level Post/run fields")
        storage = SPECS[name][2]
        if storage != "SHARED":
            wrong_post = input_read(name, "validate Post scope", lambda: df.filter(
                F.col("post_slug").isNull()
                | (F.col("post_slug") != F.lit(post))
            ).limit(1).count())
            require(wrong_post == 0, name + ": foreign/null Post row")
        else:
            df = df.filter(F.col("post_slug") == F.lit(post))
        if storage != "HISTORY":
            wrong_run = input_read(name, "validate run scope", lambda: df.filter(
                F.col("rcaip_run_id").isNull()
                | (F.col("rcaip_run_id") != F.lit(run))
            ).limit(1).count())
            require(wrong_run == 0, name + ": stale/null/incompatible run row")
        else:
            require(input_read(name, "validate historical null-run scope", lambda:
                               df.filter(F.col("rcaip_run_id").isNull()).limit(1).count()) == 0,
                    name + ": unscoped historical row")
            df = df.filter(F.col("rcaip_run_id") == F.lit(run))
        frames[name] = df
        projection = PROJECT.get(name)
        if projection is not None:
            projection = list(projection) + [
                c for c in ("generated_at_utc", "ip005_generated_at_utc")
                if c in df.columns
            ]
        data[name] = bounded(name, df, REQUIRED[name], projection)
        manifest[name]["status"] = "LOADED_PENDING_COMPATIBILITY"

    def index(name, keys=("asset_key",)):
        result = {}
        for row in data[name]:
            key = tuple(row.get(c) for c in keys)
            require(not any(blank(v) for v in key), name + ": blank declared key")
            require(key not in result, name + ": duplicate declared key")
            result[key] = row
        if keys == ("asset_key",):
            return {k[0]: v for k, v in result.items()}
        return result

    for name in CORE:
        load(name)
    asset = index("asset")
    asset_keys = set(asset)
    require(0 < len(asset) <= MAX_ASSETS, "Invalid/oversized Asset Plane")
    idx = {"asset": asset}
    for name in ("mart", "feature", "lifecycle", "operational", "investigation"):
        idx[name] = index(name)
        check(name + ": exact Asset Plane key set", set(idx[name]) == asset_keys)

    def same_fields(left, right, fields, description):
        require(set(left) <= set(right), description + ": missing members")
        for key, row in left.items():
            for field in fields:
                require(field in row and field in right[key],
                        description + ": required copied field absent: " + field)
                require(encoded(row[field]) == encoded(right[key][field]),
                        description + ": copied field mismatch: " + field)
        check(description, True)

    for name in ("lifecycle", "operational", "investigation"):
        copied = [
            field for field in [*SOURCE_FLAGS.values(), "current_observation_status"]
            if field in schema_columns[name]
        ]
        require(copied, name + ": no retained Asset Plane evidence to compare")
        same_fields(idx[name], asset, copied, name + ": retained Asset Plane evidence")
    same_fields(
        idx["feature"], idx["mart"], FEATURE_IDENTITY_FIELDS,
        "IP-006 retained canonical identity/source evidence",
    )
    for name, expected in (
        ("mart", {
            "feature_contract_version": "2.2.0",
            "feature_build_version": "0.1.0",
            "feature_build_mode": "SHADOW",
        }),
        ("feature", {
            "control_assurance_contract_version": "2.3.0-shadow",
            "control_assurance_build_version": "0.1.3",
            "control_assurance_build_mode": "SHADOW",
        }),
    ):
        check(name + ": existing product contract/build",
              all(all(row.get(k) == v for k, v in expected.items())
                  for row in data[name]))

    availability = index("availability", ("source_key",))
    for source_key, product in (
        ("ASSET_PLANE", "asset"), ("LIFECYCLE", "lifecycle"),
        ("TOOL_COVERAGE", "tool"),
    ):
        witness = availability.get((source_key,))
        require(witness is not None, "IP-005 source witness absent: " + source_key)
        check("IP-005 persisted source binding: " + source_key,
              witness["available"] is True
              and not witness["read_error"]
              and witness["table_name"] == manifest[product]["physical_table"]
              and witness["row_count"] == manifest[product]["row_count"]
              and witness["column_count"] == manifest[product]["source_column_count"])

    summary_index = index("ip005_summary", ("metric_name",))
    summary = {k[0]: r["metric_value"] for k, r in summary_index.items()}
    for metric, actual in (
        ("operational_asset_rows", len(asset)),
        ("distinct_asset_keys", len(asset)),
        ("tool_expectation_records", len(data["tool"])),
        ("cross_post_row_count", 0),
        ("cross_run_row_count", 0),
        ("missing_durable_output_count", 0),
    ):
        check("IP-005 final summary: " + metric,
              metric in summary and summary[metric] == actual)

    source_index = index("source", ("asset_key", "observed_source"))
    expected_source_keys = {
        (key, source)
        for key, row in asset.items()
        for source, field in SOURCE_FLAGS.items()
        if row[field] == "Y"
    }
    check("IP-003 source membership matches published Asset Plane presence",
          set(source_index) == expected_source_keys)
    allowed_freshness = {
        "CURRENT", "STALE", "CURRENT_BY_TIMESTAMP", "STALE_BY_TIMESTAMP", "UNKNOWN"
    }
    check("IP-003 existing freshness vocabulary",
          all(r["source_freshness_status"] in allowed_freshness for r in data["source"]))

    tool_index = index("tool", ("asset_key", "tool_name"))
    check("Tool coverage native members",
          all(k[0] in asset_keys for k in tool_index))
    tool_counts = Counter(r["asset_key"] for r in data["tool"])
    check("IP-005 per-asset tool denominator",
          all(r["tool_expectation_count"] == tool_counts.get(k, 0)
              for k, r in idx["operational"].items()))
    check("IP-002 published expectation authority",
          all(r["expectation_rule_source"] == "EXPLICIT_IP002_CURRENT_CONTRACT"
              for r in data["tool"]))

    for name in CORE:
        manifest[name]["status"] = "EXISTING_CONTRACT_CHECKS_PASSED"

    def exclude(names, reason):
        for name in names:
            data[name] = None
            idx.pop(name, None)
            frames.pop(name, None)
            omissions[name] = reason
            manifest.setdefault(name, {}).update({
                "status": "PRODUCT_NOT_AVAILABLE",
                "reason": reason,
                "included_row_count": None,
            })

    def optional_group(names, validator, required=False):
        try:
            for name in names:
                load(name)
            validator()
            for name in names:
                manifest[name]["status"] = "EXISTING_CONTRACT_CHECKS_PASSED"
        except CompatibilityError as exc:
            if required:
                raise
            exclude(names, str(exc))

    def validate_context(name):
        values = index(name)
        require(set(values) <= asset_keys, name + ": orphan context member")
        idx[name] = values

    optional_group(["post_context"], lambda: validate_context("post_context"))
    optional_group(["enterprise_context"],
                   lambda: validate_context("enterprise_context"))

    def validate_discovery():
        values = index("discovery")
        require(set(values) <= asset_keys, "Discovery contains orphan assets")
        existing_candidates = {
            k for k, r in idx["investigation"].items()
            if r["primary_case_type"] == "ENTERPRISE_ATTRIBUTION_CANDIDATE"
        }
        check("Discovery exact existing IP-005 candidate membership",
              set(values) == existing_candidates)
        require(len(data["disc_closure"]) == 1,
                "Discovery closure requires exactly one scoped record")
        closure = data["disc_closure"][0]
        check("Discovery existing closure",
              closure["status"] == "PASS" and closure["ready"] is True)
        check("Discovery closure denominator",
              closure["candidate_asset_count"] == closure["disposition_asset_count"]
              == len(values))
        check("Discovery no-signal closure count",
              closure["no_usable_signal_count"] == sum(
                  r["attribution_evidence_status"] == "NO_USABLE_ATTRIBUTION_SIGNAL"
                  for r in values.values()
              ))
        index("disc_gaps", ("gap_id",))
        idx["discovery"] = values

    optional_group(DISC_GROUP, validate_discovery, required=(mode == "CELL38_FULL"))

    def validate_mde():
        search, enrichment = index("mde_search"), index("mde")
        require(search and set(search) == set(enrichment)
                and set(enrichment) <= asset_keys,
                "MDE enrichment/Search Input population mismatch")
        search_df = frames["mde_search"].filter(
            F.col("search_eligible")
            & (F.col("search_eligibility_class")
               == F.lit("SEARCHABLE_MULTISOURCE_UNIQUE_NAME"))
        )
        require(input_read("mde_search", "check searchable-authority existence", lambda:
                           search_df.limit(1).count()) > 0,
                "No existing searchable MDE authority")
        # Existing producer fingerprint algorithm. No identity matching.
        authority = input_read(
            "mde_search", "aggregate search-authority fingerprint", lambda:
            search_df.select(*AUTHORITY_FIELDS)
            .withColumn(
                "_h",
                F.sha2(F.to_json(
                    F.struct(*[F.col(c) for c in AUTHORITY_FIELDS]),
                    {"ignoreNullFields": "false"},
                ), 256),
            )
            .agg(F.sha2(
                F.concat_ws("|", F.sort_array(F.collect_list("_h"))), 256
            ).alias("authority"))
            .first()["authority"]
        )
        snapshots = set()
        for key, row in enrichment.items():
            require(row["enrichment_contract_version"] == "1.0.0"
                    and row["enrichment_build_version"] == "0.1.0",
                    "Unrecognized existing MDE enrichment version")
            require(row["original_classification_preserved"] is True,
                    "MDE original classification not preserved")
            for field in (
                "operational_change_inferred", "assurance_case_changed",
                "routing_changed", "priority_changed", "finding_created",
                "ticket_created", "remediation_inferred",
            ):
                require(row[field] is False, "MDE additive-only guard failed: " + field)
            if search[key]["search_eligible"] is True:
                require(row["search_authority_sha256"] == authority,
                        "MDE search-authority fingerprint mismatch")
                require(all(not blank(row[f]) for f in MDE_PROVENANCE),
                        "MDE searchable member has incomplete provenance")
            if row["enterprise_mde_source_delta_version"] is not None:
                snapshots.add(tuple(row[f] for f in MDE_PROVENANCE[:-1]))
        require(len(snapshots) == 1, "MDE snapshot provenance is absent or mixed")
        idx["mde"] = enrichment
        manifest["mde"]["verified_existing_search_authority_sha256"] = authority

    optional_group(["mde_search", "mde"], validate_mde)

    # COMMAND ----------
    # FULL handoffs. DURABLE never inspects these globals.

    if mode == "CELL38_FULL":
        for number in (24, 25, 26, 31, 38):
            require(need(f"RCAIP_DIAG001_CELL{number}_READY") is True,
                    f"Cell {number} is not READY")
        for number in (31, 38):
            require(need(f"RCAIP_DIAG001_CELL{number}_POST_SLUG") == post
                    and need(f"RCAIP_DIAG001_CELL{number}_RUN_ID") == run,
                    f"Cell {number} runtime scope mismatch")
        parent_id = need("RCAIP_DIAG001_ATTRIBUTION_NEXT_COHORT_ID")
        require(
            need("RCAIP_DIAG001_CELL31_PARENT_COHORT_ID") == parent_id
            and need("RCAIP_DIAG001_CELL38_PARENT_COHORT_ID") == parent_id
            and need("RCAIP_DIAG001_ATTRIBUTION_NEXT_COHORT_CLASS") == NO_SIGNAL
            and need("RCAIP_DIAG001_ATTRIBUTION_NEXT_COHORT_GATE") == UMBRELLA
            and need("RCAIP_DIAG001_CELL38_NEXT_GATE") == UMBRELLA,
            "FULL selected-parent scalar contract mismatch",
        )
        for name, object_name in RUNTIME.items():
            df = need(object_name)
            manifest[name] = {"runtime_object": object_name, "status": "CHECKING"}
            if name == "closed_members":
                required_fields = ["asset_id", "post_profile_key", "assessment_run_id"]
            else:
                required_fields = [] if name == "history_summary" else ["asset_key"]
            data[name] = bounded(name, df, required_fields)
            for row in data[name]:
                if name == "closed_members":
                    # Cell 24 publishes str(RCAIP_POST_SLUG).strip().lower().
                    # Preserve its asset_id identity and case-member grain.
                    require(not blank(row["post_profile_key"])
                            and str(row["post_profile_key"]).strip().lower()
                            == str(RCAIP_POST_SLUG).strip().lower(),
                            "closed_members: Cell 24 Post scope mismatch")
                    require(row["assessment_run_id"] == RCAIP_RUN_ID,
                            "closed_members: Cell 24 run scope mismatch")
                    require(not blank(row["asset_id"]),
                            "closed_members: blank asset_id")
                    require(row["asset_id"] in asset_keys,
                            "closed_members: asset_id outside certified Asset Plane")
                for field, value in (("post_slug", post), ("rcaip_run_id", run)):
                    if field in row:
                        require(row[field] == value, name + ": runtime scope mismatch")
            if name not in ("closed_members", "history_summary"):
                idx[name] = index(name)

        branch, attr = idx["branch"], idx["attribution"]
        parent, treatment = idx["parent"], idx["treatment"]
        intune, history = idx["intune"], idx["history"]
        check("Cell 25 exact Asset Plane membership", set(branch) == asset_keys)
        require(all(not blank(r.get("next_analysis_branch"))
                    and not blank(r.get("branch_resume_state")) for r in branch.values()),
                "Cell 25 branch partition incomplete")
        closed_keys = {
            k for k, r in branch.items()
            if r["branch_resume_state"] == "CLOSED_PRIOR_IDENTITY_BRANCH"
        }
        member_keys = {r["asset_id"] for r in data["closed_members"]}
        require(not any(blank(k) for k in member_keys), "Blank prior closed member")
        check("Prior closed identity membership preserved", closed_keys == member_keys)

        same_fields(
            branch, idx["feature"], FEATURE_IDENTITY_FIELDS,
            "Cell 25 retained certified feature evidence",
        )
        branch_attr = {
            k for k, r in branch.items()
            if r["next_analysis_branch"] == "CURRENT_ATTRIBUTION_OWNERSHIP_RECONCILIATION"
            and r["branch_resume_state"] != "CLOSED_PRIOR_IDENTITY_BRANCH"
        }
        check("Cell 26 complete attribution branch", set(attr) == branch_attr)
        siblings = {}
        for key, row in attr.items():
            require(all(f in row for f in ATTR_FIELDS), "Incomplete Cell 26 schema")
            require(not blank(row["qualification_class"]), "Blank sibling class")
            siblings.setdefault(row["qualification_class"], set()).add(key)
        union = set()
        for keys in siblings.values():
            require(not union.intersection(keys), "Overlapping attribution siblings")
            union.update(keys)
        check("Attribution sibling exact union", union == set(attr))
        same_fields(attr, branch, FEATURE_IDENTITY_FIELDS,
                    "Cell 26 retained branch evidence")
        check("Complete selected Cell 26 parent membership",
              set(parent) == selected_parent_keys(attr, parent_id))
        same_fields(parent, attr, list(data["parent"][0]) if data["parent"] else [],
                    "Selected parent remains unchanged in Cell 26")
        check("Cell 38 exact selected-parent membership", set(treatment) == set(parent))
        require(parent, "Cell 38 selected parent is empty")
        same_fields(parent, treatment, list(data["parent"][0]),
                    "Cell 38 inherited parent fields")
        check("Cell 38 published parent count",
              need("RCAIP_DIAG001_CELL38_PARENT_ASSET_COUNT") == len(parent))
        require(all(r["qualification_class"] == NO_SIGNAL
                    and r["qualification_next_gate"] == UMBRELLA
                    for r in parent.values()), "Unexpected selected parent")

        check("Exact Intune preservation",
              set(intune) == {k for k, r in treatment.items()
                              if r.get("source_family") == "INTUNE_ONLY"})
        carried_history = {
            k for k, r in treatment.items() if r.get("c31_history_candidate") is True
        }
        check("Exact history overlap", carried_history == set(history)
              and set(history) <= set(intune))

        for key, row in treatment.items():
            require(
                not blank(row.get("source_family"))
                and not blank(row.get("c38_qualified_continuation_gate"))
                and row.get("c38_parent_contract_state") == "CELL26_PARENT_CONTRACT_VALIDATED"
                and row.get("c38_customer_determination_state") == "NOT_YET_DETERMINED"
                and row.get("workload_promotion_state") == "NOT_READY_FOR_CASE_FORMATION"
                and row.get("c31_parent_cohort_id") == parent_id,
                "Cell 38 treatment/customer/workload contract differs",
            )
            if key in intune:
                prior = intune[key]
                require(
                    row["c31_intune_qualification_class"] == prior["qualification_class"]
                    and row["c38_qualified_continuation_gate"] == prior["next_gate"]
                    and row["c31_history_candidate"] == prior["c31_history_candidate"]
                    and row["c31_history_dependency"] == prior["c31_history_dependency"],
                    "Intune qualification/continuation changed",
                )
            else:
                require(row["c38_qualified_continuation_gate"] == UMBRELLA
                        and row["c31_history_candidate"] is False
                        and row["c31_history_dependency"] is None,
                        "Non-Intune continuation/history contract changed")
        for row in history.values():
            require(
                row["c31_dependency_state"] == "UNRESOLVED_ENGINEERING_EVIDENCE_BOUNDARY"
                and row["c31_blocks_asset_delivery"] is False
                and row["c31_blocks_paris"] is False
                and row["c31_history_search_authorized"] is False,
                "History dependency changed or search authorized",
            )
        dimensions = [
            "c38_qualified_continuation_gate",
            "c31_history_candidate", "c31_history_dependency",
        ]
        expected = Counter(tuple(r[d] for d in dimensions) for r in treatment.values())
        observed = {}
        for row in data["history_summary"]:
            key = tuple(row[d] for d in dimensions)
            require(key not in observed, "Duplicate history-summary grain")
            require(row["c38_history_blocks_delivery"] is False,
                    "History incorrectly blocks delivery")
            observed[key] = row["asset_count"]
        check("History cross-tab exact reconciliation", dict(expected) == observed)

        for key, row in attr.items():
            if row.get("_d002_present") is True:
                require(key in idx["discovery"], "Missing durable Discovery member")
                for field in DISC_FIELDS:
                    alias = "d002_" + field
                    if alias in row and field in idx["discovery"][key]:
                        require(encoded(row[alias]) == encoded(idx["discovery"][key][field]),
                                "Cell 26/durable Discovery evidence mismatch")
        for name in RUNTIME:
            manifest[name]["status"] = "EXISTING_CONTRACT_CHECKS_PASSED"
    else:
        for name, object_name in RUNTIME.items():
            data[name] = None
            manifest[name] = {
                "runtime_object": object_name,
                "status": NOT_EXECUTED, "included_row_count": None,
            }

    # COMMAND ----------
    # Presentation-only records, labels and limitation occurrences.

    def producer(name, field):
        if name == "treatment":
            for number in (31, 32, 34, 36, 38):
                if field.startswith(f"c{number}_"):
                    return f"DIAG_CELL{number}"
        return name

    def display(name, field, value):
        p = producer(name, field)
        if isinstance(value, str):
            return LABELS.get((p, field, value), value)
        if value is None:
            return "NULL_SOURCE_VALUE"
        if isinstance(value, (dict, list, tuple)):
            return encoded(value).decode()
        return str(value)

    limitations, asset_limitations = [], {}
    limitation_keys = set()

    def limitation(key, p, field, value, ordinal=0, native_key=None):
        grain = (key, p, field, encoded(value), ordinal, native_key)
        require(grain not in limitation_keys, "Duplicate limitation presentation grain")
        limitation_keys.add(grain)
        record = {
            "asset_key": key, "Originating Product": p,
            "Source Field": field, "Raw Value": value,
            "Display Label": display(p, field, value),
            "Element Ordinal": ordinal, "Native Source Key": native_key,
            "Record Scope": "ASSET" if key is not None else "PRODUCT",
        }
        limitations.append(record)
        if key is not None:
            asset_limitations.setdefault(key, []).append(
                p + "." + field + ": " + display(p, field, value)
            )

    if mode == "CELL38_FULL":
        limitation_fields = {
            "c32_evidence_limitations", "c34_context_limitations",
            "c34_interpretation_limit", "c36_context_limitations",
            "c36_interpretation_limit",
        }
        for key, row in idx["treatment"].items():
            for field in limitation_fields:
                value = row.get(field)
                if value is None or value == "" or value == []:
                    continue
                if field == "c32_evidence_limitations" and value == "NONE_AT_SCHEMA_LEVEL":
                    continue
                values = value if isinstance(value, list) else [value]
                for ordinal, item in enumerate(values):
                    limitation(key, producer("treatment", field), field, item, ordinal)
            if row.get("c31_history_dependency") == "UNRESOLVED_ENGINEERING_EVIDENCE_BOUNDARY":
                limitation(key, "DIAG_CELL31", "c31_history_dependency",
                           row["c31_history_dependency"])
            if row.get("c32_host_attribution_boundary_state") == "NETWORK_ATTRIBUTION_EVIDENCE_UNAVAILABLE":
                limitation(key, "DIAG_CELL32", "c32_host_attribution_boundary_reason",
                           row.get("c32_host_attribution_boundary_reason"))
    for row in data["source"]:
        if row["source_freshness_status"] in {"STALE", "STALE_BY_TIMESTAMP", "UNKNOWN"}:
            limitation(
                row["asset_key"], "source", "source_freshness_status",
                row["source_freshness_status"], native_key=row["observed_source"],
            )
    for row in data["availability"]:
        if row["available"] is False or row["read_error"]:
            limitation(
                None, "availability", "published_source_boundary", row,
                native_key=row["source_key"],
            )
    if data.get("disc_availability") is not None:
        for n, row in enumerate(data["disc_availability"]):
            if row["availability_status"] != "AVAILABLE" or row["metadata_error"]:
                limitation(None, "disc_availability", "published_source_boundary", row, n)
        for row in data["disc_gaps"]:
            if row["gap_status"] not in {"NOT_OBSERVED", "NO_SEEDS", "AVAILABLE", "CLOSED"}:
                limitation(None, "disc_gaps", "published_gap", row,
                           native_key=row["gap_id"])
    for name, reason in omissions.items():
        limitation(None, name, "publication_availability", reason)
    if mode == "DURABLE_EVIDENCE_V1":
        limitation(None, "DIAG", "edition_availability", NOT_EXECUTED)

    def technical(name, record):
        output = {}
        for field, value in record.items():
            prefix = producer(name, field) + "." + field
            output[prefix + " [raw]"] = value
            p = producer(name, field)
            if isinstance(value, str) and (p, field, value) in LABELS:
                output[prefix + " [display]"] = LABELS[(p, field, value)]
        return output

    context_fields = {
        "c31_intune_qualification_class",
        "c32_qualification_state", "c32_scan_perspectives",
        "c32_network_roles", "c32_existing_network_states",
        "c34_qualification_state", "c34_observation_context",
        "c34_existing_review_context",
        "c36_qualification_state", "c36_observation_context",
        "c36_existing_review_context",
    }
    central = []
    central_index = {}
    for key in sorted(asset_keys):
        f, life = idx["feature"][key], idx["lifecycle"][key]
        a = idx.get("attribution", {}).get(key)
        t = idx.get("treatment", {}).get(key)
        b = idx.get("branch", {}).get(key)
        evidence_context = [
            "feature.attribution_confidence: " +
            display("feature", "attribution_confidence", f.get("attribution_confidence"))
        ]
        if a:
            evidence_context.append(
                "attribution.qualification_class: " + str(a["qualification_class"])
            )
        if t:
            for field in sorted(context_fields):
                if t.get(field) is not None:
                    evidence_context.append(
                        producer("treatment", field) + "." + field + ": " +
                        display("treatment", field, t[field])
                    )
        continuation = (
            display("treatment", "c38_qualified_continuation_gate",
                    t["c38_qualified_continuation_gate"]) if t else
            display("attribution", "qualification_next_gate",
                    a["qualification_next_gate"]) if a else
            b["next_analysis_branch"] if b else NOT_EXECUTED
        )
        row = {
            "Asset": f.get("canonical_name") or key,
            "Device Context": display("feature", "device_family", f.get("device_family")),
            "Current Observation": display(
                "feature", "current_observation_status", f.get("current_observation_status")
            ),
            "Lifecycle": display("lifecycle", "lifecycle_state", life["lifecycle_state"]),
            "Active Sources": f.get("source_mask_active"),
            "Stale Sources": f.get("source_mask_stale"),
            "Identity Evidence": (
                display("feature", "identity_basis", f.get("identity_basis")) + "; " +
                display("feature", "identity_strength", f.get("identity_strength"))
            ),
            "Attribution Evidence": "\n".join(evidence_context),
            "Evidence Limitation": "\n".join(asset_limitations.get(key, [])),
            "Analytical Branch": b["next_analysis_branch"] if b else NOT_EXECUTED,
            "Analytical Continuation": continuation,
            "Customer Determination": (
                display("treatment", "c38_customer_determination_state",
                        t["c38_customer_determination_state"]) if t else
                "No final customer determination established in consumed products"
            ),
            "Workload Authorization": (
                display("attribution", "workload_promotion_state",
                        a["workload_promotion_state"]) if a else
                "No workload authorization established in consumed products"
            ),
            "asset_key": key, "post_slug": post, "rcaip_run_id": run,
        }
        missing = {}
        for name in (
            "asset", "feature", "lifecycle", "operational", "investigation",
            "discovery", "mde",
        ):
            source_row = idx.get(name, {}).get(key)
            if source_row is not None:
                row.update(technical(name, source_row))
                for field, value in source_row.items():
                    if value is None:
                        missing[name + "." + field] = "NULL_SOURCE_VALUE"
                for field in PROJECT.get(name, []):
                    if field not in source_row:
                        missing[name + "." + field] = "OPTIONAL_FIELD_ABSENT"
            else:
                missing[name] = (
                    "PRODUCT_NOT_AVAILABLE" if data.get(name) is None
                    else "OUTSIDE_PRODUCT_MEMBERSHIP"
                )
        if b:
            row.update(technical("branch", {
                field: b[field] for field in ("next_analysis_branch", "branch_resume_state")
            }))
        if a:
            row.update(technical("attribution", {field: a[field] for field in ATTR_FIELDS}))
        if t:
            annotations = {
                field: value for field, value in t.items()
                if field == "source_family"
                or field.startswith(("c31_", "c32_", "c34_", "c36_", "c38_"))
            }
            row.update(technical("treatment", annotations))
            if t.get("c32_relationship_evaluation_performed") is False:
                for field in annotations:
                    if field.startswith("c32_") and field.endswith("_count"):
                        missing["DIAG_CELL32." + field] = "NOT_EVALUATED"
        if mode == "DURABLE_EVIDENCE_V1":
            missing["DIAG"] = NOT_EXECUTED
        row["Field Availability"] = missing
        central.append(row)
        central_index[key] = row
    check("Central Asset Plane denominator preserved",
          len(central) == len(asset_keys) and set(central_index) == asset_keys)

    attribution_keys = (
        set(idx["attribution"]) if mode == "CELL38_FULL"
        else set(idx.get("discovery", {}))
    )
    attribution_rows = [dict(central_index[k]) for k in sorted(attribution_keys)]
    attribution_note = (
        "Complete executed Cell 26 attribution branch"
        if mode == "CELL38_FULL" else
        "Durable Discovery candidates only; DIAG branch accounting " + NOT_EXECUTED
    )
    post_reasons, enterprise_reasons = {}, {}

    def include(target, key, reason):
        require(key in asset_keys, "Orphan context inclusion")
        target.setdefault(key, []).append(reason)

    for key, row in asset.items():
        value = row["requires_post_validation"]
        if value is True or (isinstance(value, str) and value.upper() in {"Y", "YES", "TRUE"}):
            include(post_reasons, key, "asset.requires_post_validation")
    for name, reasons in (
        ("post_context", post_reasons),
        ("enterprise_context", enterprise_reasons),
        ("mde", enterprise_reasons),
    ):
        for key in idx.get(name, {}):
            include(reasons, key, name + ": existing product membership")
    if mode == "CELL38_FULL":
        for key, row in idx["attribution"].items():
            if row["qualification_class"] == POST_CLASS:
                include(post_reasons, key, POST_CLASS)
            if row["qualification_class"] == CROSS_CLASS:
                include(enterprise_reasons, key, CROSS_CLASS)

    def context_rows(reasons, products):
        rows = []
        for key in sorted(reasons):
            row = dict(central_index[key])
            row["Inclusion Basis"] = reasons[key]
            for name in products:
                row[name + " [published context]"] = idx.get(name, {}).get(key)
            rows.append(row)
        return rows

    def native_rows(name, sort_field):
        rows = []
        for source_row in data[name]:
            row = {"Asset": central_index[source_row["asset_key"]]["Asset"]}
            row.update(technical(name, source_row))
            rows.append(row)
        return sorted(rows, key=lambda r: (
            str(r.get(name + ".asset_key [raw]")),
            str(r.get(name + "." + sort_field + " [raw]")),
        ))

    source_rows = native_rows("source", "observed_source")
    tool_rows = native_rows("tool", "tool_name")
    generated = datetime.now(timezone.utc).isoformat()
    overview = []

    def metric(section, measure, actual=None, reference=None, status="AVAILABLE", detail=None):
        overview.append({
            "Section": section, "Measure": measure, "Actual Runtime Value": actual,
            "Historical Reference Only": reference, "Status": status, "Detail": detail,
        })

    for name, value in (
        ("Publication mode", mode), ("Post", post), ("Run", run),
        ("Generated UTC", generated), ("Certified profile hash", RCAIP_PROFILE_HASH),
        ("Asset denominator", len(asset)),
    ):
        metric("Publication", name, value)
    metric("Interpretation", "Authority",
           "Evidence publication only; unresolved attribution is not workload")
    metric("Provenance", "Certification scope",
           "Existing producer contracts and cross-product consistency checks",
           detail="No uniform producer-code/complete-upstream-snapshot attestation is claimed")

    if mode == "CELL38_FULL":
        dimensions = [
            ("branch", ["next_analysis_branch", "branch_resume_state"]),
            ("attribution", ["qualification_class"]),
            ("treatment", ["source_family"]),
            ("treatment", ["c38_qualified_continuation_gate"]),
        ]
        for name, fields in dimensions:
            counts = Counter(tuple(r[f] for f in fields) for r in data[name])
            for key, count in sorted(counts.items(), key=lambda item: str(item[0])):
                metric(name, " / ".join(map(str, key)), count)
        metric("History overlap", "Distinct dependency members", len(idx["history"]),
               detail="Overlap; never added to exclusive treatment totals")
    else:
        for section in ("Cell 25 branches", "Cell 26 attribution", "Cell 38 treatment"):
            metric(section, "Runtime accounting", status=NOT_EXECUTED)
    for name, fields in (
        ("feature", ["source_mask_active"]),
        ("feature", ["source_mask_stale"]),
        ("source", ["observed_source", "source_freshness_status"]),
        ("tool", ["tool_name", "coverage_status"]),
    ):
        counts = Counter(tuple(r.get(f) for f in fields) for r in data[name])
        for key, count in sorted(counts.items(), key=lambda item: str(item[0])):
            metric("Coverage: " + name, " / ".join(map(str, key)), count)
    metric("Limitations", "Occurrences", len(limitations))
    metric("Limitations", "Distinct affected assets",
           len({r["asset_key"] for r in limitations if r["asset_key"] is not None}))

    # Historical values are display data only; never acceptance predicates.
    # Publication copy: operational reference-only display totals excluded.

    manifest["asset"]["asset_key_sha256"] = key_fingerprint(asset_keys)
    for name, record in manifest.items():
        metric("Product manifest", name, record.get("row_count"),
               status=record.get("status", "PINNED_REGISTRY"), detail=record)
    for item in checks:
        metric("Certification", item["check"], status=item["status"])
    for name in SHEET_NAMES[1:]:
        metric("Navigation", name, detail="Open worksheet")

    package = {
        "Overview": (overview, "Publication metadata and denominator-aware accounting"),
        "Asset Review Detail": (central, "One row per certified Asset Plane identity"),
        "Attribution Evidence Review": (attribution_rows, attribution_note),
        "Post Validation Context": (
            context_rows(post_reasons, ["post_context"]),
            "Existing validation context; no assigned Post workload implied",
        ),
        "Enterprise Crosswalk Context": (
            context_rows(enterprise_reasons, ["enterprise_context", "mde"]),
            "Existing reconciliation context; no approved referral implied",
        ),
        "Evidence Limitations": (
            limitations, "Explicit limitations/dependencies; not an action queue",
        ),
        "Source Evidence": (source_rows, "Native asset × observed-source grain"),
        "Tool Coverage": (tool_rows, "Native asset × tool-expectation grain"),
    }

    # COMMAND ----------
    # Excel-safe presentation encoding and established publication method.

    escape_count = 0

    def escape_text(value):
        nonlocal escape_count
        parts = []
        for character in value:
            number = ord(character)
            if character == "\\":
                parts.append("\\\\")
            elif (
                number < 32 and number not in (9, 10, 13)
                or 0xD800 <= number <= 0xDFFF
                or number in (0xFFFE, 0xFFFF)
            ):
                parts.append(
                    f"\\x{number:02X}" if number <= 255 else f"\\u{number:04X}"
                )
                escape_count += 1
            else:
                parts.append(character)
        result = "".join(parts)
        require(len(result) <= 32767, "Excel cell exceeds limit; no truncation permitted")
        return result

    def excel_value(value):
        if value is None or isinstance(value, bool):
            return value
        if isinstance(value, int):
            return value if abs(value) < 10**15 else str(value)
        if isinstance(value, float):
            require(math.isfinite(value), "Non-finite Excel number")
            # Python's shortest round-trip representation preserves the exact
            # finite float, including signed zero, without Excel numeric rounding.
            # Evidence fingerprints are computed before presentation conversion.
            return repr(value)
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if isinstance(value, (dict, list, tuple, bytes)):
            # Canonical JSON is already an explicit reversible representation.
            return escape_text(encoded(value).decode())
        return escape_text(str(value))

    def columns(rows, initial):
        result = list(initial)
        seen = set(result)
        for row in rows:
            for field in row:
                if field not in seen:
                    result.append(field)
                    seen.add(field)
        return result

    workbook = Workbook()
    workbook.remove(workbook.active)
    validation, sheet_manifest = {}, []
    export_cells = 0
    for number, name in enumerate(SHEET_NAMES, 1):
        rows, description = package[name]
        if name in SHEET_NAMES[1:5]:
            initial = OPERATOR_COLUMNS + ["asset_key", "post_slug", "rcaip_run_id"]
        elif name == "Overview":
            initial = [
                "Section", "Measure", "Actual Runtime Value",
                "Historical Reference Only", "Status", "Detail",
            ]
        elif name == "Evidence Limitations":
            initial = [
                "asset_key", "Originating Product", "Source Field",
                "Raw Value", "Display Label", "Element Ordinal",
                "Native Source Key", "Record Scope",
            ]
        else:
            native = "source" if name == "Source Evidence" else "tool"
            initial = ["Asset"] + [native + "." + c + " [raw]"
                                   for c in schema_columns[native]]
        headers = columns(rows, initial)
        require(len(headers) <= 16384 and len(rows) + 5 <= 1048576,
                name + ": Excel dimensions exceeded")
        export_cells += (len(rows) + 5) * len(headers)
        require(export_cells <= MAX_EXPORT_CELLS, "Export cell safety limit exceeded")
        ws = workbook.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.cell(1, 1, name)
        ws.cell(2, 1, description)
        ws.cell(3, 1, f"Post={post}; Run={run}; Mode={mode}; Generated={generated}")
        ws.cell(4, 1, "Overview")
        ws.cell(4, 1).hyperlink = "#'Overview'!A1"
        ws.cell(4, 1).style = "Hyperlink"
        ws.row_dimensions[1].height = 25
        ws.row_dimensions[2].height = 32
        ws.row_dimensions[5].height = 42
        ws.cell(1, 1).font = Font(name="Aptos", size=16, bold=True, color="FFFFFF")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="17365D")
        for r in (2, 3):
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        for column, header in enumerate(headers, 1):
            cell = ws.cell(5, column, header)
            cell.data_type = "s"
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            colour = "66577A" if header in {
                "Customer Determination", "Workload Authorization"
            } else "1F4E78"
            cell.fill = PatternFill("solid", fgColor=colour)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(column)].width = (
                42 if header in {
                    "Attribution Evidence", "Evidence Limitation",
                    "Analytical Continuation", "Customer Determination",
                    "Workload Authorization", "Detail",
                } else 27
            )
        expected = []
        for row_number, record in enumerate(rows, 6):
            values = [excel_value(record.get(h)) for h in headers]
            expected.append(values)
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_number, column, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Aptos", size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if row_number % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F5F8")
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "#,##0.############"
            if name == "Overview" and record.get("Section") == "Navigation":
                ws.cell(row_number, 2).hyperlink = "#'" + record["Measure"] + "'!A1"
                ws.cell(row_number, 2).style = "Hyperlink"
        last = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A5:{last}{max(5, len(rows) + 5)}"
        ws.freeze_panes = "A6" if name == "Overview" else "C6"
        if name in SHEET_NAMES[1:5] and len(headers) > len(OPERATOR_COLUMNS):
            ws.column_dimensions.group(
                get_column_letter(len(OPERATOR_COLUMNS) + 1), last, hidden=True
            )
        if rows:
            table = Table(displayName=f"Report002Table{number}",
                          ref=f"A5:{last}{len(rows) + 5}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=False
            )
            ws.add_table(table)
        validation[name] = (headers, expected)
        sheet_manifest.append({
            "sheet": name, "order": number, "row_count": len(rows),
            "column_count": len(headers), "description": description,
        })

    filename = f"RCAIP_{post}_{run}_{mode}_Asset_Enterprise_Visibility.xlsx"
    directory = None
    for root in (Path("/local_disk0/tmp"), Path(tempfile.gettempdir())):
        try:
            root.mkdir(parents=True, exist_ok=True)
            directory = Path(tempfile.mkdtemp(
                prefix=f"rcaip_report002_{post}_", dir=str(root)
            ))
            break
        except OSError:
            continue
    require(directory is not None, "No writable driver-local directory")
    local_path = directory / filename
    workbook.save(local_path)
    workbook.close()
    require(0 < local_path.stat().st_size <= MAX_FILE_BYTES, "Invalid/oversized workbook")

    reopened = load_workbook(local_path, read_only=True, data_only=False)
    try:
        require(reopened.sheetnames == SHEET_NAMES, "Saved worksheet order differs")
        for name in SHEET_NAMES:
            ws = reopened[name]
            headers, expected = validation[name]
            require(ws.max_row == len(expected) + 5
                    and ws.max_column == len(headers),
                    name + ": saved dimensions differ")
            actual_headers = [c.value for c in next(ws.iter_rows(min_row=5, max_row=5))]
            require(actual_headers == headers, name + ": saved headers differ")
            if expected:
                actual_rows = ws.iter_rows(min_row=6, max_row=5 + len(expected))
                seen = 0
                for actual, wanted in zip(actual_rows, expected):
                    seen += 1
                    for cell, value in zip(actual, wanted):
                        require(cell.value == (None if value == "" else value),
                                name + ": sanitized round-trip value differs")
                        require(cell.data_type not in ("f", "e"),
                                name + ": formula/error cell introduced")
                require(seen == len(expected), name + ": saved row parity differs")
    finally:
        reopened.close()

    # Refuse stale publication if a consumed durable binding changed during build.
    for name, record in manifest.items():
        if record.get("status") == "PRODUCT_NOT_AVAILABLE":
            continue
        if "physical_table" in record:
            _, current = pin(record["physical_table"])
            require(current["delta_version"] == record["delta_version"]
                    and current["delta_table_id"] == record["delta_table_id"],
                    name + ": durable product changed during workbook build")

    content = local_path.read_bytes()
    workbook_sha = sha(content)
    client = WorkspaceClient()
    principal = client.current_user.me().user_name
    require(isinstance(principal, str) and principal.strip(),
            "Authenticated Workspace principal unavailable")
    workspace_directory = f"/Users/{principal}/RCAIP_Reports/{post}"
    published_path = workspace_directory + "/" + filename
    client.workspace.mkdirs(workspace_directory)
    client.workspace.upload(
        path=published_path, content=io.BytesIO(content),
        format=ImportFormat.AUTO, overwrite=True,
    )
    with client.workspace.download(published_path, format=ExportFormat.AUTO) as stream:
        published = stream.read(MAX_FILE_BYTES + 1)
    require(len(published) <= MAX_FILE_BYTES and sha(published) == workbook_sha,
            "Published bytes differ; no PASS certificate issued")

    publication_id = sha(encoded([post, run, mode, workbook_sha]))
    certificate = {
        "report": "RCAIP_REPORT_002",
        "publication_id": publication_id,
        "publication_mode": mode,
        "post_slug": post,
        "rcaip_run_id": run,
        "profile_hash": RCAIP_PROFILE_HASH,
        "published_path": published_path,
        "workbook_sha256": workbook_sha,
        "published_sha256": sha(published),
        "generated_at_utc": generated,
        "certification_status": "PASS",
        "certification_scope": "EXISTING_RCAIP_CONTRACT_AND_PUBLICATION_CHECKS",
        "upstream_code_snapshot_attestation_claimed": False,
        "escaped_illegal_character_count": escape_count,
        "asset_count": len(asset_keys),
        "asset_key_sha256": key_fingerprint(asset_keys),
        "checks": checks,
        "new_analytical_authority": False,
    }
    metadata_schema = (
        "publication_id string, post_slug string, rcaip_run_id string, "
        "publication_mode string, generated_at_utc string, payload_json string"
    )
    for stem, payload in (
        ("report002_workbook_manifest", {
            "sources": manifest, "sheets": sheet_manifest,
            "omissions": omissions, "workbook_sha256": workbook_sha,
        }),
        ("report002_workbook_certificate", certificate),
    ):
        metadata = spark.createDataFrame(
            [(publication_id, post, run, mode, generated, encoded(payload).decode())],
            schema=metadata_schema,
        )
        metadata.write.format("delta").mode("append").saveAsTable(
            f"{namespace}.{post}_{stem}"
        )
    return published_path, workbook_sha, certificate


(
    RCAIP_REPORT002_PUBLISHED_PATH,
    RCAIP_REPORT002_SHA256,
    RCAIP_REPORT002_CERTIFICATE,
) = build_and_publish_report002()
RCAIP_REPORT002_READY = True
print("REPORT 002: PASS")
print("Publication:", RCAIP_REPORT002_PUBLISHED_PATH)
print("SHA-256:", RCAIP_REPORT002_SHA256)
print("NO NEW ANALYTICAL AUTHORITY INTRODUCED")
