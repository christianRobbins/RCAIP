"""Additive retained IP-population presentation; no producer execution or Spark.

The four Cell 24 sheets remain owned by the accepted renderer. This companion
reads saved decisions, never repeats identity matching or assigns new customers.
Would this exact logic run for Paris? YES: scope is explicit, counts are evidence.
"""
import ast
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re

import rcaip_cell24_retained as retained

consumer = retained.consumer
require = consumer.require
# Reviewed repository executable ASTs, using retained.producer_fingerprint.
FINGERPRINTS = {
    16: '75566cd187bef0ea0dd7dba7a328c87a8bdc00322bcebb3a8d4b978112fe011d',
    17: 'b797e260376321794725d5642f5120a955eb571565c071ad8ef8c4f0b2ce473e',
    20: '5cc9f8df861486b7623e881c6c2bd6e343fda1ea923fe54377ad2624b1d2cb16',
    21: '41ed29c6b6d88798727751464268c085f6c851a8366b7b9cc7222de8606f629e',
    22: '71f3b3f27f3f8e968ebbfa21e596c7b0714dfe914373a73b72f54345d897dbb7',
    23: 'd6405cbd928b86dbc6043390e9f2d51c5ab1f2cef22dcc82a917cc362707c43c',
}
PRODUCTS = {
    16: ('active_unresolved_ip_member_detail_df',),
    17: ('tenable_uuid_member_detail_df',),
    20: ('linked_cohort_member_detail_df', 'no_link_context_df'),
    21: ('no_link_qualified_member_detail_df',),
    22: ('temporal_member_detail_df', 'currentness_validation_df', 'residual_current_enumeration_df'),
    23: ('customer_qualification_member_df',),
}
SHEETS = ['RCSO Summary', 'IP Population Review', 'Case review', 'Case evidence', 'Members', 'Branch closure']
FAMILIES = ('Stronger UUID relationship', 'Stronger exact-IP relationship',
            'Device / recency qualification', 'Currentness validation',
            'Current residual identity', 'Temporal evidence / timing review')

# Cell 16/17 TENABLE_NORM selects -> Cell 20 candidate_tenable_df select.
# These are direct inherited aliases, not transformed features. The IP uses
# coalesce separately; normalized UUID group keys are checked separately. Stale
# flag/DNS/hostname lack a counterpart in BOTH displayed products and are not
# invented. Cell 20 -> 21/22 shared-name checks continue the evidence chain.
INHERITED_ALIASES = (
    ('tenable_host_uuid', 'candidate_tenable_uuid'),
    ('tenable_acr_score', 'candidate_tenable_acr_score'),
    ('tenable_key_drivers', 'candidate_tenable_key_drivers'),
    ('tenable_organization_name', 'candidate_tenable_organization_name'),
    ('tenable_event_time', 'candidate_tenable_event_time'),
    ('tenable_last_observed_timestamp', 'candidate_tenable_last_observed'),
    ('tenable_days_since_observed', 'candidate_tenable_age_days'),
    ('tenable_observation_age_bucket', 'candidate_tenable_age_bucket'),
)

# Code-to-code contracts from the pinned Cell 21 withColumn expressions.
# Validate saved decisions ONLY: never evaluate raw features to classify a row.
QUALIFICATION_DECISIONS = {
    'MULTI_IP_SOURCE_LOCAL_IDENTITY_NO_CORROBORATION': ('SOURCE_LOCAL_LIFECYCLE_RECONCILIATION_BEFORE_ENUMERATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'STALE_SINGLE_SOURCE_IDENTITY_LIFECYCLE_REVIEW': ('LIFECYCLE_RECONCILIATION_BEFORE_ENUMERATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'TENABLE_PRINT_DEVICE_CONTEXT_CANDIDATE': ('DEVICE_CLASS_AND_INVENTORY_VALIDATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'TENABLE_INFRASTRUCTURE_DEVICE_CONTEXT_CANDIDATE': ('DEVICE_CLASS_AND_INVENTORY_VALIDATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'FRESH_GENERAL_PURPOSE_COMPUTE_UNCORROBORATED': ('TARGETED_IDENTITY_ENUMERATION', 'RESIDUAL_IDENTITY_REVIEW_CANDIDATE'),
    'AGING_GENERAL_PURPOSE_COMPUTE_UNCORROBORATED': ('RECENCY_VALIDATION_THEN_TARGETED_ENUMERATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'FRESH_UNKNOWN_DEVICE_SINGLE_SOURCE': ('TARGETED_DEVICE_CONTEXT_ENUMERATION', 'RESIDUAL_DEVICE_CONTEXT_REVIEW_CANDIDATE'),
    'AGING_UNKNOWN_DEVICE_SINGLE_SOURCE': ('RECENCY_VALIDATION_OR_MONITORING_BEFORE_ENUMERATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'OTHER_DEVICE_CONTEXT_REVIEW': ('MODEL_OR_DEVICE_CONTEXT_REVIEW', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
}
# Cell 20 source_local_link_df / ip_cohort_assignment_df: family, lane, status.
LINK_DECISIONS = {
    'TENABLE_UUID_ALREADY_REPRESENTED_BY_STRONGER_ASSET': ('SOURCE_LOCAL_IDENTITY_MODEL_RECONCILIATION', 'RCAIP_ENGINEERING', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'TENABLE_UUID_LINKS_MULTIPLE_STRONGER_ASSETS_REVIEW': ('SOURCE_LOCAL_IDENTITY_MODEL_RECONCILIATION', 'RCAIP_ENGINEERING', 'MODEL_REVIEW_REQUIRED'),
    'MULTI_TARGET_IP_ASSIGNMENT_RECONCILIATION': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_DATA_AND_SERVICE_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'CURRENT_TENABLE_IP_IDENTITY_RECONCILIATION': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_SERVICE_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'STALE_TENABLE_IP_IDENTITY_RECONCILIATION': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_SERVICE_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'CURRENT_TARGET_INVENTORY_IP_REASSIGNMENT': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_DATA_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'STALE_TENABLE_TARGET_INVENTORY_IP_REASSIGNMENT': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_DATA_AND_SERVICE_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'CURRENT_TARGET_INVENTORY_IP_VALIDATION': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_DATA_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'NONCURRENT_INVENTORY_IP_LIFECYCLE_RECONCILIATION': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'ENTERPRISE_DATA_RECONCILIATION', 'RECONCILIATION_PATTERN_SUPPORTED'),
    'LINKED_RELATIONSHIP_MODEL_REVIEW_REQUIRED': ('ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION', 'RCAIP_ENGINEERING_REVIEW', 'MODEL_REVIEW_REQUIRED'),
}
TEMPORAL_DECISIONS = {
    'NO_MATCHING_CERTIFIED_NMAP_COLLECTION': ('TEMPORAL_EVIDENCE_LIMITATION_REVIEW', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'MATCHING_NMAP_PROVENANCE_TIMESTAMP_UNAVAILABLE': ('TEMPORAL_EVIDENCE_LIMITATION_REVIEW', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'TENABLE_LAST_OBSERVED_TIMESTAMP_UNAVAILABLE': ('TEMPORAL_EVIDENCE_LIMITATION_REVIEW', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'LATER_NMAP_COLLECTION_AFTER_TENABLE_LAST_OBSERVED': ('CURRENTNESS_VALIDATION_BEFORE_IDENTITY_ENUMERATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'TENABLE_TIMESTAMP_OVERLAPS_NMAP_COLLECTION_WINDOW': ('COLLECTION_TIMING_REVIEW', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
    'TENABLE_OBSERVED_AFTER_ALL_RELEVANT_NMAP_COLLECTIONS': ('FOCUSED_IDENTITY_DEVICE_ENUMERATION', 'RESIDUAL_CURRENT_IDENTITY_REVIEW_CANDIDATE'),
    'TEMPORAL_RELATIONSHIP_REVIEW_REQUIRED': ('TEMPORAL_EVIDENCE_LIMITATION_REVIEW', 'NOT_READY_FOR_WORKLOAD_PROMOTION'),
}

# Display text is scoped by producer AND field. Unknown codes never receive a
# guessed benign explanation. Original codes and full native records remain.
DISPLAY = {
    (20, 'explanation_class'): {
        'TENABLE_UUID_ALREADY_REPRESENTED_BY_STRONGER_ASSET': 'The Tenable UUID is already represented by a stronger asset identity; identity precedence still needs review.',
        'TENABLE_UUID_LINKS_MULTIPLE_STRONGER_ASSETS_REVIEW': 'The Tenable UUID relates to multiple stronger assets; the identity model requires review.',
        'MULTI_TARGET_IP_ASSIGNMENT_RECONCILIATION': 'The IP relates to multiple stronger asset identities; current, historical and duplicate assignments need reconciliation.',
        'CURRENT_TENABLE_IP_IDENTITY_RECONCILIATION': 'Different current or recent Tenable identities share the IP; their source identities need reconciliation.',
        'STALE_TENABLE_IP_IDENTITY_RECONCILIATION': 'A current Tenable identity shares the IP with stale Tenable identity evidence; lifecycle reconciliation remains.',
        'CURRENT_TARGET_INVENTORY_IP_REASSIGNMENT': 'The stronger asset has current Tenable evidence at another IP; the inventory assignment needs reconciliation.',
        'STALE_TENABLE_TARGET_INVENTORY_IP_REASSIGNMENT': 'The stronger asset has stale Tenable evidence at another IP; inventory and source lifecycle need reconciliation.',
        'CURRENT_TARGET_INVENTORY_IP_VALIDATION': 'An inventory IP link to a currently observed stronger asset needs validation before it can support identity.',
        'NONCURRENT_INVENTORY_IP_LIFECYCLE_RECONCILIATION': 'The inventory asset has no current observation while a different Tenable identity uses the associated IP; lifecycle reconciliation remains.',
        'LINKED_RELATIONSHIP_MODEL_REVIEW_REQUIRED': 'The retained link requires model review before customer treatment; no resolved or benign outcome is asserted.',
    },
    (21, 'member_qualification_class'): {
        'MULTI_IP_SOURCE_LOCAL_IDENTITY_NO_CORROBORATION': 'Several IP records belong to one source-local identity without stronger corroboration.',
        'STALE_SINGLE_SOURCE_IDENTITY_LIFECYCLE_REVIEW': 'The single-source identity has stale observation evidence; lifecycle review remains.',
        'TENABLE_PRINT_DEVICE_CONTEXT_CANDIDATE': 'Tenable supplied a print-device hint; the device class and inventory still need validation.',
        'TENABLE_INFRASTRUCTURE_DEVICE_CONTEXT_CANDIDATE': 'Tenable supplied an infrastructure or specialty-device hint; device and inventory validation remains.',
        'FRESH_GENERAL_PURPOSE_COMPUTE_UNCORROBORATED': 'A recently observed general-purpose compute hint lacks stronger identity corroboration.',
        'AGING_GENERAL_PURPOSE_COMPUTE_UNCORROBORATED': 'An aging general-purpose compute observation lacks stronger corroboration; recency needs validation.',
        'FRESH_UNKNOWN_DEVICE_SINGLE_SOURCE': 'A recent single-source observation has no device-type hint; device context remains unresolved.',
        'AGING_UNKNOWN_DEVICE_SINGLE_SOURCE': 'An aging single-source observation has no device-type hint; recency and device context remain unresolved.',
        'OTHER_DEVICE_CONTEXT_REVIEW': 'The recorded device context requires model or device review before further treatment.',
    },
    (21, 'member_next_gate'): {
        'SOURCE_LOCAL_LIFECYCLE_RECONCILIATION_BEFORE_ENUMERATION': 'Reconcile the source-local identity lifecycle before identity enumeration.',
        'LIFECYCLE_RECONCILIATION_BEFORE_ENUMERATION': 'Reconcile lifecycle evidence before identity enumeration.',
        'DEVICE_CLASS_AND_INVENTORY_VALIDATION': 'Validate the device class and inventory context; the hint alone is not a confirmed device identity.',
        'TARGETED_IDENTITY_ENUMERATION': 'Perform the governed targeted identity review after its remaining qualification gates.',
        'RECENCY_VALIDATION_THEN_TARGETED_ENUMERATION': 'Validate observation recency, then perform targeted identity review as supported.',
        'TARGETED_DEVICE_CONTEXT_ENUMERATION': 'Perform targeted device-context review; source absence alone does not establish a control failure.',
        'RECENCY_VALIDATION_OR_MONITORING_BEFORE_ENUMERATION': 'Validate recency or monitor before promoting identity enumeration.',
        'MODEL_OR_DEVICE_CONTEXT_REVIEW': 'Review the model or device context before further treatment.',
    },
    (22, 'temporal_next_gate'): {
        'CURRENTNESS_VALIDATION_BEFORE_IDENTITY_ENUMERATION': 'Validate whether the observation remains current before requesting identity enumeration; non-observation is not proof of absence.',
        'FOCUSED_IDENTITY_DEVICE_ENUMERATION': 'Proceed to focused identity/device review using the retained customer-qualified action.',
        'COLLECTION_TIMING_REVIEW': 'Review overlapping collection timing before interpreting the identity evidence.',
        'TEMPORAL_EVIDENCE_LIMITATION_REVIEW': 'Review the missing or insufficient temporal evidence; no workload promotion is supported by this gate.',
    },
}
PROPOSED_LABEL = 'PROPOSED RESPONSE - RCSO REVIEW PENDING'
TIME_BASIS = 'Retained Tenable last-observed text, exactly as exported. No timezone inferred; not report-generation time.'


def _decision(row, cell, primary, companions, allowed, required=()):
    for field in (primary, *companions, *required):
        require(consumer.text(row.get(field)), f'Cell {cell}: missing decision field {field}')
    require(row[primary] in allowed, f'Cell {cell}: unsupported {primary}')
    require(tuple(row[f] for f in companions) == allowed[row[primary]],
            f'Cell {cell}: contradictory decision combination for {primary}')


def display_code(cell, field, value):
    require(value in DISPLAY.get((cell, field), {}), f'Cell {cell}: unsupported display value for {field}')
    return DISPLAY[cell, field][value]


def source_columns(source, product):
    """Read the actual select contract AFTER its pinned source hash is checked."""
    matches = []
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == product for t in node.targets):
            selects = [n for n in ast.walk(node.value) if isinstance(n, ast.Call)
                       and isinstance(n.func, ast.Attribute) and n.func.attr == 'select']
            require(len(selects) == 1 and all(isinstance(a, ast.Constant) and isinstance(a.value, str)
                                            for a in selects[0].args), 'Unsupported saved select contract: ' + product)
            matches.append([a.value for a in selects[0].args])
    require(len(matches) == 1, 'Ambiguous saved select contract: ' + product)
    return matches[0]


def _field(saved, label):
    values = re.findall(r'^' + re.escape(label) + r':\s*([^\r\n]+)$', saved, re.M)
    require(len(values) == 1, 'Missing/ambiguous saved field: ' + label)
    return values[0].strip()


def _cell_products(book, number, post, run):
    marker = 'RCAIP DIAG 001 - CELL ' + str(number) + (' v1.1' if number == 22 else '') + ' CERTIFICATION'
    candidates = [c for c in book['cells'] if c.get('cell_type') == 'code'
                  and marker in retained._joined(c.get('source', ''))]
    require(len(candidates) == 1, f'Cell {number}: missing/ambiguous producer')
    cell = candidates[0]; source = retained._joined(cell['source'])
    require(retained.producer_fingerprint(source) == FINGERPRINTS[number], f'Cell {number}: changed executable source')
    schemas = {p: source_columns(source, p) for p in PRODUCTS[number]}
    outputs = cell.get('outputs'); require(isinstance(outputs, list), f'Cell {number}: missing outputs')
    saved = ''; found = {}; references = {}
    for index, output in enumerate(outputs):
        require(isinstance(output, dict) and output.get('output_type') != 'error', f'Cell {number}: saved error')
        if output.get('output_type') == 'stream':
            require(output.get('name') == 'stdout', f'Cell {number}: unexpected stream')
            saved += retained._joined(output.get('text', ''))
        payload = output.get('metadata', {}).get(retained.MIME)
        if payload is None: continue
        require(isinstance(payload, dict), f'Cell {number}: invalid table envelope')
        fields = payload.get('schema')
        if isinstance(fields, dict):
            require(fields.get('type') == 'struct', f'Cell {number}: unsupported schema')
            fields = fields.get('fields')
        require(isinstance(fields, list) and all(isinstance(f, dict) and isinstance(f.get('name'), str)
                    and 'type' in f for f in fields), f'Cell {number}: invalid fields')
        columns = [f['name'] for f in fields]
        require(len(columns) == len(set(columns)), f'Cell {number}: duplicate columns')
        names = [p for p, cols in schemas.items() if columns == cols]
        if not names: continue  # Other producer summaries are not input products.
        name = names[0]
        require(name not in found, 'Duplicate product: ' + name)
        require(payload.get('overflow') is False, 'Overflow/missing completeness: ' + name)
        kinds = [retained._schema_type(f['type']) for f in fields]
        rows = payload.get('data')
        require(isinstance(rows, list) and len(rows) <= 100000, 'Missing/oversized rows: ' + name)
        records = []
        for row in rows:
            require(isinstance(row, list) and len(row) == len(columns), 'Positional row mismatch: ' + name)
            for value, kind in zip(row, kinds): retained._typed(value, kind)
            records.append(dict(zip(columns, row)))
        found[name] = records
        references[name] = {'producer_cell': number, 'output_index': index, 'schema': fields}
    require(set(found) == set(schemas), f'Cell {number}: missing required product/schema')
    require(saved.count(marker) == 1 and _field(saved, 'Status') == 'PASS', f'Cell {number}: certification not unambiguous PASS')
    require(_field(saved, 'Post').strip().lower() == post.strip().lower(), f'Cell {number}: foreign Post')
    require(_field(saved, 'Run ID') == run, f'Cell {number}: foreign run')
    return found, saved, {'source_ast_sha256': FINGERPRINTS[number], 'execution_count': cell.get('execution_count'),
                         'cell_metadata': cell.get('metadata', {}), 'products': references,
                         'saved_certification': saved}


def _index(rows, key, label):
    result = {}
    for row in rows:
        value = row[key]
        require(consumer.text(value) and value not in result, label + ': duplicate/blank identity')
        result[value] = row
    return result


def _same(left, right, label, aliases=()):
    # Decimal spelling/sign and temporal text must agree, not just numeric equality.
    for a, b in [(k, k) for k in left.keys() & right.keys()] + list(aliases):
        require(retained._json(left[a]) == retained._json(right[b]), label + ': conflicting ' + a)


def _group(row):
    """Exact Cell 17/20 key recipe; association only, no new identity matching."""
    uuid = row['tenable_host_uuid']
    require(uuid is None or isinstance(uuid, str), 'Invalid source UUID')
    return 'TENABLE_UUID::' + uuid.strip(' ').lower() if uuid and uuid.strip(' ') else 'NO_UUID::' + row['asset_key']


def _ip(row):
    # Match the producer's coalesce, which does not treat an empty string as null.
    return row['canonical_ip'] if row['canonical_ip'] is not None else row['tenable_ip']


def load_population(path, *, expected_post, expected_run, proposed_asset_keys=()):
    cell24 = retained.load_retained_cell24(path, expected_post=expected_post, expected_run=expected_run)
    # Reuse the reviewed parser, including exact duplicate-output handling. Compare
    # bytes with the first read so file replacement cannot join two artifacts.
    with Path(path).open('rb') as stream: raw = stream.read(64 * 1024 * 1024 + 1)
    require(hashlib.sha256(raw).hexdigest() == json.loads(cell24['historical_provenance'])['input_sha256'], 'Input changed between reads')
    def invalid(_): raise consumer.InputContractError('Nonfinite JSON number')
    book = json.loads(raw.decode('utf-8-sig'), object_pairs_hook=retained._ObjectPairs,
                      parse_float=Decimal, parse_int=lambda s: Decimal(s) if s == '-0' else int(s), parse_constant=invalid)
    book = retained._resolve_notebook_json(book, (), [])
    products = {}; certs = {}; provenance = {}
    for number in PRODUCTS:
        found, certs[number], provenance[number] = _cell_products(book, number, expected_post, expected_run)
        products.update(found)
    rows = lambda n, i=0: products[PRODUCTS[n][i]]
    start = _index(rows(16), 'asset_key', 'Starting records')
    uuid = _index(rows(17), 'asset_key', 'UUID records')
    require(start.keys() == uuid.keys(), 'Starting/UUID membership mismatch')
    groups = defaultdict(set)
    for key in start:
        _same(start[key], uuid[key], 'Cell 16/17')
        groups[_group(uuid[key])].add(key)
    for key, row in uuid.items():
        require(row['candidate_asset_key_count'] == len(groups[_group(row)]), 'UUID group record count mismatch')
    links = _index(rows(20), 'analysis_subject_key', 'Linked groups')
    no_link = _index(rows(20, 1), 'asset_key', 'No-link members')
    qualified = _index(rows(21), 'asset_key', 'Qualified members')
    temporal = _index(rows(22), 'asset_key', 'Temporal members')
    currentness = _index(rows(22, 1), 'asset_key', 'Currentness members')
    residual = _index(rows(22, 2), 'asset_key', 'Residual members')
    customer = _index(rows(23), 'asset_key', 'Customer members')
    members = _index(cell24['members']['rows'], 'asset_id', 'Cell 24 members')
    require(no_link.keys() == qualified.keys(), 'No-link/qualification membership mismatch')
    linked_assets = set()
    group_family = {}
    for group, row in links.items():
        _decision(row, 20, 'explanation_class', ('explanation_family', 'provisional_treatment_lane', 'cohort_status'),
                  LINK_DECISIONS, ('recommended_treatment',))
        keys = row['candidate_asset_keys']
        require(group in groups and isinstance(keys, list) and len(keys) == len(set(keys))
                and set(keys) == groups[group] and row['candidate_asset_count'] == len(keys), 'Linked group membership mismatch')
        require(isinstance(row['candidate_ips'], list) and set(row['candidate_ips']) ==
                {_ip(start[k]) for k in keys}, 'Linked IP membership mismatch')
        require(not linked_assets.intersection(keys), 'Overlapping linked groups')
        linked_assets.update(keys)
        family = row['explanation_family']
        require(family in ('SOURCE_LOCAL_IDENTITY_MODEL_RECONCILIATION', 'ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION'), 'Unsupported linked explanation family')
        group_family[group] = FAMILIES[0 if family == 'SOURCE_LOCAL_IDENTITY_MODEL_RECONCILIATION' else 1]
        for key in keys:
            require(_group({'asset_key': key, 'tenable_host_uuid': row['candidate_tenable_uuid']}) == group,
                    'Linked UUID reference mismatch')
            is_uuid_link = uuid[key]['tenable_uuid_reconciliation_class'] == 'UUID_LINKS_EXISTING_DURABLE_ASSET'
            require(is_uuid_link == (group_family[group] == FAMILIES[0]), 'UUID/linked outcome conflict')
            if is_uuid_link:
                require(set(row['target_asset_keys']) == set(uuid[key]['stronger_non_ip_asset_keys']), 'Stronger UUID target conflict')
        require(isinstance(row['target_asset_keys'], list) and len(set(row['target_asset_keys'])) ==
                len(row['target_asset_keys']) == row['target_asset_count'], 'Target count/membership mismatch')
    require(not linked_assets.intersection(no_link) and linked_assets | no_link.keys() == start.keys(), 'Incomplete/overlapping starting partition')
    for key, row in no_link.items():
        require(row['tenable_identity_group_key'] == _group(uuid[key]), 'No-link group association mismatch')
        require(row['tenable_identity_group_key'] not in links, 'Group split between linked/no-link outcomes')
        _same(start[key], row, 'Cell 16/20 inherited evidence', INHERITED_ALIASES)
        require(row['candidate_ip'] == _ip(start[key]), 'Candidate IP conflict')
        _same(row, qualified[key], 'Cell 20/21')
    for row in qualified.values():
        _decision(row, 21, 'member_qualification_class', ('member_next_gate', 'workload_promotion_state'), QUALIFICATION_DECISIONS)
    for group in {r['tenable_identity_group_key'] for r in qualified.values()}:
        signatures = {(r['member_qualification_class'], r['member_next_gate'], r['workload_promotion_state'])
                      for r in qualified.values() if r['tenable_identity_group_key'] == group}
        require(len(signatures) == 1, 'Mixed qualification group requires its saved group-level product; no inferred gate')
        require(all(consumer.text(v) for v in next(iter(signatures))), 'Missing qualification decision')
    targeted = {k for k, r in qualified.items() if r['member_next_gate'] == 'TARGETED_IDENTITY_ENUMERATION'
                and r['workload_promotion_state'] == 'RESIDUAL_IDENTITY_REVIEW_CANDIDATE'}
    require(temporal.keys() == targeted, 'Temporal input membership mismatch')
    for key, row in temporal.items():
        _decision(row, 22, 'temporal_collection_class', ('temporal_next_gate', 'workload_promotion_state'),
                  TEMPORAL_DECISIONS, ('temporal_interpretation',))
        # Promotion state deliberately changes at this producer boundary.
        _same({k: v for k, v in qualified[key].items() if k != 'workload_promotion_state'}, row, 'Cell 21/22')
        _same(no_link[key], row, 'Cell 20/22')
        require(row['temporal_next_gate'] in ('CURRENTNESS_VALIDATION_BEFORE_IDENTITY_ENUMERATION',
                'FOCUSED_IDENTITY_DEVICE_ENUMERATION', 'COLLECTION_TIMING_REVIEW', 'TEMPORAL_EVIDENCE_LIMITATION_REVIEW'),
                'Unsupported temporal next gate')
        require((row['temporal_next_gate'] == 'FOCUSED_IDENTITY_DEVICE_ENUMERATION') ==
                (row['workload_promotion_state'] == 'RESIDUAL_CURRENT_IDENTITY_REVIEW_CANDIDATE'), 'Temporal promotion conflict')
    expected_current = {k for k, r in temporal.items() if r['temporal_next_gate'] == 'CURRENTNESS_VALIDATION_BEFORE_IDENTITY_ENUMERATION'}
    expected_residual = {k for k, r in temporal.items() if r['workload_promotion_state'] == 'RESIDUAL_CURRENT_IDENTITY_REVIEW_CANDIDATE'}
    require(currentness.keys() == expected_current and residual.keys() == expected_residual, 'Temporal subproduct membership mismatch')
    require(not expected_current.intersection(expected_residual), 'Overlapping temporal outcomes')
    for subset in (currentness, residual):
        for key, row in subset.items(): _same(temporal[key], row, 'Cell 22 subproduct')
    require(customer.keys() == residual.keys() == members.keys(), 'Residual/customer/Cell 24 membership mismatch')
    for key, row in customer.items():
        _same(residual[key], row, 'Cell 22/23')
        _same(qualified[key], row, 'Cell 21/23')
        _same(row, members[key], 'Cell 23/24')
        require(members[key]['member_id'] == row['tenable_identity_group_key'], 'Cell 24 member/group identity mismatch')
    selected = list(proposed_asset_keys)
    require(all(consumer.text(k) for k in selected) and len(selected) == len(set(selected))
            and set(selected) <= members.keys(), 'Proposed response selection must be distinct existing residual members')
    population = []
    for key in sorted(start):
        group = _group(uuid[key]); evidence = {PRODUCTS[16][0]: start[key], PRODUCTS[17][0]: uuid[key]}
        if group in links:
            row = links[group]; family = group_family[group]
            outcome = row['explanation_class']; reason = display_code(20, 'explanation_class', outcome)
            question = row['recommended_treatment']; next_code = None
            owner = row['provisional_treatment_lane'] + ' (provisional; not a case customer)'
            targets = row['target_asset_keys']; evidence[PRODUCTS[20][0]] = row
        else:
            row = qualified[key]; targets = []
            evidence.update({PRODUCTS[20][1]: no_link[key], PRODUCTS[21][0]: row})
            family = FAMILIES[2]; outcome = row['member_qualification_class']
            reason = display_code(21, 'member_qualification_class', outcome)
            next_code = row['member_next_gate']; question = display_code(21, 'member_next_gate', next_code)
            owner = 'UNESTABLISHED: no case routing at this gate'
            if key in temporal:
                row = temporal[key]; evidence[PRODUCTS[22][0]] = row
                family = FAMILIES[3] if key in currentness else FAMILIES[4] if key in residual else FAMILIES[5]
                outcome, reason = row['temporal_collection_class'], row['temporal_interpretation']
                next_code = row['temporal_next_gate']; question = display_code(22, 'temporal_next_gate', next_code)
                if key in currentness: evidence[PRODUCTS[22][1]] = currentness[key]
                if key in residual:
                    evidence[PRODUCTS[22][2]] = residual[key]; evidence[PRODUCTS[23][0]] = customer[key]
                    question = customer[key]['recommended_next_action']; owner = customer[key]['primary_customer']
            group_family.setdefault(group, family)
            require(group_family[group] == family, 'Mixed terminal outcomes within identity group; native group contract required')
        population.append({'asset_key': key, 'identity_group': group, 'observed_ip': _ip(start[key]),
            'outcome_family': family, 'original_outcome': outcome, 'reason': reason, 'remaining_question': question,
            'original_next_step': next_code,
            'customer_or_lane': owner, 'proposed_response': key in selected, 'stronger_asset_keys': targets,
            'observation_time': start[key]['tenable_last_observed_timestamp'], 'case_id': members[key]['case_instance_id'] if key in members else None,
            'evidence': evidence})
    counts = {f: (sum(r['outcome_family'] == f for r in population), sum(v == f for v in group_family.values())) for f in FAMILIES}
    closure = {r['stage_order']: r['count'] for r in cell24['closure']['rows']}
    expected = {1: len(start), 2: len(groups), 3: counts[FAMILIES[0]][1], 4: counts[FAMILIES[1]][1],
                5: len(groups) - len(links), 6: counts[FAMILIES[2]][1], 7: len({r['tenable_identity_group_key'] for r in temporal.values()}),
                8: counts[FAMILIES[3]][1], 9: counts[FAMILIES[4]][1]}
    require(all(closure[k] == v for k, v in expected.items()), 'Population/Cell 24 closure mismatch')
    checks = {16: {'Active unresolved-IP candidates': len(start), 'Distinct candidate asset keys': len(start)},
              17: {'Active unresolved-IP candidate rows': len(start), 'Tenable identity groups after UUID grouping': len(groups)},
              20: {'Starting unresolved-IP asset rows': len(start), 'Tenable source-local identity groups': len(groups),
                   'No-link candidate asset rows advancing': len(no_link)},
              21: {'No-link asset rows represented': len(no_link), 'True no-link Tenable identity groups entering Cell 21': expected[5]},
              22: {'Targeted candidate asset rows': len(temporal)},
              23: {'Cell 22 residual current identity rows': len(residual), 'Cell 22 residual distinct identity groups': expected[9]}}
    checks[17].update({'UUID groups linking an existing stronger Asset Plane identity': expected[3],
                      'Candidate rows in existing-asset UUID link groups': counts[FAMILIES[0]][0]})
    checks[20].update({'Source-local UUID groups linked to stronger Asset Plane identities': expected[3],
                      'Remaining groups with stronger exact-IP relationships': expected[4],
                      'Remaining groups with no stronger exact-IP relationship': expected[5],
                      'Linked identity groups assigned to explanatory cohorts': len(links),
                      'Linked/no-link overlap': 0, 'Duplicate explanatory cohort assignments': 0})
    checks[21].update({'Groups not ready for workload promotion': expected[6],
                      'Residual targeted-enumeration candidates': expected[7]})
    checks[22].update({'Currentness-validation candidates before identity enumeration': len(currentness),
                      'Residual current identity-enumeration candidates': len(residual),
                      'Temporal evidence-limitation / timing-review candidates': len(temporal) - len(currentness) - len(residual)})
    checks[23].update({'Residual identities missing Cell 21 context': 0, 'Invalid Cell 22 temporal-gate rows': 0,
                      'Identity groups assigned to multiple customers': 0})
    for number, fields in checks.items():
        for label, value in fields.items(): require(_field(certs[number], label) == str(value), f'Cell {number}: saved count mismatch: {label}')
    return {'cell24': cell24, 'population': population, 'counts': counts, 'provenance': provenance,
            'input_sha256': hashlib.sha256(raw).hexdigest()}


def build_population_workbook(path, *, expected_post, expected_run, proposed_asset_keys=(), synthetic=False, generated_at=None):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    data = load_population(path, expected_post=expected_post, expected_run=expected_run, proposed_asset_keys=proposed_asset_keys)
    generated_at = generated_at or datetime.now(timezone.utc).isoformat()
    wb = consumer.build_cell24_workbook(data['cell24'], synthetic=synthetic, generated_at=generated_at)
    summary = wb.create_sheet('RCSO Summary', 0); detail = wb.create_sheet('IP Population Review', 1)
    navy = '16324F'; teal = '087F8C'
    def put(ws, row, col, value):
        c = ws.cell(row, col, consumer.evidence_value(value))
        if isinstance(c.value, str): c.data_type = 's'
        c.alignment = Alignment(vertical='top', wrap_text=True)
        return c
    def band(ws, row, title, end):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end)
        c = put(ws, row, 1, title); c.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=navy); ws.row_dimensions[row].height = 30
    def link(cell, sheet, row):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A{row}", display=str(cell.value))
        cell.font = Font(name='Calibri', color='0563C1', underline='single')
    label = consumer.SYNTHETIC_LABEL if synthetic else 'REPRODUCTION OF RETAINED NOTEBOOK RESULTS / NOT A LIVE ASSESSMENT'
    band(summary, 1, label, 6); band(detail, 1, label, 12)
    band(summary, 3, 'IP identity review: complete population and remaining decisions', 6)
    for row, text in enumerate([
        f'Scope: {expected_post} / {expected_run}',
        f'Report generated: {generated_at}. Observation times remain source text in the population view.',
        'Current means current in the captured assessment, not at report generation.',
        'Internal review pending. Proposed follow-up has not been approved or sent. Original producer routing is unchanged.',
        'Starting analytical records and Tenable identity groups are different units; neither asserts a physical-device or distinct-IP total.',
        'Unrecorded facts remain unestablished. Source absence is not proof of a missing control or an unsafe device.',
    ], 5):
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        put(summary, row, 1, text); summary.row_dimensions[row].height = 34
    for col, title in enumerate(('Retained outcome', 'Starting records', 'Identity groups'), 1): put(summary, 12, col, title)
    for row, family in enumerate(FAMILIES, 13):
        for col, value in enumerate((family, *data['counts'][family]), 1): put(summary, row, col, value)
    for col, value in enumerate(('Complete population', len(data['population']), sum(v[1] for v in data['counts'].values())), 1): put(summary, 19, col, value)
    band(summary, 21, 'Residual questions and deferred work', 6)
    notes = [f'Currentness checks outstanding: {data["counts"][FAMILIES[3]][0]} starting records. See the retained temporal interpretation and next gate.',
             f'Proposed response selection: {sum(r["proposed_response"] for r in data["population"])} residual records. This selection does not create or reroute cases.',
             ('Remaining question: which identity and managing team are supported by the captured evidence? Original owner-specific actions are in Case review.'
              if data['cell24']['cases']['rows'] else 'No residual cases in this retained result. Currentness and other deferred treatments remain visible.'),
             'Open complete population', 'Open original case review',
             'Cross-cell checks validate saved source, scope, membership and shared evidence. An artifact hash is not a signed certificate or current live authority.',
             'Input SHA-256: ' + data['input_sha256']]
    for row, text in enumerate(notes, 23):
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        put(summary, row, 1, text); summary.row_dimensions[row].height = 36
    summary['A26'].hyperlink = "#'IP Population Review'!A7"; summary['A27'].hyperlink = "#'Case review'!A8"
    for i, width in enumerate((47, 21, 21, 15, 15, 15), 1): summary.column_dimensions[get_column_letter(i)].width = width
    summary.freeze_panes = 'A12'; summary.print_options.horizontalCentered = True
    band(detail, 3, 'Every starting analytical record; filters do not alter producer decisions', 12)
    detail.merge_cells('A4:L4'); put(detail, 4, 1, 'Arrays retain native relationships: no element-by-element IP/key pairing. Evidence columns retain full saved records and source schema notes.')
    detail.row_dimensions[4].height = 32
    detail.merge_cells('A5:L5'); put(detail, 5, 1, 'Unrecorded facts are unestablished. Times and exact numbers are lossless text where needed. Proposed response = internal-review selection only.')
    detail.row_dimensions[5].height = 32
    headers = ['Observed IP', 'Starting asset key', 'Tenable identity group', 'Outcome family', 'Original outcome',
               'Explanation', 'Remaining / deferred question', 'Customer / provisional lane', 'Communication review status',
               'Recorded observation time', 'Stronger asset keys', 'Residual case', 'Device family', 'Network scope',
               'Management model', 'Management owner', 'Source pattern', 'Tenable UUID', 'Time basis',
               'Original next-step code', 'Proposed response']
    product_names = [p for ps in PRODUCTS.values() for p in ps]
    evidence_start = len(headers) + 1
    headers += [p + ' (native evidence)' for p in product_names]
    for col, title in enumerate(headers, 1):
        c = put(detail, 7, col, title); c.fill = PatternFill('solid', fgColor=teal); c.font = Font(bold=True, color='FFFFFF')
        detail.column_dimensions[get_column_letter(col)].width = 32 if col != 6 and col != 7 else 65
    detail.row_dimensions[7].height = 72
    # Locate actual accepted evidence rows rather than assuming the renderer sort.
    case_rows = {wb['Case evidence'].cell(i, 1).value: i for i in range(8, wb['Case evidence'].max_row + 1)}
    for line, row in enumerate(data['population'], 8):
        original = row['evidence'][PRODUCTS[16][0]]
        values = [row[k] for k in ('observed_ip', 'asset_key', 'identity_group', 'outcome_family', 'original_outcome',
                  'reason', 'remaining_question', 'customer_or_lane', 'proposed_response', 'observation_time', 'stronger_asset_keys', 'case_id')]
        values[8] = PROPOSED_LABEL if row['proposed_response'] else 'NOT SELECTED FOR THIS PROPOSAL'
        values += [original[k] for k in ('device_family', 'network_scope', 'management_model', 'management_owner', 'source_mask_active', 'tenable_host_uuid')]
        values += [TIME_BASIS, row['original_next_step'], row['proposed_response']]
        values += [retained._json(row['evidence'][p]) if p in row['evidence'] else 'NOT IN THIS BRANCH' for p in product_names]
        for col, value in enumerate(values, 1): put(detail, line, col, value)
        if row['case_id'] is not None:
            detail.cell(line, 12).hyperlink = "#'Case evidence'!A" + str(case_rows[consumer.evidence_value(row['case_id'])])
        if row['proposed_response']:
            for cell in detail[line]: cell.fill = PatternFill('solid', fgColor='FFF2CC')
            detail.cell(line, 9).font = Font(bold=True, color='7F6000')
        detail.row_dimensions[line].height = 105
    for offset, product in enumerate(product_names, evidence_start):
        number = next(n for n, ps in PRODUCTS.items() if product in ps)
        detail.cell(7, offset).comment = Comment(consumer.evidence_value(retained._json(data['provenance'][number])), 'RCAIP source reference')
    if data['population']:
        table = Table(displayName='IPPopulationReview', ref=f'A7:{get_column_letter(len(headers))}{7 + len(data["population"])}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        detail.add_table(table)
    # Communication selection is distinct from every original case decision.
    # Existing question/customer are copied from validated Cell 23/24 evidence.
    band(summary, 31, 'Proposed response items for internal RCSO review', 6)
    selected_rows = [(i, row) for i, row in enumerate(data['population'], 8) if row['proposed_response']]
    if not selected_rows:
        summary.merge_cells('A33:F33')
        put(summary, 33, 1, 'No items selected for this proposal. Original cases and deferred questions remain unchanged.')
        summary.row_dimensions[33].height = 36
    case_cards = {ws_row[1].value: ws_row[0].row for ws_row in wb['Case review'].iter_rows(max_col=2)
                  if ws_row[0].value == 'Case ID'}
    for position, (population_line, item) in enumerate(selected_rows):
        first = 33 + position * 10
        band(summary, first, PROPOSED_LABEL, 6)
        summary.cell(first, 1).fill = PatternFill('solid', fgColor='FFF2CC')
        summary.cell(first, 1).font = Font(name='Calibri', size=12, bold=True, color='7F6000')
        for offset, (title, value) in enumerate([
            ('Observed IP', item['observed_ip']), ('Starting analytical record', item['asset_key']),
            ('Retained observation time', item['observation_time']), ('Observation time basis', TIME_BASIS),
            ('Existing question / action', item['remaining_question']), ('Existing customer', item['customer_or_lane']),
            ('Original case', item['case_id']),
        ], 1):
            line = first + offset
            put(summary, line, 1, title)
            summary.merge_cells(start_row=line, start_column=2, end_row=line, end_column=6)
            value_cell = put(summary, line, 2, value)
            height = max(30, ((len(str(value_cell.value)) + 74) // 75) * 16 + 12)
            require(height <= 409, 'Selected question exceeds readable summary row; no truncation')
            summary.row_dimensions[line].height = height
        summary.merge_cells(start_row=first + 8, start_column=1, end_row=first + 8, end_column=3)
        summary.merge_cells(start_row=first + 8, start_column=4, end_row=first + 8, end_column=6)
        link(put(summary, first + 8, 1, 'Open this population record'), 'IP Population Review', population_line)
        link(put(summary, first + 8, 4, 'Open original case review'), 'Case review', case_cards[consumer.evidence_value(item['case_id'])])
        summary.row_dimensions[first + 8].height = 30
    detail.freeze_panes = 'D8'; detail.sheet_view.showGridLines = False; summary.sheet_view.showGridLines = False
    for ws in (summary, detail):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    require(wb.sheetnames == SHEETS, 'Six-sheet contract mismatch')
    return wb
