# Databricks notebook source
"""IP Population Review entry point. Defining this function performs no I/O.

After separate approval/manual deployment, invoke publish_ip_population with
explicit module_folder, source_uri, expected_input_sha256, expected_post,
expected_run, report_dir and proposed_asset_keys. No DIAG live globals are used.
Requires the already-approved enterprise Python environment, openpyxl 3.1.5 and
databricks-sdk. No installation, Config execution or source acquisition occurs.
"""
import hashlib
import importlib.util
import io
from pathlib import Path
import re
import sys
import tempfile
from datetime import datetime, timezone

EXPECTED_MODULES = {
    'rcaip_cell24_presentation': 'c2d54f4b019a1630965a6c0b947d7fab54bf86cf76f236b4e352ddc59f8736d6',
    'rcaip_cell24_retained': '326fb26c2c8abbbf52af76fe104e0fec119ac083efa3abb954c8f2ca1276b55e',
    'rcaip_ip_population': '1470f6dda996f9e2e9687d3820ace86d2ac808f46c9b834acceb4fe1c250801f',
}


def publish_ip_population(*, module_folder, source_uri, expected_input_sha256,
                          expected_post, expected_run, report_dir, proposed_asset_keys=()):
    """Explicit enterprise-only invocation; retained reproduction, not assessment.

    Example configuration is supplied separately in the review return. Never run
    this complete caller in local tests; tests extract the pagination statements.
    """
    from openpyxl import load_workbook
    from databricks.sdk import WorkspaceClient
    from databricks.sdk.service.workspace import ExportFormat, ImportFormat
    if not re.fullmatch(r'[a-z0-9][a-z0-9_-]*', expected_post) or not expected_run:
        raise ValueError('STOP: explicit canonical Post slug and run required')
    if not source_uri.startswith('dbfs:/') or not report_dir.startswith('/Users/'):
        raise ValueError('STOP: explicit retained DBFS input and Workspace user report directory required')
    if not re.fullmatch(r'[a-f0-9]{64}', expected_input_sha256):
        raise ValueError('STOP: reviewed input artifact SHA-256 required')
    # Same reviewed source-verification/import pattern as the Cell 24 caller.
    folder = Path(module_folder)
    for name, expected in EXPECTED_MODULES.items():
        path = folder / (name + '.py')
        text = path.read_bytes().decode('utf-8-sig').replace('\r\n', '\n').replace('\r', '\n')
        actual = hashlib.sha256((text.rstrip('\n') + '\n').encode('utf-8')).hexdigest()
        if actual != expected: raise RuntimeError('STOP: deployed source hash mismatch: ' + name)
    modules = {}
    for name in EXPECTED_MODULES:
        spec = importlib.util.spec_from_file_location(name, folder / (name + '.py'))
        if spec is None or spec.loader is None: raise RuntimeError('STOP: module loader unavailable: ' + name)
        module = importlib.util.module_from_spec(spec); sys.modules[name] = module
        spec.loader.exec_module(module); modules[name] = module
    client = WorkspaceClient(); source_path = source_uri[len('dbfs:'):]
    max_bytes = 64 * 1024 * 1024
    info = client.dbfs.get_status(source_path)
    if info.is_dir or type(info.file_size) is not int or not 0 < info.file_size <= max_bytes:
        raise RuntimeError('STOP: retained input is not a nonempty file within the 64 MiB bound.')
    # PAGINATED_DBFS_READ_START
    # Identical read statements to the accepted Cell 24 caller; no new transport.
    notebook_buffer = bytearray()
    with client.dbfs.download(source_path) as stream:
        while True:
            request_bytes = min(1024 * 1024, max_bytes + 1 - len(notebook_buffer))
            chunk = stream.read(request_bytes)
            if not chunk:
                break
            notebook_buffer.extend(chunk)
            if len(notebook_buffer) > max_bytes:
                raise RuntimeError("STOP: retained notebook exceeds the byte bound.")
    notebook_bytes = bytes(notebook_buffer)
    del notebook_buffer
    # PAGINATED_DBFS_READ_END
    if len(notebook_bytes) != info.file_size:
        raise RuntimeError('STOP: downloaded notebook size does not match DBFS metadata.')
    input_sha = hashlib.sha256(notebook_bytes).hexdigest()
    if input_sha != expected_input_sha256: raise RuntimeError('STOP: retained input artifact hash mismatch')
    builder = modules['rcaip_ip_population']
    with tempfile.TemporaryDirectory(prefix='rcaip_ip_population_') as folder:
        local_input = Path(folder) / 'retained.ipynb'; local_input.write_bytes(notebook_bytes)
        workbook = builder.build_population_workbook(local_input, expected_post=expected_post,
            expected_run=expected_run, proposed_asset_keys=proposed_asset_keys, synthetic=False)
        with io.BytesIO() as buffer:
            workbook.save(buffer); workbook_bytes = buffer.getvalue()
        workbook.close()
    check = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
    try:
        if check.sheetnames != builder.SHEETS: raise RuntimeError('STOP: six-sheet contract mismatch')
        for ws in check.worksheets[2:]:
            if 'NOT LIVE AUTHORITY' not in str(ws['A1'].value): raise RuntimeError('STOP: missing historical authority label')
        for ws in check.worksheets[:2]:
            if 'REPRODUCTION OF RETAINED NOTEBOOK RESULTS' not in str(ws['A1'].value):
                raise RuntimeError('STOP: missing population historical label')
        for ws in check:
            for table in ws.tables.values():
                if table.autoFilter is None or table.ref != table.autoFilter.ref:
                    raise RuntimeError('STOP: table/filter metadata mismatch')
    finally:
        check.close()
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    report_path = f'{report_dir.rstrip("/")}/RCAIP_{expected_post}_IP_Population_RETAINED_{stamp}.xlsx'
    client.workspace.mkdirs(report_dir)
    client.workspace.upload(path=report_path, content=io.BytesIO(workbook_bytes), format=ImportFormat.AUTO, overwrite=False)
    with client.workspace.download(path=report_path, format=ExportFormat.AUTO) as stream:
        published = stream.read(len(workbook_bytes) + 1)
    workbook_sha = hashlib.sha256(workbook_bytes).hexdigest()
    if hashlib.sha256(published).hexdigest() != workbook_sha:
        raise RuntimeError('STOP: published bytes differ from generated workbook')
    return {'workspace_path': report_path, 'input_sha256': input_sha, 'workbook_sha256': workbook_sha,
            'published_bytes_verified': True, 'status': 'Retained reproduction; operator/desktop Excel review remains'}
