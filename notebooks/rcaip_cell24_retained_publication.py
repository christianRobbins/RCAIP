import hashlib
import importlib.util
import io
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

# =============================================================================
# RCAIP CELL 24 RETAINED RESULTS REPRODUCTION
# Invokes the reviewed modules from commit 7bee1eb.
# Publication caller: proposed repository integration; not an analytical module.
# No DIAG execution / no Spark analytical replay.
# =============================================================================

# ---------------------------------------------------------------------------
# 1. Dependency
# ---------------------------------------------------------------------------

try:
    import openpyxl
except ImportError:
    subprocess.check_call([
        sys.executable,
        "-m",
        "pip",
        "install",
        "--quiet",
        "openpyxl==3.1.5",
    ])
    import openpyxl

from openpyxl import load_workbook

print("openpyxl:", openpyxl.__version__)

# ---------------------------------------------------------------------------
# 2. Locate and verify deployed RCAIP modules
# ---------------------------------------------------------------------------

workspace_folder = Path(
    "/Workspace/Users/operator@example.invalid/"
    "RCAIP_V3_RELATIONSHIP_ASSURANCE"
)

presentation_path = workspace_folder / "rcaip_cell24_presentation.py"
retained_path = workspace_folder / "rcaip_cell24_retained.py"

expected_hashes = {
    presentation_path:
        "c2d54f4b019a1630965a6c0b947d7fab54bf86cf76f236b4e352ddc59f8736d6",
    retained_path:
        "326fb26c2c8abbbf52af76fe104e0fec119ac083efa3abb954c8f2ca1276b55e",
}


def lf_normalized_sha256(path):
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = text.rstrip("\n") + "\n"
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


for path, expected in expected_hashes.items():
    if not path.is_file():
        raise RuntimeError(f"STOP: deployed module not found: {path}")
    actual = lf_normalized_sha256(path)
    print(path.name, "LF-normalized SHA-256:", actual)
    if actual != expected:
        raise RuntimeError(
            f"STOP: {path.name} does not match reviewed commit 7bee1eb."
        )

print("PASS: deployed module hashes match reviewed source.")

# ---------------------------------------------------------------------------
# 3. Load exact Workspace modules directly
# ---------------------------------------------------------------------------


def load_workspace_module(module_name, path):
    spec = importlib.util.spec_from_file_location(module_name, str(path))
    if spec is None or spec.loader is None:
        raise RuntimeError(f"STOP: could not create Python loader for {path}")
    module = importlib.util.module_from_spec(spec)
    # Makes presentation available when retained.py imports it.
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


presentation = load_workspace_module("rcaip_cell24_presentation", presentation_path)
retained = load_workspace_module("rcaip_cell24_retained", retained_path)

print("PASS: reviewed Cell 24 modules loaded.")

# ---------------------------------------------------------------------------
# 4. Read retained example_post_a notebook artifact from DBFS
# ---------------------------------------------------------------------------

from databricks.sdk import WorkspaceClient
from databricks.sdk.service.workspace import ExportFormat, ImportFormat

source_uri = (
    "dbfs:/FileStore/tables/example_rcaip_inputs/"
    "example_post_a_example_run/"
    "RCAIP_DIAG_001_Asset_Visibility_Funnel__1_.ipynb"
)
source_path = source_uri[len("dbfs:"):]
client = WorkspaceClient()
max_bytes = 64 * 1024 * 1024
info = client.dbfs.get_status(source_path)

if (
    info.is_dir
    or type(info.file_size) is not int
    or not 0 < info.file_size <= max_bytes
):
    raise RuntimeError(
        "STOP: retained input is not a nonempty file within the 64 MiB bound."
    )

print("Reading retained notebook artifact.")
print("DIAG source will NOT be executed.")

# PAGINATED_DBFS_READ_START
# DBFS allows at most 1 MiB per read. The 64 MiB bound is for the WHOLE
# artifact, not one request. Read until EOF, allowing one overflow byte.
notebook_buffer = bytearray()
with client.dbfs.download(source_path) as stream:
    while True:
        request_bytes = min(1024 * 1024, max_bytes + 1 - len(notebook_buffer))
        chunk = stream.read(request_bytes)
        if not chunk:
            break
        notebook_buffer.extend(chunk)
        if len(notebook_buffer) > max_bytes:
            raise RuntimeError("STOP: retained notebook exceeds the byte bound.")
notebook_bytes = bytes(notebook_buffer)
del notebook_buffer
# PAGINATED_DBFS_READ_END

if len(notebook_bytes) != info.file_size:
    raise RuntimeError(
        "STOP: downloaded notebook size does not match DBFS metadata."
    )
if len(notebook_bytes) > max_bytes:
    raise RuntimeError("STOP: retained notebook exceeds the byte bound.")

input_sha = hashlib.sha256(notebook_bytes).hexdigest()
print("Retained notebook bytes:", len(notebook_bytes))
print("Input SHA-256:", input_sha)

# ---------------------------------------------------------------------------
# 5. Validate retained Cell 24 evidence and build workbook
# ---------------------------------------------------------------------------

with tempfile.TemporaryDirectory(prefix="rcaip_cell24_retained_") as folder:
    local_input = Path(folder) / "RCAIP_DIAG_001_Asset_Visibility_Funnel__1_.ipynb"
    local_input.write_bytes(notebook_bytes)
    print("Validating saved Cell 24 certification and products.")
    workbook = retained.workbook_from_retained_cell24(
        local_input,
        expected_post="example_post_a",
        expected_run="example_post_a_example_run",
        synthetic=False,
    )
    with io.BytesIO() as buffer:
        workbook.save(buffer)
        workbook_bytes = buffer.getvalue()
    workbook.close()

print("PASS: retained Cell 24 results validated and workbook generated.")

# ---------------------------------------------------------------------------
# 6. Reopen generated workbook before publication
# ---------------------------------------------------------------------------

check = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
try:
    expected_sheets = ["Case review", "Case evidence", "Members", "Branch closure"]
    if check.sheetnames != expected_sheets:
        raise RuntimeError(
            "STOP: workbook sheet structure mismatch. "
            f"Found={check.sheetnames}"
        )
    for sheet in check.worksheets:
        if "RETAINED NOTEBOOK RESULTS REPRODUCTION" not in str(sheet["A2"].value):
            raise RuntimeError(
                "STOP: retained-results label missing on " + sheet.title
            )
        if "NOT LIVE AUTHORITY" not in str(sheet["A1"].value):
            raise RuntimeError(
                "STOP: historical-authority warning missing on " + sheet.title
            )
finally:
    check.close()

print("PASS: generated workbook reopened and validated.")

# ---------------------------------------------------------------------------
# 7. Publish exact bytes to established RCAIP Workspace report location
# ---------------------------------------------------------------------------

principal = client.current_user.me().user_name
if not principal:
    raise RuntimeError("STOP: authenticated Workspace user could not be resolved.")
report_dir = f"/Users/{principal}/RCAIP_Reports/example_post_a"
stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
report_path = f"{report_dir}/RCAIP_London_Cell24_RETAINED_{stamp}.xlsx"
client.workspace.mkdirs(report_dir)
client.workspace.upload(
    path=report_path,
    content=io.BytesIO(workbook_bytes),
    format=ImportFormat.AUTO,
    overwrite=False,
)
print("Workbook uploaded. Verifying published bytes.")

# ---------------------------------------------------------------------------
# 8. Verify publication bytes
# ---------------------------------------------------------------------------

with client.workspace.download(path=report_path, format=ExportFormat.AUTO) as stream:
    published_bytes = stream.read(len(workbook_bytes) + 1)
workbook_sha = hashlib.sha256(workbook_bytes).hexdigest()
published_sha = hashlib.sha256(published_bytes).hexdigest()
if published_sha != workbook_sha:
    raise RuntimeError(
        "STOP: published workbook bytes do not match the validated generated workbook."
    )

# ---------------------------------------------------------------------------
# 9. Final result
# ---------------------------------------------------------------------------

print()
print("=" * 100)
print("example_post_a CELL 24 RETAINED WORKBOOK: CREATED")
print("=" * 100)
print("Workspace file:", report_path)
print("Input SHA-256:", input_sha)
print("Workbook SHA-256:", workbook_sha)
print("Published SHA-256:", published_sha)
print("Published bytes verified: YES")
print("Status: historical retained-results reproduction; desktop Excel acceptance remains.")
print("=" * 100)
