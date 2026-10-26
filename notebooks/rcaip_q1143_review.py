"""Bounded Q1143 presentation of two captured IP001 products. No action on import.

No classifier, ownership inference, XML parsing or source acquisition. The union
of scoped observation groups is a review denominator, NOT physical-device count.
"""
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
import math
from pathlib import Path
import re
from uuid import uuid4

PRODUCTS = {'ports':'IP001_Q1143_PRINTER_MFD_PORTS_WITH_PROVENANCE',
            'devices':'IP001_Q1143_PRINTER_MFD_DEVICE_REVIEW_WITH_PROVENANCE'}
HOST = ('post_slug','rcaip_run_id','scan_id','host_key','ip_address','asset_domain','scan_scope','scan_location','scan_source_file')
SHEETS = ['Printer Review','Assessor Checks','Port Evidence','Device Evidence','Read Me']
AUDIT_SHA = '4896cfd6c1933bfe3773bd0837b63b1cf4b2b055a5663f0b1a0580ce38df373f'
PRODUCER_SHA = '7fbd254de65d6bb050094da9f5f2aad0249584ca1e14cd71307f92196fe2746b'
LIMIT = 20000


class ContractError(ValueError):
    pass


def require(ok, code):
    if not ok:
        raise ContractError(code)


def strict_json(text):
    def pairs(values):
        result = {}
        for key,value in values:
            require(key not in result,'DUPLICATE_JSON_KEY')
            result[key] = value
        return result
    return json.loads(text, object_pairs_hook=pairs, parse_float=Decimal, parse_constant=lambda _: (_ for _ in ()).throw(ContractError('NONFINITE_JSON')))


def evidence(value):
    """Existing population-presentation encoding pattern; reversible text escapes.
    Null is a blank; empty and leading-quote text is JSON-quoted. Integers stay numeric when Excel-exact.
    No formula interpretation, silent truncation, rounding or temporal inference.
    """
    if value is None or type(value) is bool:
        return value
    if type(value) is int:
        return value if abs(value)<10**15 else str(value)
    if type(value) is float:
        require(math.isfinite(value),'NONFINITE_EVIDENCE');return repr(value)
    if isinstance(value,Decimal):
        require(value.is_finite(),'NONFINITE_EVIDENCE');return str(value)
    if isinstance(value,(date,datetime)):
        value=value.isoformat()
    if isinstance(value,(list,tuple)):
        require(all(x is None or isinstance(x,str) for x in value),'UNSUPPORTED_ARRAY_EVIDENCE')
        value=json.dumps(value,ensure_ascii=False,separators=(',',':'))
    require(isinstance(value,str),'UNSUPPORTED_EVIDENCE_TYPE')
    if value=='' or value.startswith('"') or re.match(r'^\d{4,}-\d\d-\d\d(?:T| |$)',value):
        value=json.dumps(value,ensure_ascii=False)
    value=value.replace('\\','\\\\')
    value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]',lambda m:'\\u%04x'%ord(m[0]),value)
    require(not re.search(r'[\ud800-\udfff]',value),'UNPAIRED_SURROGATE')
    require(len(value.encode('utf-16-le'))//2<=32767,'EXCEL_CELL_LIMIT')
    return value


def contracts(execution, post, run):
    require(isinstance(execution,dict) and execution.get('state')=='CERTIFIED','CERTIFIED_EXECUTION_REQUIRED')
    require(execution.get('post_slug')==post and execution.get('rcaip_run_id')==run and bool(post) and bool(run),'EXECUTION_SCOPE_MISMATCH')
    require(execution.get('sources',{}).get('producer',{}).get('sha256')==PRODUCER_SHA,'PRODUCER_SOURCE_MISMATCH')
    records={}
    for kind,logical in PRODUCTS.items():
        record=deepcopy(execution.get('outputs',{}).get(logical))
        require(isinstance(record,dict),'MISSING_CAPTURED_'+kind.upper())
        required={'logical_name','table','version','table_id','schema','generation','execution_id','producer_sha256','write_token'}
        require(required<=record.keys(),'INCOMPLETE_CAPTURED_BINDING')
        require(record['logical_name']==logical and record['execution_id']==execution.get('execution_id') and bool(record['execution_id']),'CAPTURED_EXECUTION_MISMATCH')
        require(record['producer_sha256']==PRODUCER_SHA and record['generation']==execution.get('generation'),'CAPTURED_SOURCE_OR_GENERATION_MISMATCH')
        token=strict_json(record['write_token'])
        require(isinstance(token,dict) and all(token.get(k)==v for k,v in dict(logical_name=logical,execution_id=record['execution_id'],producer_sha256=PRODUCER_SHA,generation=record['generation'],post_slug=post,rcaip_run_id=run).items()) and bool(token.get('write_id')),'CAPTURED_TOKEN_MISMATCH')
        require(type(record['version']) is int and record['version']>=0 and bool(record['table_id']),'CAPTURED_IDENTITY_VERSION_REQUIRED')
        names=[x[0] for x in record['schema']]
        require(len(names)==len(set(names)) and set(HOST+('generated_at_utc',))<=set(names),'SCOPED_SCHEMA_REQUIRED')
        need=('protocol','port','port_state','port_observation') if kind=='ports' else ('port_members',)
        require(set(need)<=set(names),'PRODUCT_SCHEMA_MISSING')
        records[kind]=record
    require(records['ports']['table']!=records['devices']['table'] and records['ports']['table_id']!=records['devices']['table_id'],'DISTINCT_PRODUCTS_REQUIRED')
    return records


def read_inputs(spark, execution, binding_api, *, expected_post, expected_run, max_rows=LIMIT):
    """Explicit enterprise read: only captured versions of the two named products.
    No registry/latest lookup. Temporal transport and bounds use reviewed reader.
    """
    require(type(max_rows) is int and 0<max_rows<=LIMIT,'ROW_BOUND')
    require(hashlib.sha256(Path(binding_api.__file__).read_bytes().replace(b'\r\n',b'\n')).hexdigest()==AUDIT_SHA,'BINDING_API_SOURCE_MISMATCH')
    before=deepcopy(execution)
    records=contracts(before,expected_post,expected_run)
    data={}
    for kind,record in records.items():
        binding=binding_api.Binding(record['table'],record['version'],record['table_id'],tuple(map(tuple,record['schema'])),record['producer_sha256'],record['generation'])
        data[kind]=binding_api.read_delta_rows(spark,binding,expected_post,expected_run,max_rows=max_rows)
        require(binding_api.table_identity(spark,record['table'])==record['table_id'],'TABLE_IDENTITY_CHANGED_AFTER_READ')
    require(execution==before,'EXECUTION_CHANGED_DURING_READ')
    return validate(data,records,post=expected_post,run=expected_run,generation_key=binding_api.generation_key,max_rows=max_rows)


def validate(data, records, *, post, run, generation_key, max_rows=LIMIT):
    """Pure materialized-input path; caller retains full native evidence unchanged."""
    require(set(data)==set(PRODUCTS) and set(records)==set(PRODUCTS),'MISSING_INPUT')
    require(type(max_rows) is int and 0<max_rows<=LIMIT,'ROW_BOUND')
    groups=defaultdict(lambda: {'ports':[],'devices':[]})
    saved=deepcopy(data)
    for kind in PRODUCTS:
        require(isinstance(saved[kind],list) and len(saved[kind])<=max_rows,'ROW_BOUND')
        names=[x[0] for x in records[kind]['schema']]
        for row in saved[kind]:
            require(isinstance(row,dict) and set(row)==set(names),'ROW_SCHEMA_MISMATCH')
            require(all(isinstance(row.get(k),str) and row[k].strip() for k in HOST),'GROUP_KEY_MISSING')
            require(row['post_slug']==post and row['rcaip_run_id']==run,'ROW_SCOPE_MISMATCH')
            require('runtime_post_slug' not in row or row['runtime_post_slug']==post,'ROW_SCOPE_ALIAS_MISMATCH')
            require(generation_key(row['generated_at_utc'])==generation_key(records[kind]['generation']),'ROW_GENERATION_MISMATCH')
            for value in row.values(): evidence(value)  # reject unsupported/lossy transport before rendering
            if kind=='ports':
                require(type(row['port']) is int and 0<row['port']<=65535 and row['protocol'] in ('tcp','udp','sctp'),'ENDPOINT_INVALID')
                require(row['port_state'] in ('open','closed','filtered','unfiltered','open|filtered','closed|filtered'),'PORT_STATE_INVALID')
                native=strict_json(row['port_observation'])
                require(isinstance(native,dict) and all(native.get(k)==row[k] for k in HOST+('protocol','port','port_state')),'NATIVE_OBSERVATION_MISMATCH')
            else:
                require(isinstance(row['port_members'],list) and all(isinstance(v,str) for v in row['port_members']),'MEMBER_ARRAY_REQUIRED')
                for member in row['port_members']:
                    parsed=strict_json(member)
                    require(isinstance(parsed,dict) and all(parsed.get(k)==row[k] for k in HOST),'MEMBER_GROUP_MISMATCH')
            groups[tuple(row[k] for k in HOST)][kind].append(row)
    ordered=[]
    # Compare the producer's exact native member strings; do not reserialize or
    # round JSON numeric evidence merely to force agreement.
    for index,(key,rows) in enumerate(sorted(groups.items()),1):
        ports,devices=rows['ports'],rows['devices']
        if not ports:reconciliation='Device input only - port evidence unavailable; reconcile inputs'
        elif not devices:reconciliation='Port input only - device review unavailable; reconcile inputs'
        else:
            a=Counter(r['port_observation'] for r in ports)
            b=Counter(v for r in devices for v in r['port_members'])
            reconciliation=('Both inputs - native member occurrences agree' if a==b else 'Both inputs - member discrepancy; reconcile inputs')
        ordered.append(dict(reference=f'G{index:04d}',key=key,**rows,reconciliation=reconciliation))
    require(sum(len(g['ports']) for g in ordered)==len(saved['ports']) and sum(len(g['devices']) for g in ordered)==len(saved['devices']),'ACCOUNTING_FAILURE')
    return dict(post=post,run=run,records=deepcopy(records),data=saved,groups=ordered)


CHECKS = [
    ('Applicability / scope','Confirm this observed group is an applicable printer/MFD.','NOT_YET_CONFIRMED'),
    ('Assessor sample / rationale','Select sample only after applicability; record rationale. No mandatory size inferred.','NOT_YET_CONFIRMED'),
    ('Controls selected','Record the sample of controls actually chosen for this group.','NOT_YET_CONFIRMED'),
    ('Port closure: TCP 21 / 23 / 80','Test closure separately; absence of an open observation does not verify closure.','NOT_TESTED'),
    ('Web administration','Test actual accessibility and policy separately; redirect/authentication is not port closure.','NOT_TESTED'),
    ('Password / unique credentials','Record verification outcome/reference only. Never record passwords.','NOT_TESTED'),
    ('Guide application','Record guide application evidence. Supplied CSA alternative: confirmed guide not applied should be Fail.','NOT_TESTED'),
    ('Stakeholder response','Identify actual owner/contact and record response/remediation evidence.','NOT_YET_CONFIRMED'),
    ('Secondary scan','Record separately authorized Nmap reference and outcome; this workbook does not scan.','NOT_TESTED'),
    ('Final assessor judgement','Assessor determines CSA judgement from scope, sampled controls and evidence.','NOT_YET_CONFIRMED'),
]


def build_workbook(model, *, synthetic=False):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise ContractError('OPENPYXL_REQUIRED_NO_AUTOMATIC_INSTALL') from exc
    book=Workbook();book.remove(book.active)
    matrix={};links={}
    label='SYNTHETIC ONLY - ' if synthetic else ''
    def sheet(name,headers,rows,widths):
        ws=book.create_sheet(name);ws.sheet_view.showGridLines=False
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
        ws['A1']=label+name+' | Q1143 candidate work package'
        ws['A1'].font=Font(name='Calibri',size=17,bold=True,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor='15364A');ws.row_dimensions[1].height=32
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
        ws['A2']='Applicability and assessor sample NOT_YET_CONFIRMED. No automated CSA verdict. Scoped observations are not unique physical devices.'
        ws['A2'].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=34
        table_rows=[headers]+rows;matrix[name]=[]
        for ri,row in enumerate(table_rows,4):
            encoded=[evidence(v) for v in row];matrix[name].append(encoded)
            for ci,value in enumerate(encoded,1):
                cell=ws.cell(ri,ci,value)
                if isinstance(value,str):cell.data_type='s'
                cell.font=Font(name='Calibri',size=11,color='15364A',bold=ri==4)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                cell.fill=PatternFill('solid',fgColor='DCE9EF' if ri==4 else ('FFF2CC' if (name=='Assessor Checks' and ci in (4,5)) or (name=='Printer Review' and ci==6) else ('F2F6F8' if ri%2 else 'FFFFFF')))
            ws.row_dimensions[ri].height=28 if ri==4 else (104 if name=='Printer Review' else 54 if name=='Assessor Checks' else 32)
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        ws.freeze_panes='C5' if name in SHEETS[:2] else 'C5'
        table_range=f'A4:{get_column_letter(len(headers))}{max(4,4+len(rows))}'
        if not rows:
            ws.auto_filter.ref=table_range
        if rows:
            table=Table(displayName='Q1143_'+name.replace(' ','_'),ref=table_range)
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(table)
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:4'
        return ws
    port_rows=[];device_rows=[];review=[];checks=[];anchors={}
    columns={k:[v[0] for v in model['records'][k]['schema']] for k in PRODUCTS}
    for g in model['groups']:
        ref=g['reference'];native=(g['ports'] or g['devices'])[0]
        anchors[ref]={'review':len(review)+5,'checks':len(checks)+5,'ports':len(port_rows)+5 if g['ports'] else None,'devices':len(device_rows)+5 if g['devices'] else None}
        for kind,target in [('ports',port_rows),('devices',device_rows)]:
            for row in g[kind]:target.append([ref]+[row[c] for c in columns[kind]]+['Working row'])
        required=[]
        for port in (21,23,80):
            count=sum(r['protocol']=='tcp' and r['port']==port and r['port_state']=='open' for r in g['ports'])
            required.append(f'TCP {port}: {count} open occurrence(s)' if count else f'TCP {port}: Not observed open; closure not verified')
        all_ports=Counter((r['protocol'],r['port'],r['port_state']) for r in g['ports'])
        other='; '.join(f'{proto}/{port} {state} x{count}' for (proto,port,state),count in sorted(all_ports.items())) or 'No port rows in this input'
        names=sorted({r['hostname'] for r in g['ports']+g['devices'] if r.get('hostname')})
        device=ref+' | '+native['ip_address']+'\n'+(' / '.join(names) if names else 'Name not supplied')
        context=native['scan_id']+'\n'+native['scan_scope']+' / '+native['scan_location']+'\n'+g['reconciliation']
        request='Confirm printer applicability and assessor sample. Check TCP 21/23/80 closure, web administration and password policy separately; validate necessity of other services.'
        review.append([device,context,'\n'.join(required)+'\nOther evidence: '+other,request,'RCSO + Post technical contact; confirm actual owner','AWAITING POST RESPONSE','Open assessor checks'])
        for title,instruction,initial in CHECKS:checks.append([ref+' | '+native['ip_address'],title,instruction,initial,None])
    sheet('Printer Review',['Candidate observation','Scan / reconciliation','Observed port evidence','Requested check','Suggested review team','Post response','Supporting evidence'],review,[27,35,45,48,27,30,25])
    sheet('Assessor Checks',['Candidate observation','Check','What to record','Assessor outcome / response','Evidence / sample rationale / reference'],checks,[27,30,65,30,48])
    sheet('Port Evidence',['Candidate group']+columns['ports']+['Back to review'],port_rows,[18]+[26]*len(columns['ports'])+[20])
    sheet('Device Evidence',['Candidate group']+columns['devices']+['Back to review'],device_rows,[18]+[26]*len(columns['devices'])+[20])
    readme=[['Purpose','Sample-based Q1143 work package; manual assessor judgement. No automatic ownership, vulnerability or compliance finding.'],
        ['Scope',model['post']+' / '+model['run']],['Candidate groups (union)',len(model['groups'])],
        ['Native port occurrences',len(model['data']['ports'])],['Native device-review rows',len(model['data']['devices'])],
        ['Groups in both inputs',sum(bool(g['ports']) and bool(g['devices']) for g in model['groups'])],
        ['Port-only groups',sum(bool(g['ports']) and not g['devices'] for g in model['groups'])],
        ['Device-only groups',sum(bool(g['devices']) and not g['ports'] for g in model['groups'])],
        ['Empty inputs','Both captured inputs are present but empty; no closure/compliance established.' if not model['groups'] else 'No input counts were forced to match.'],
        ['Grouping','Exact producer host grain: '+', '.join(HOST)+'. Group labels are navigation only, not physical-device identities.'],
        ['Control basis','Assessor review wording: assessor tests a sample of controls on a sample of applicable printers. Review 21/23/80 closure; web administration/password policy and unique non-default credentials; stakeholder remediation and secondary Nmap verification.'],
        ['Guide alternative','Supplied alternative: confirmation that the guide is not being applied should be Fail. No such confirmation is inferred here.'],
        ['Interpretation','No open row means Not observed open; closure not verified. TCP80 does not prove a webpage or default credentials. Authentication/redirect does not automatically meet port closure. Other protocols require necessity review, not automatic vulnerability labels.'],
        ['Use','Start with Printer Review; follow links to checks and native evidence. Applicability, sample and all outcomes start unconfirmed/not tested. Never enter passwords. Source evidence sheets retain values; edits are manual assessor records, not new producer certification.'],
        ['Encoding','Native timestamps are transported as original text under reviewed reader UTC guards. Backslashes are doubled, illegal controls shown as \\u tokens; empty, temporal-looking and leading-quote strings are JSON-quoted before escaping. Decode string fields by undoing doubled backslashes and illegal-control \\u tokens, then JSON-decode exactly once only if the recovered text starts with a quotation mark. Other text stays literal; use the native schema for numeric/array types. Finite floats/decimals and large integers are lossless text. Blank native evidence remains null. Formula/error-like text is literal.'],
        ['Preview','Operator preview counts are context only; not constants or certified exact expectations.'],
        ['Delivery status','SYNTHETIC REVIEW; not enterprise execution' if synthetic else 'Captured pinned products; applicability/sample/testing remain assessor decisions.']]
    for port in (21,23,80):
        readme.append([f'TCP {port} open occurrences',sum(r['protocol']=='tcp' and r['port']==port and r['port_state']=='open' for r in model['data']['ports'])])
    for kind,record in model['records'].items():
        for key in ('logical_name','table','version','table_id','generation','execution_id','producer_sha256','write_token','schema'):
            value=record[key];readme.append([kind+' / '+key,json.dumps(value,ensure_ascii=False) if isinstance(value,(list,tuple)) else value])
    sheet('Read Me',['Item','Meaning / provenance'],readme,[34,140])
    def link(sheetname,row,col,target):
        book[sheetname].cell(row,col).hyperlink=target;links[(sheetname,row,col)]=target
        book[sheetname].cell(row,col).font=Font(name='Calibri',size=11,color='0563C1',underline='single')
    for ref,a in anchors.items():
        link('Printer Review',a['review'],7,f"#'Assessor Checks'!A{a['checks']}")
        for row in range(a['checks'],a['checks']+len(CHECKS)):link('Assessor Checks',row,1,f"#'Printer Review'!A{a['review']}")
        for name,kind,rows in [('Port Evidence','ports',port_rows),('Device Evidence','devices',device_rows)]:
            if a[kind] is not None:
                # Direct navigation from readable working cells to each evidence grain.
                column=3 if kind=='ports' else 1
                link('Printer Review',a['review'],column,f"#'{name}'!A{a[kind]}")
            for rownum,row in enumerate(rows,5):
                if row[0]==ref:link(name,rownum,len(row),f"#'Printer Review'!A{a['review']}")
    buffer=io.BytesIO();book.save(buffer);book.close();raw=buffer.getvalue()
    check=load_workbook(io.BytesIO(raw),data_only=False)
    try:
        require(check.sheetnames==SHEETS,'REOPEN_SHEETS')
        for name,rows in matrix.items():
            for ri,row in enumerate(rows,4):
                for ci,value in enumerate(row,1):
                    cell=check[name].cell(ri,ci)
                    require(cell.value==value and (not isinstance(value,str) or cell.data_type=='s'),'REOPEN_VALUE')
            for table in check[name].tables.values():require(table.autoFilter is not None and table.autoFilter.ref==table.ref,'REOPEN_FILTER')
        for (name,row,col),target in links.items():
            require(check[name].cell(row,col).hyperlink.target==target,'REOPEN_LINK')
            target_sheet,coord=target[1:].split('!');require(check[target_sheet.strip("'")][coord].value is not None,'LINK_TARGET')
    finally:check.close()
    return raw


def publish(spark, execution, binding_api, *, expected_post, expected_run, approved_report_dir, workspace_client=None):
    """Explicit later enterprise call. New Workspace file, never overwrite the full report."""
    require(re.fullmatch(r'[a-z0-9][a-z0-9_-]*',expected_post) is not None,'POST_SLUG_REQUIRED')
    require(approved_report_dir.startswith('/Users/') and '..' not in approved_report_dir.split('/'),'APPROVED_WORKSPACE_REPORT_DIR_REQUIRED')
    try:
        from databricks.sdk import WorkspaceClient
        from databricks.sdk.service.workspace import ImportFormat, ExportFormat
    except ImportError as exc:raise ContractError('DATABRICKS_SDK_REQUIRED_NO_INSTALL') from exc
    client=workspace_client or WorkspaceClient()
    model=read_inputs(spark,execution,binding_api,expected_post=expected_post,expected_run=expected_run)
    raw=build_workbook(model)
    name='RCAIP_'+expected_post+'_Q1143_Printer_Review_'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')+'_'+uuid4().hex+'.xlsx'
    target=approved_report_dir.rstrip('/')+'/'+name
    client.workspace.upload(path=target,content=io.BytesIO(raw),format=ImportFormat.AUTO,overwrite=False)
    with client.workspace.download(path=target,format=ExportFormat.AUTO) as stream:
        result=bytearray()
        while len(result)<=len(raw):
            chunk=stream.read(min(1024*1024,len(raw)+1-len(result)))
            if not chunk:break
            result.extend(chunk)
    require(bytes(result)==raw,'PUBLISHED_BYTES_MISMATCH')
    return {'workspace_path':target,'sha256':hashlib.sha256(raw).hexdigest(),'published_bytes_verified':True,
            'status':'Q1143 candidate review; no automatic CSA judgement'}
