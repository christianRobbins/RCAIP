"""MDE population presentation of explicitly declared reconstruction inputs.

Only presentation: decisions, narratives, cases and routing come from the pinned
rules and original results. No data acquisition, matching, currentness refresh or
import-time publication. Python/openpyxl follows the accepted IP delivery path.
"""
from collections import Counter
from io import BytesIO
from pathlib import Path

import rcaip_mde_reconstruction as reconstruction

bridge = reconstruction.bridge
require = bridge.require
evidence_value = bridge.retained.consumer.evidence_value
SHEETS = ('RCSO Summary', 'MDE Population', 'Case review', 'Case evidence',
          'Members', 'Shared support', 'Model review', 'Reconciliation')
CMDB = ('UNESTABLISHED: paired CMDB/tool assignments, governed identity links, '
        'observation times and assignment meanings are not supplied. Names and '
        'machine groups alone do not prove Post assignment. No discrepancy or '
        'CMDB action is inferred; reverse inventory remains separate work.')
NEXT_LABELS = {
    'CURRENT_IDENTITY_EVIDENCE_ENRICHMENT': 'Resolve insufficient identity evidence before drawing a control conclusion.',
    'SET_ASIDE_FROM_MDE_ABSENCE_WORKLOAD': 'Set aside from missing-MDE workload; preserve original source state.',
    'LIFECYCLE_IDENTITY_DATA_RECONCILIATION': 'Reconcile lifecycle, identity and historical administrative evidence.',
    'CURRENT_IDENTITY_ATTRIBUTION_RESOLUTION': 'Resolve the identity and attribution relationship first.',
    'CURRENT_CONTROL_APPLICABILITY_AND_OWNER_TEST': 'Use the recorded applicability and owner gate shown here.',
}
COHORT_LABELS = {
    'CURRENT_MDE_EVIDENCE_CONFIRMED': 'Current MDE evidence recovered',
    'NONCURRENT_LIFECYCLE_OR_EVIDENCE_LIMITATION': 'Noncurrent or lifecycle evidence limitation',
    'CURRENT_IDENTITY_TEMPORAL_REVIEW': 'Identity and time relationship unresolved',
    'CURRENT_IDENTITY_EVIDENCE_LIMITATION': 'Identity evidence insufficient',
    'CURRENT_STALE_OR_HISTORICAL_MDE_ONLY': 'Only stale or historical MDE evidence',
}


def build_workbook(bundle, *, original_raw, synthetic):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.utils import get_column_letter

    require(type(synthetic) is bool, 'Explicit synthetic label required')
    b = reconstruction.validate_bundle(bundle, original_raw)
    q = sorted(b['qualification'], key=lambda r: r['asset_key'])
    gates = {r['asset_key']: r for r in b['current_gate']}
    cases, members, support, review, accounting = [b['original_products'][p] for p in bridge.check.PRODUCTS[14]]
    cases = sorted(cases, key=lambda r: r['case_id'])
    members = sorted(members, key=lambda r: (r['case_id'], r['asset_key']))
    member_index = {r['asset_key']: r for r in members}
    review_index = {r['asset_key']: r for r in review}
    columns = bridge.check.contracts({int(k): v for k, v in b['original_sources'].items()})
    wb = Workbook(); wb.remove(wb.active)
    label = 'SYNTHETIC TEST DATA / NOT A example_post_a ASSESSMENT' if synthetic else 'DECLARED RECONSTRUCTION / RETAINED MDE RESULTS'
    expected = {}

    def put(ws, r, c, value):
        value = evidence_value(value)
        cell = ws.cell(r, c, value)
        if isinstance(value, str): cell.data_type = 's'
        cell.font = Font(name='Arial', size=11, color='18354A')
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        expected[(ws.title, r, c)] = value
        return cell

    def link(cell, sheet, row):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A{row}", display=str(cell.value))
        cell.font = Font(name='Arial', size=11, color='0563C1', underline='single')

    def base(name):
        ws = wb.create_sheet(name); ws.sheet_view.showGridLines = False
        for i in range(1, 9): ws.column_dimensions[get_column_letter(i)].width = 24
        put(ws, 2, 1, name).font = Font(name='Arial', size=16, bold=True, color='18354A')
        ws.merge_cells('A2:H2')
        put(ws, 3, 1, label); ws.merge_cells('A3:H3'); ws.row_dimensions[3].height = 28
        put(ws, 4, 1, reconstruction.LIMITATION); ws.merge_cells('A4:H4'); ws.row_dimensions[4].height = 56
        put(ws, 5, 1, 'Historical context only. No current-device or reopened formal CSA certification.'); ws.merge_cells('A5:H5')
        put(ws, 6, 1, 'Scope: ' + b['manifest']['post'] + ' / ' + b['manifest']['run']); ws.merge_cells('A6:H6')
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def table(ws, headers, rows, start=8, name=None):
        require(len(headers) == len(set(headers)) and all(isinstance(h, str) and h for h in headers), 'Presentation headers invalid')
        for c, header in enumerate(headers, 1):
            cell = put(ws, start, c, header)
            cell.font = Font(name='Arial', size=11, color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='18354A')
            ws.column_dimensions[get_column_letter(c)].width = 26 if c > 8 else ws.column_dimensions[get_column_letter(c)].width
        ws.row_dimensions[start].height = 42
        for r, row in enumerate(rows, start + 1):
            require(len(row) == len(headers), 'Presentation row width mismatch')
            for c, value in enumerate(row, 1): put(ws, r, c, value)
            ws.row_dimensions[r].height = 100 if ws.title == 'MDE Population' else 60
        if rows:
            ref = f'A{start}:{get_column_letter(len(headers))}{start + len(rows)}'
            t = Table(displayName=name or ws.title.replace(' ', ''), ref=ref)
            t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(t)
        else:
            put(ws, start + 2, 1, 'No retained rows in this product. This is not evidence of a new live assessment.'); ws.merge_cells(start_row=start+2,start_column=1,end_row=start+2,end_column=8)
        ws.freeze_panes = f'B{start+1}'

    for name in SHEETS: base(name)
    pop_rows = {r['asset_key']: i + 9 for i, r in enumerate(q)}
    case_rows = {r['case_id']: 9 + 13*i for i, r in enumerate(cases)}
    member_rows = {r['asset_key']: i + 9 for i, r in enumerate(members)}
    review_rows = {r['asset_key']: i + 9 for i, r in enumerate(review)}
    case_evidence_rows = {r['case_id']: i + 9 for i, r in enumerate(cases)}
    summary = wb['RCSO Summary']
    counts = Counter(r['mde_qualification_cohort'] for r in q)
    summary_rows = [['Starting members with original MDE MISSING state', len(q),
                     'Every Search Input candidate remains in the population. MISSING is the original source state, not an enterprise absence conclusion.']]
    for code, n in sorted(counts.items()):
        meaning = next(r['mde_qualification_meaning'] for r in q if r['mde_qualification_cohort'] == code)
        require(code in COHORT_LABELS, 'Unsupported population display code')
        summary_rows.append([COHORT_LABELS[code], n, meaning])
    table(summary, ['Population outcome', 'Members', 'What the evidence established'], summary_rows, name='MDEOutcomeSummary')
    summary.column_dimensions['A'].width = 46; summary.column_dimensions['B'].width = 12; summary.column_dimensions['C'].width = 85
    for r in range(9, 9+len(summary_rows)): summary.row_dimensions[r].height = 78
    cursor = 11 + len(summary_rows)
    put(summary, cursor, 1, 'Remaining original cases and recipients').font = Font(name='Arial', size=13, bold=True)
    table(summary, ['Original case', 'Affected members', 'Residual question', 'Recorded recipient', 'Recorded next action'],
          [[r['case_id'], r['affected_assets'], r['residual_concern'], r['primary_customer'], r['recommended_treatment']] for r in cases],
          start=cursor+1, name='MDECaseRecipients')
    summary.column_dimensions['D'].width=30; summary.column_dimensions['E'].width=70
    for i, r in enumerate(cases, cursor+2):
        link(summary.cell(i,1), 'Case review', case_rows[r['case_id']]); summary.row_dimensions[i].height=130
    cursor += len(cases)+5
    notes = [CMDB,
             'Recipients above are original Cell 14 routing. RCSO review and Post awareness do not transfer service-owner work to the Post or CMDB.',
             'Model-review remainder: ' + str(len(review)) + ' retained members. No replacement cases or owners are assigned.',
             'Recovered current MDE evidence means current relative to the pinned historical evidence, not device health today.',
             'Parked MDE-history expansion and the full reverse inventory/customer comparison remain outside this delivery.']
    for n, note in enumerate(notes, cursor):
        put(summary,n,1,note); summary.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5); summary.row_dimensions[n].height=54
    link(put(summary,cursor+6,1,'Review every starting member'),'MDE Population',8)
    link(put(summary,cursor+7,1,'Read accounting and provenance'),'Reconciliation',8)

    ws=wb['MDE Population']
    qcols=[f['name'] for f in b['derived_schemas']['qualification']]
    gcols=[f['name'] for f in b['derived_schemas']['current_gate']]
    headings=['Member', 'Observed name', 'Original condition', 'Enterprise evidence explanation', 'Recorded gate explanation',
              'Supported next question', 'Original recipient', 'Original case', 'Retained MDE evidence', 'Inventory comparison']
    rows=[]
    for r in q:
        key=r['asset_key']; g=gates.get(key); m=member_index.get(key); rev=review_index.get(key)
        require(r['mde_next_gate'] in NEXT_LABELS, 'Unsupported next-step display code')
        # The native model-review projection has no formed-case treatment or
        # recipient. Management context must not become a customer assignment.
        action = (m['recommended_treatment'] if m else
                  'Original case model left this condition unresolved. No formed case, case treatment or assigned case recipient was supplied.' if rev else
                  g['mde_current_gate_meaning'] if g else NEXT_LABELS[r['mde_next_gate']])
        customer = m['primary_customer'] if m else None if rev else r['case_primary_customer']
        rows.append([key,r['canonical_name'],r['original_mde_assurance_state'],r['mde_qualification_meaning'],
                     g['mde_current_gate_meaning'] if g else 'Outside the source-defined Cell 11 current subset',
                     action,customer,m['case_id'] if m else None,
                     'Latest evidence time (source text): '+str(r['latest_mde_evidence_timestamp'])+'\n'
                     'MDE entities - current: '+str(r['current_mde_entity_count'])+', stale: '+str(r['stale_mde_entity_count'])
                     +', historical: '+str(r['historical_mde_entity_count'])+'\nNo timezone or present-day status inferred.',
                     'UNESTABLISHED - see RCSO Summary']
                    +[r[c] for c in qcols]+[g[c] if g else None for c in gcols])
    table(ws,headings+['qualification.'+c for c in qcols]+['gate.'+c for c in gcols],rows)
    for c,width in [('A',38),('B',30),('D',65),('E',60),('F',65),('G',30),('H',26),('I',55),('J',30)]: ws.column_dimensions[c].width=width
    for i,r in enumerate(q,9):
        if r['asset_key'] in member_index:
            link(ws.cell(i,8),'Case review',case_rows[member_index[r['asset_key']]['case_id']])
        if r['asset_key'] in review_rows:
            link(ws.cell(i,1),'Model review',review_rows[r['asset_key']])
        if r['mde_qualification_cohort']=='CURRENT_MDE_EVIDENCE_CONFIRMED': ws.cell(i,4).fill=PatternFill('solid',fgColor='E5F0E9')
        else: ws.cell(i,4).fill=PatternFill('solid',fgColor='FFF0D7')

    cards=wb['Case review']
    for c in cases:
        r=case_rows[c['case_id']]
        fields=[('Original case',c['case_id']),('Recipient',c['primary_customer']),('Condition',c['relationship_outcome']),
                ('Residual concern',c['residual_concern']),('Supported next decision',c['recommended_treatment']),
                ('Priority basis',c['priority_basis']),('Affected members',c['affected_assets']),
                ('Original rule',c['case_rule_key']),('Case evidence','Open native case evidence'),
                ('Members','Open first retained member'),('Routing boundary','Original enterprise referral; Post awareness does not assign Post action.')]
        for offset,(label,value) in enumerate(fields):
            put(cards,r+offset,1,label).font=Font(name='Arial',size=11,bold=True)
            put(cards,r+offset,2,value);cards.merge_cells(start_row=r+offset,start_column=2,end_row=r+offset,end_column=8)
            cards.row_dimensions[r+offset].height=44 if offset in (2,3,4,5,10) else 26
        link(cards.cell(r+8,2),'Case evidence',case_evidence_rows[c['case_id']])
        first=next(m for m in members if m['case_id']==c['case_id'])
        link(cards.cell(r+9,2),'Members',member_rows[first['asset_key']])
    if not cases: put(cards,9,1,'No original cases. No new case formation.')
    for sheet, product, data in zip(('Case evidence','Members','Shared support','Model review'),bridge.check.PRODUCTS[14][:4],(cases,members,support,review)):
        cols=columns[product]; table(wb[sheet],cols,[[r[c] for c in cols] for r in data])
        for i,r in enumerate(data,9):
            if r.get('case_id') in case_rows and 'case_id' in cols:
                link(wb[sheet].cell(i,cols.index('case_id')+1),'Case review',case_rows[r['case_id']])
            key=r.get('asset_key',r.get('case_member_asset_key'))
            field='asset_key' if 'asset_key' in cols else 'case_member_asset_key'
            if key in pop_rows and field in cols: link(wb[sheet].cell(i,cols.index(field)+1),'MDE Population',pop_rows[key])
    recon=wb['Reconciliation']
    acols=columns[bridge.check.PRODUCTS[14][4]]
    table(recon,acols,[[r[c] for c in acols] for r in sorted(accounting,key=lambda r:r['stage_order'])],name='OriginalAccounting')
    recon.column_dimensions['B'].width=65;recon.column_dimensions['D'].width=85
    for n,r in enumerate(sorted(accounting,key=lambda r:r['stage_order']),9):
        link(recon.cell(n,2),'MDE Population' if r['stage_order']<6 else 'Case evidence',8)
    facts=[('Provenance mode',b['authority']),('Limitation',b['limitation']),('Original notebook SHA-256',b['original_notebook_sha256']),
           ('Rule version',b['rule_version']),('Original products',b['original_kind']),('Member maps',b['member_map_kind']),
           ('Selection time (not observation time)',b['selection']['selection_provenance']['selected_at_utc']),
           ('CMDB pairing',CMDB),('Definition: member','One candidate asset_key in Search Input. Entity counts are not asset counts.'),
           ('Definition: case','Original Cell 14 deterministic enterprise referral. Cases are not the starting population.'),
           ('Definition: support','Shared durable-identity relationship; not an identity merge or an additional case member.'),
           ('Definition: original accounting','Stage counts have distinct grains. Do not sum sequential stages or cases with members.'),
           ('Definition: time','All source timestamp text retained. No report-time substitution or timezone inference.')]
    for name,v in b['producer_fingerprints'].items(): facts.append(('Pinned Cell '+name+' fingerprint',v))
    for name,binding in b['selection']['inputs'].items():
        for field in ('table','table_id','version','schema_sha256'): facts.append((name+' / '+field,binding[field]))
        facts.append((name+' / payload SHA-256',b['input_payload_sha256'][name]))
    table(recon,['Evidence / definition','Value'],facts,start=21,name='ReconstructionProvenance')
    for i in range(22,22+len(facts)): recon.row_dimensions[i].height=72
    recon.column_dimensions['A'].width=52;recon.column_dimensions['B'].width=100
    wb._rcaip_expected_cells=expected
    return wb


def save_workbook(bundle, destination, *, approved_directory, original_raw, synthetic):
    """Explicit non-overwriting output; exact reopened values and table/filter checks."""
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    wb=build_workbook(bundle,original_raw=original_raw,synthetic=synthetic)
    stream=BytesIO();wb.save(stream);raw=stream.getvalue()
    reopened=load_workbook(BytesIO(raw),data_only=False)
    require(tuple(reopened.sheetnames)==SHEETS,'Workbook sheet contract')
    for (sheet,r,c),value in wb._rcaip_expected_cells.items():
        cell=reopened[sheet].cell(r,c)
        require(cell.value==value and cell.data_type not in ('f','e'),'Workbook exact representation mismatch')
    for ws in reopened:
        for t in ws.tables.values():
            require(t.autoFilter is not None and t.ref==t.autoFilter.ref,'Workbook filter/table mismatch')
            x,y,z,_=range_boundaries(t.ref)
            require([c.name for c in t.tableColumns]==[ws.cell(y,i).value for i in range(x,z+1)],'Workbook table header mismatch')
        for row in ws:
            for cell in row:
                if cell.hyperlink:
                    match=bridge.re.fullmatch(r"'([^']+)'!([A-Z]+[0-9]+)",cell.hyperlink.location or '')
                    require(match and match[1] in reopened and reopened[match[1]][match[2]].value is not None,'Workbook navigation mismatch')
    dest,root=Path(destination).resolve(),Path(approved_directory).resolve()
    require(root.is_dir() and dest.parent==root and dest.suffix=='.xlsx' and not dest.exists(),'New approved workbook destination required')
    with dest.open('xb') as out: out.write(raw)
    require(dest.read_bytes()==raw,'Workbook downloaded-byte mismatch')
    return dict(state='RECONSTRUCTION_WORKBOOK_WRITTEN',sha256=bridge.digest(raw),
                mode=reconstruction.MODE,limitation=reconstruction.LIMITATION,
                desktop_excel='NOT_TESTED',live_assessment='NOT_CERTIFIED',formal_csa='NOT_REOPENED')
