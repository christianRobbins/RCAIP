"""Cell 24 diagnostic presentation. Python/openpyxl is the sole implementation route.

D0026-D0029 / D0033-D0034 / D0063-D0064: retain producer decisions and native
case/member/closure grains. No source acquisition, producer execution, automatic
publication or import-time activity. The JavaScript prototype is reference only.
Would this exact analytical logic run for Paris? YES; no Post-specific counts.
"""
from collections import Counter
from collections.abc import Mapping
from copy import deepcopy
from datetime import date, datetime, timezone
from decimal import Decimal
import json
import math
import re


PRODUCTS = {'cases': 'RCAIP_DIAG001_UNRESOLVED_IP_CASE_DF', 'members': 'RCAIP_DIAG001_UNRESOLVED_IP_CASE_MEMBER_DF', 'closure': 'RCAIP_DIAG001_UNRESOLVED_IP_BRANCH_CLOSURE_DF'}
CASE_COLUMNS = ('case_instance_id', 'case_family', 'case_signature', 'case_class', 'case_subtype', 'primary_customer', 'secondary_customer_scope', 'priority_band', 'confidence_band', 'condition_title', 'executive_summary', 'why_it_matters', 'evidence_summary', 'benign_explanations_considered', 'residual_concern', 'recommended_action', 'action_owner', 'priority_basis', 'confidence_basis', 'affected_member_count', 'affected_asset_count', 'source_record_count', 'relationship_count', 'related_assessment_question_ids', 'trend_state', 'disposition_status', 'post_profile_key', 'assessment_run_id', 'case_provenance', 'affected_asset_keys', 'candidate_ips', 'candidate_tenable_uuids', 'network_scopes', 'active_source_patterns', 'identity_bases', 'identity_strengths', 'earliest_member_tenable_observation', 'latest_member_tenable_observation', 'latest_relevant_nmap_scan_end_time_utc')
MEMBER_COLUMNS = ('case_instance_id', 'member_type', 'member_id', 'asset_id', 'membership_reason', 'relationship_outcome', 'identity_confidence', 'evidence_confidence', 'member_priority_modifier', 'source_provenance', 'post_profile_key', 'assessment_run_id', 'candidate_ip', 'candidate_tenable_uuid', 'candidate_tenable_last_observed', 'candidate_tenable_age_days', 'latest_relevant_nmap_scan_end_time_utc', 'tenable_vs_latest_nmap_hours', 'candidate_tenable_acr_score', 'tenable_device_type_hint', 'network_scope', 'operational_boundary', 'management_model', 'management_owner', 'identity_basis', 'identity_strength', 'source_mask_active', 'current_observation_status', 'post_visibility_state', 'attribution_confidence', 'attribution_conflict_flag', 'primary_customer', 'residual_treatment_class', 'case_subtype', 'customer_routing_basis', 'recommended_next_action')
CLOSURE_COLUMNS = ['stage_order', 'stage', 'measurement_unit', 'count', 'interpretation', 'treatment_or_customer', 'active_workload']
STAGES = [['STARTING_ACTIVE_UNRESOLVED_IP', 'ANALYTICAL_ASSET_ROW', 'ANALYTICAL_INPUT', 'NO'], ['TENABLE_SOURCE_LOCAL_IDENTITIES', 'IDENTITY_GROUP', 'IDENTITY_RECONCILIATION', 'NO'], ['STRONGER_SOURCE_LOCAL_UUID_RELATIONSHIP', 'IDENTITY_GROUP', 'EXPLANATORY_RECONCILIATION', 'NO'], ['STRONGER_EXACT_IP_RELATIONSHIP', 'IDENTITY_GROUP', 'EXPLANATORY_RECONCILIATION', 'NO'], ['TRUE_NO_LINK_TENABLE_IDENTITIES', 'IDENTITY_GROUP', 'DEVICE_LIFECYCLE_QUALIFICATION', 'NO'], ['DEVICE_RECENCY_NOT_READY_FOR_WORKLOAD', 'IDENTITY_GROUP', 'QUALIFIED_NON_WORKLOAD', 'NO'], ['FRESH_GENERAL_PURPOSE_TEMPORAL_GATE', 'IDENTITY_GROUP', 'TEMPORAL_QUALIFICATION', 'NO'], ['CURRENTNESS_VALIDATION_BEFORE_ENUMERATION', 'IDENTITY_GROUP', 'CURRENTNESS_VALIDATION', 'NO'], ['TEMPORALLY_CURRENT_RESIDUAL', 'IDENTITY_GROUP', 'CUSTOMER_QUALIFICATION', 'NO'], ['RCSO_ROUTED_RESIDUAL_MEMBERS', 'CASE_MEMBER', 'RCSO', 'YES'], ['POST_ROUTED_RESIDUAL_MEMBERS', 'CASE_MEMBER', 'POST_DT', 'YES'], ['ENTERPRISE_ROUTED_RESIDUAL_MEMBERS', 'CASE_MEMBER', 'ENTERPRISE', 'YES'], ['FINAL_UNRESOLVED_IP_ASSURANCE_CASES', 'ASSURANCE_CASE', 'CASE_SET', 'YES'], ['FINAL_RCSO_CASES', 'ASSURANCE_CASE', 'RCSO', 'YES']]

class InputContractError(ValueError):
    """A required producer contract is missing or inconsistent; publication stops."""


def require(ok, reason):
    if not ok:
        raise InputContractError(reason)


def text(value):
    return isinstance(value, str) and bool(value.strip())


def count(value):
    return type(value) is int and value >= 0


def evidence_value(value):
    """Lossless display encoding. No formulas, rounding, dropped controls or null fill.

    Numeric arrays are intentionally unsupported, including [-0.0]. Cell 24
    publishes string arrays. Reject rather than silently stringify numeric arrays.
    Timestamps retain source text; native Python dates are ISO text. Spark temporal
    values are cast BEFORE collect, avoiding Python datetime range constraints.
    """
    if value is None or type(value) is bool:
        return value
    if type(value) is int:
        return value if abs(value) < 10**15 else str(value)
    if type(value) is float:
        require(math.isfinite(value), 'Nonfinite floating-point evidence')
        return repr(value)
    if isinstance(value, Decimal):
        require(value.is_finite(), 'Nonfinite decimal evidence')
        return str(value)
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    if isinstance(value, (list, tuple)):
        require(all(x is None or isinstance(x, str) for x in value),
                'Unsupported numeric/nested array transport; only string/null arrays are supported')
        value = json.dumps(value, ensure_ascii=False, separators=(',', ':'))
    require(isinstance(value, str), f'Unsupported evidence type: {type(value).__name__}')
    if value == '' or re.match(r'^\d{4,}-\d\d-\d\d(?:T| |$)', value):
        value = json.dumps(value, ensure_ascii=False)
    value = value.replace('\\', '\\\\')
    value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]',
                   lambda m: '\\u%04x' % ord(m[0]), value)
    require(not re.search(r'[\ud800-\udfff]', value), 'Unpaired Unicode surrogate')
    require(len(value.encode('utf-16-le')) // 2 <= 32767, 'Evidence exceeds Excel cell limit; no truncation')
    return value


READY = 'RCAIP_DIAG001_CELL24_READY'
STATUS = 'RCAIP_DIAG001_UNRESOLVED_IP_BRANCH_STATUS'
POST = 'RCAIP_POST_SLUG'
RUN = 'RCAIP_RUN_ID'
CUSTOMERS = {'RCSO', 'POST_DT', 'ENTERPRISE_SERVICE_OWNER', 'ENTERPRISE_DATA_OWNER', 'RCAIP_ENGINEERING', 'MONITORING'}
CLASSES = {'ARCHITECTURE_REVIEW', 'FOCUSED_POST_VALIDATION', 'MONITORING_OR_TREND', 'ENTERPRISE_RECONCILIATION_REFERRAL', 'EVIDENCE_LIMITATION'}
SYNTHETIC_LABEL = 'SYNTHETIC TEST DATA / NOT A example_post_a ASSESSMENT'


def _authority(runtime, expected_post, expected_run):
    require(isinstance(runtime, Mapping), 'Producer namespace/mapping required')
    require(runtime.get(READY) is True, f'{READY}: must be True')
    require(runtime.get(STATUS) == 'CLOSED_FOR_CURRENT_RUN', f'{STATUS}: must be CLOSED_FOR_CURRENT_RUN')
    require(text(expected_post) and text(expected_run), 'Expected Post/run authority required')
    require(text(runtime.get(POST)) and runtime[POST].strip().lower() == expected_post.strip().lower(), 'Runtime Post mismatch')
    require(runtime.get(RUN) == expected_run, 'Runtime run mismatch')


def validate_cell24(runtime, expected_post, expected_run, *, max_rows=100000):
    """Validate already-materialized inputs. Never reads Spark or changes decisions."""
    _authority(runtime, expected_post, expected_run)
    require(type(max_rows) is int and 0 < max_rows <= 100000, 'Invalid presentation row bound')
    data = {}
    for kind, required in [('cases', CASE_COLUMNS), ('members', MEMBER_COLUMNS), ('closure', CLOSURE_COLUMNS)]:
        name = PRODUCTS[kind]
        p = runtime.get(name)
        require(isinstance(p, Mapping) and isinstance(p.get('columns'), (list, tuple)) and isinstance(p.get('rows'), list), f'{name}: missing product/schema')
        cols = p['columns']
        require(all(text(c) for c in cols) and len(cols) == len(set(cols)), f'{name}: invalid/duplicate columns')
        require(set(required) <= set(cols), f'{name}: incomplete producer schema')
        require(len(p['rows']) <= max_rows, f'{name}: row limit exceeded; no rows dropped')
        for index, row in enumerate(p['rows']):
            require(isinstance(row, Mapping) and set(row) == set(cols), f'{name}: row {index} schema mismatch')
            for field, value in row.items():
                try:
                    evidence_value(value)
                except InputContractError as exc:
                    raise InputContractError(f'{name}.{field}: {exc}') from exc
            if kind != 'closure':
                require(text(row['post_profile_key']) and row['post_profile_key'].strip().lower() == expected_post.strip().lower(), f'{name}: foreign/missing Post')
                require(row['assessment_run_id'] == expected_run, f'{name}: foreign/missing run')
        data[kind] = deepcopy(p)
    cases, members, closure = [data[k]['rows'] for k in ('cases', 'members', 'closure')]
    by_case = {}
    for c in cases:
        require(text(c['case_instance_id']) and c['case_instance_id'] not in by_case, 'Duplicate/blank case ID')
        for key in ('case_family', 'case_signature', 'case_subtype', 'priority_band', 'confidence_band',
                    'condition_title', 'executive_summary', 'why_it_matters', 'evidence_summary',
                    'benign_explanations_considered', 'residual_concern', 'recommended_action',
                    'action_owner', 'priority_basis', 'confidence_basis', 'case_provenance'):
            require(text(c[key]), f'Case missing {key}')
        require(c['primary_customer'] in CUSTOMERS and c['action_owner'] == c['primary_customer'], 'Case customer/action-owner mismatch')
        require(c['case_class'] in CLASSES, 'Unresolved identity cannot become a control finding/unknown class')
        require(c['source_record_count'] is None, 'Cell 24 raw source-record count must remain null')
        require(all(count(c[k]) for k in ('affected_member_count', 'affected_asset_count', 'relationship_count')), 'Invalid case counts')
        by_case[c['case_instance_id']] = []
    ids, pairs = {}, set()
    for m in members:
        require(m['case_instance_id'] in by_case, 'Orphan case member')
        require(text(m['member_id']) and text(m['asset_id']), 'Missing member/asset identity')
        require(m['member_type'] == 'PROBABLE_ASSET_IDENTITY', 'Unexpected member type')
        pair = (m['case_instance_id'], m['member_id'], m['asset_id'])
        require(pair not in pairs, 'Duplicate case/member/asset relationship')
        pairs.add(pair)
        require(m['member_id'] not in ids or ids[m['member_id']] == m['case_instance_id'], 'Member assigned multiple cases')
        ids[m['member_id']] = m['case_instance_id']
        by_case[m['case_instance_id']].append(m)
        for key in ('membership_reason', 'source_provenance', 'customer_routing_basis', 'recommended_next_action'):
            require(text(m[key]), f'Member missing {key}')
    for c in cases:
        ms = by_case[c['case_instance_id']]
        require(ms and len(ms) == c['affected_member_count'] == c['relationship_count'], 'Case member/relationship counts do not reconcile')
        assets = {m['asset_id'] for m in ms}
        require(isinstance(c['affected_asset_keys'], list) and all(text(x) for x in c['affected_asset_keys']), 'Invalid affected asset keys')
        require(len(assets) == c['affected_asset_count'] == len(c['affected_asset_keys']) and assets == set(c['affected_asset_keys']), 'Case asset membership does not reconcile')
        for m in ms:
            require(m['primary_customer'] == c['primary_customer'] and m['case_subtype'] == c['case_subtype'] and
                    m['recommended_next_action'] == c['recommended_action'] and m['evidence_confidence'] == c['confidence_band'], 'Case/member decision mismatch')
    # Producer equivalent: member_sum == member_rows == asset_sum. Do NOT
    # deduplicate or demand global asset uniqueness across DIFFERENT cases.
    require(sum(c['affected_member_count'] for c in cases) == len(members), 'Aggregate affected_member_count != member-row count')
    require(sum(c['affected_asset_count'] for c in cases) == len(members), 'Aggregate affected_asset_count != member-row count (Cell 24 certification)')
    require(len(closure) == len(STAGES), 'Incomplete/extra closure stages')
    require(all(type(s['stage_order']) is int for s in closure), 'Invalid closure stage order')
    stages = sorted(closure, key=lambda s: s['stage_order'])
    for index, s in enumerate(stages):
        require(s['stage_order'] == index + 1 and [s[k] for k in ('stage', 'measurement_unit', 'treatment_or_customer', 'active_workload')] == STAGES[index], 'Closure stage/unit/treatment mismatch')
        require(count(s['count']) and text(s['interpretation']), 'Invalid closure count/interpretation')
    n = lambda i: stages[i - 1]['count']
    require(n(2) == n(3) + n(4) + n(5) and n(5) == n(6) + n(7), 'Closure identity stages do not reconcile')
    # Cell 22 temporal-other groups are not published in closure. No invented zero.
    require(n(7) >= n(8) + n(9) and n(1) >= n(2), 'Inconsistent closure qualification counts')
    require(n(9) == len(ids) and n(13) == len(cases) and n(14) == sum(c['primary_customer'] == 'RCSO' for c in cases), 'Closure residual/case counts do not reconcile')
    for stage, owners in [(10, {'RCSO'}), (11, {'POST_DT'}), (12, {'ENTERPRISE_SERVICE_OWNER', 'ENTERPRISE_DATA_OWNER'})]:
        require(n(stage) == sum(m['primary_customer'] in owners for m in members), 'Closure customer-member counts do not reconcile')
    data.update(post=expected_post.strip().lower(), run=expected_run,
                authority={key: runtime[key] for key in (READY, STATUS, POST, RUN)})
    return data


def collect_cell24_session(producer_namespace, *, max_rows=100000):
    """Explicit same-session adapter. Invoke ONLY in the verified certified producer
    namespace, not from reconstructed dictionaries or arbitrary retained products.

    Reads only existing named outputs/scalars; no tables, SQL source queries,
    Config/producers, toPandas or toLocalIterator. Spark-session identity is checked
    before collection; references and authority must remain unchanged afterward.
    This does not invent a certificate for unscoped closure. The caller must first
    establish that this is the original certified Cell 24 session/result. Existing
    flags and Spark-session identity alone cannot detect a manually substituted
    old closure object. Cached lineage may reevaluate if evicted: authorization
    must establish safe retained inputs before calling; this adapter never replays
    a notebook to obtain them. No readiness flags are set.
    """
    require(isinstance(producer_namespace, Mapping), 'Certified producer namespace required')
    names = [*PRODUCTS.values(), READY, STATUS, POST, RUN]
    for name in names:
        require(name in producer_namespace, f'Missing producer-session object: {name}')
    refs = {name: producer_namespace[name] for name in names}
    _authority(refs, refs[POST], refs[RUN])
    require(type(max_rows) is int and 0 < max_rows <= 100000, 'Invalid collection row bound')
    frames = [refs[name] for name in PRODUCTS.values()]
    for name, frame in zip(PRODUCTS.values(), frames):
        require(hasattr(frame, 'sparkSession') and hasattr(frame, 'schema'), f'{name}: not a Spark DataFrame')
    session = frames[0].sparkSession
    require(all(frame.sparkSession is session for frame in frames), 'Cell 24 products belong to different Spark sessions')
    zone = session.conf.get('spark.sql.session.timeZone')
    require(text(zone), 'Spark timestamp session timezone unavailable')
    # Inspect every schema before collecting any product. Native timestamps need
    # an unambiguous UTC clock; timestamp_ntz, dates and strings are zone-free here.
    timestamp_fields = [f'{name}.{field.name}'
                        for name, frame in zip(PRODUCTS.values(), frames)
                        for field in frame.schema.fields
                        if field.dataType.typeName() == 'timestamp']
    if timestamp_fields:
        require(zone in ('UTC', 'Etc/UTC', 'Z', '+00:00'),
                f'Native timestamp transport requires UTC, Etc/UTC, Z or +00:00; '
                f'session timezone {zone!r} is unsupported: {", ".join(timestamp_fields)}')
    runtime = {k: refs[k] for k in (READY, STATUS, POST, RUN)}
    for name, frame in zip(PRODUCTS.values(), frames):
        fields = frame.schema.fields
        columns = [field.name for field in fields]
        require(len(columns) == len(set(columns)), f'{name}: duplicate schema columns')
        projections, numeric_text = [], []
        for field in fields:
            kind = field.dataType.typeName()
            escaped = '`' + field.name.replace('`', '``') + '`'
            if kind in ('float', 'double', 'decimal', 'timestamp', 'timestamp_ntz', 'date'):
                projections.append(f'CAST({escaped} AS STRING) AS {escaped}')
                if kind in ('float', 'double', 'decimal'):
                    numeric_text.append(field.name)
            elif kind == 'array':
                require(field.dataType.elementType.typeName() == 'string', f'{name}.{field.name}: unsupported numeric/nested array transport')
                projections.append(escaped)
            else:
                require(kind in ('string', 'boolean', 'byte', 'short', 'integer', 'long', 'void'), f'{name}.{field.name}: unsupported Spark type {kind}')
                projections.append(escaped)
        try:
            rows = frame.selectExpr(*projections).limit(max_rows + 1).collect()
        except (TimeoutError, ConnectionError, OSError) as exc:
            raise InputContractError(f'{name}: bounded collection failed ({type(exc).__name__}): {exc}') from exc
        # Other Spark/programming/resource exceptions propagate, never become nulls.
        require(len(rows) <= max_rows, f'{name}: collection exceeds {max_rows} rows; no truncation allowed')
        runtime[name] = {'columns': columns, 'rows': [r.asDict(recursive=True) for r in rows]}
        for row in runtime[name]['rows']:
            for col in numeric_text:
                if row[col] is not None:
                    require(isinstance(row[col], str), f'{name}.{col}: numeric cast did not return text')
                    try:
                        finite = Decimal(row[col]).is_finite()
                    except ArithmeticError as exc:
                        raise InputContractError(f'{name}.{col}: invalid numeric text') from exc
                    require(finite, f'{name}.{col}: nonfinite numeric evidence')
    for name in names:
        require(name in producer_namespace and (producer_namespace[name] is refs[name] if name in PRODUCTS.values() else producer_namespace[name] == refs[name]), f'Producer-session object changed during collection: {name}')
    _authority(producer_namespace, refs[POST], refs[RUN])
    result = validate_cell24(runtime, refs[POST], refs[RUN], max_rows=max_rows)
    # Read-only runtime metadata, not invented provenance. Timestamp casts use it.
    require(session.conf.get('spark.sql.session.timeZone') == zone, 'Spark timestamp timezone changed during collection')
    result['timestamp_session_timezone'] = zone
    return result


CARD_FIELDS = [('What appeared wrong', 'condition_title'), ('Summary', 'executive_summary'),
               ('Why it matters', 'why_it_matters'), ('What evidence established', 'evidence_summary'),
               ('Explanations considered', 'benign_explanations_considered'), ('Residual uncertainty', 'residual_concern'),
               ('Supported next action', 'recommended_action'), ('Responsible customer', 'primary_customer'),
               ('Action owner', 'action_owner'), ('Priority', 'priority_band'), ('Priority basis', 'priority_basis'),
               ('Confidence', 'confidence_band'), ('Confidence basis', 'confidence_basis')]


def build_cell24_workbook(data, *, synthetic, generated_at=None):
    """Render validated materialized products with openpyxl. Returns Workbook only.
    Revalidates the products so a mutated/forged intermediate cannot bypass guards.
    Caller explicitly saves; no filesystem writes or Spark operations here.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    require(type(synthetic) is bool, 'Explicit synthetic flag required')
    require(isinstance(data.get('authority'), Mapping), 'Original producer authority snapshot required')
    runtime = {**data['authority'], **{PRODUCTS[k]: data[k] for k in PRODUCTS}}
    checked = validate_cell24(runtime, data['post'], data['run'])
    generated_at = generated_at or datetime.now(timezone.utc).isoformat()
    require(text(generated_at), 'Generation timestamp required')
    wb = Workbook()
    wb.remove(wb.active)
    label = SYNTHETIC_LABEL if synthetic else 'Cell 24 diagnostic presentation / not a production case queue'
    historical = data.get('historical_provenance')
    if historical:
        require(text(historical), 'Historical provenance text required')
        label += ' / RETAINED NOTEBOOK RESULTS REPRODUCTION'
    dark = PatternFill('solid', fgColor='18354A')
    white = Font(name='Arial', size=11, bold=True, color='FFFFFF')

    def put(ws, row, col, value):
        value = evidence_value(value)
        cell = ws.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = 's'  # Includes '=...', '#N/A', apostrophes: literal, no prefix mutation.
        cell.font = Font(name='Arial', size=11)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        return cell

    def link(cell, sheet, row):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A{row}", display=str(cell.value))
        cell.font = Font(name='Arial', size=11, color='0563C1', underline='single')

    def base(name):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        put(ws, 2, 1, label).font = Font(name='Arial', size=14, bold=True, color='8A4B08')
        put(ws, 3, 1, f"{checked['post']} / {checked['run']} / Generated {generated_at}")
        if historical:
            from openpyxl.comments import Comment
            put(ws, 1, 1, 'HISTORICAL EVIDENCE; NOT LIVE AUTHORITY. Provenance in this cell comment.')
            evidence_value(historical)  # Validate size; preserve the already escaped JSON record exactly.
            ws['A1'].comment = Comment(historical, 'RCAIP')
        ws.row_dimensions[2].height = 24
        return ws

    review = base('Case review')
    review.column_dimensions['A'].width = 34
    review.column_dimensions['B'].width = 110
    # Presentation-only merges span otherwise empty cells and avoid clipped scope labels.
    for r in (2, 3, 4):
        review.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    put(review, 4, 1, 'Cell 24 diagnostic cases; producer decisions retained. This is not full Asset delivery.')
    put(review, 5, 1, 'Existing cases'); put(review, 5, 2, len(checked['cases']['rows']))
    put(review, 6, 1, 'Supporting member records'); put(review, 6, 2, len(checked['members']['rows']))
    row = 8
    card_rows, case_rows, member_rows = {}, {}, {}
    for index, case in enumerate(checked['cases']['rows'], 8):
        case_rows[case['case_instance_id']] = index
    for index, member in enumerate(checked['members']['rows'], 8):
        member_rows[(member['case_instance_id'], member['member_id'], member['asset_id'])] = index
    for case in sorted(checked['cases']['rows'], key=lambda c: c['case_instance_id']):
        cid = case['case_instance_id']; card_rows[cid] = row
        put(review, row, 1, 'Case ID'); put(review, row, 2, cid)
        for cell in review[row][:2]:
            cell.fill = dark; cell.font = white
        row += 1
        for title, field in CARD_FIELDS:
            put(review, row, 1, title); cell = put(review, row, 2, case[field])
            height = max(30, (sum(max(1, math.ceil(len(line) / 95)) for line in str(cell.value).split('\n')) * 16 + 10))
            require(height <= 409, f'Case narrative too long for readable Excel row: {field}; no truncation')
            review.row_dimensions[row].height = height
            row += 1
        put(review, row, 1, 'Affected members / distinct assets')
        put(review, row, 2, f"{case['affected_member_count']} / {case['affected_asset_count']}"); row += 1
        for member in checked['members']['rows']:
            if member['case_instance_id'] != cid:
                continue
            put(review, row, 1, 'Supporting member')
            cell = put(review, row, 2, f"{member['member_id']}\nAsset: {member['asset_id']}\n{member['membership_reason']}\nEvidence: {member['source_provenance']}")
            link(cell, 'Members', member_rows[(cid, member['member_id'], member['asset_id'])])
            review.row_dimensions[row].height = 110; row += 1
        put(review, row, 1, 'Full case evidence')
        link(put(review, row, 2, 'Open unchanged producer case fields'), 'Case evidence', case_rows[cid])
        row += 2
    require(row <= 1048576, 'Case presentation exceeds Excel row limit')
    for kind, name in [('cases', 'Case evidence'), ('members', 'Members'), ('closure', 'Branch closure')]:
        p = checked[kind]; ws = base(name); cols = p['columns']
        for r in (2, 3, 4, 5):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=min(7, len(cols)))
        put(ws, 4, 1, PRODUCTS[kind])
        note = ('Original closure stages and units; do not sum mixed grains. Scope comes from the certified producer session.' if kind == 'closure' else
                'Literal evidence. Null=blank; empty/array/temporal text=JSON; fractional numerics=text; controls/backslashes escaped. Click case ID to return to its explanation.')
        if historical and kind == 'closure':
            note = 'Original closure stages and units; do not sum mixed grains. Scope bound to same-cell saved certification, not live-session authority.'
        if data.get('timestamp_session_timezone'):
            note += f" Spark timestamp cast timezone: {data['timestamp_session_timezone']} (timestamp_ntz remains zone-free)."
        put(ws, 5, 1, note); ws.row_dimensions[5].height = 36
        rows = sorted(p['rows'], key=lambda s: s['stage_order']) if kind == 'closure' else p['rows']
        for j, col in enumerate(cols, 1):
            cell = put(ws, 7, j, col); cell.fill = dark; cell.font = white
            ws.column_dimensions[get_column_letter(j)].width = 28
        ws.row_dimensions[7].height = 46
        for r, record in enumerate(rows, 8):
            height = 30
            for j, col in enumerate(cols, 1):
                cell = put(ws, r, j, record[col])
                height = max(height, min(400, math.ceil(len(str(cell.value or '')) / 26) * 15 + 10))
                if col == 'case_instance_id':
                    link(cell, 'Case review', card_rows[record[col]])
            ws.row_dimensions[r].height = height
        if rows:
            table = Table(displayName=f'Cell24_{kind}', ref=f'A7:{get_column_letter(len(cols))}{7 + len(rows)}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
        ws.freeze_panes = 'B8'
    return wb


def workbook_from_cell24_session(producer_namespace, *, synthetic=False, max_rows=100000):
    """Explicit approved-environment entry point; caller saves the returned workbook."""
    return build_cell24_workbook(collect_cell24_session(producer_namespace, max_rows=max_rows), synthetic=synthetic)
