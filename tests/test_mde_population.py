"""Synthetic reconstruction -> retention/reopening -> workbook, no enterprise I/O."""
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import hashlib
import json
import subprocess
import sys
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'notebooks'),str(ROOT/'tests')]
import test_mde_durable_bridge as fixtures
import test_mde_retained_check as native
import rcaip_mde_reconstruction as core
import rcaip_mde_population as population
import RCAIP_MDE_Population_Review as entry

bridge=core.bridge


def setup(shared=True, limitation=False):
    # Reuse accepted positive examples, not invented analytical conclusions.
    if limitation:
        book=native.limitation_fixture()
        with patch.object(native,'fixture',return_value=book): args=fixtures.setup(False)
    else: args=fixtures.setup(shared)
    manifest,_,_,inputs,raw=args
    # The old fixture has four-row display percentages. Reconcile synthetic
    # summary percentages when the existing five-row positive control is used.
    if limitation:
        doc=bridge.strict_json(raw)
        payload=native.payload(doc,0)
        cols=[f['name'] for f in payload['schema']]
        for row in payload['data']:row[cols.index('share_of_mde_missing_pct')]=Decimal('20.00')
        raw=bridge.encoded(doc)
    run=dict(run_id=native.RUN,post_profile_key='SYNTHETIC_PROFILE',profile_hash='SYNTHETIC_PROFILE_HASH',results_namespace='synthetic',active=False,run_status='CLOSED')
    profile=dict(post_slug=native.POST,post_profile_key='SYNTHETIC_PROFILE',profile_hash='SYNTHETIC_PROFILE_HASH',active=False)
    evidence=dict(collected_at_utc='2026-09-27T00:00:00+00:00',products={},
        run_registry=dict(origin=core.metadata.RUN_REGISTRY,metadata_snapshot_version=2,record=run),
        profile_registry=dict(origin=core.metadata.PROFILE_REGISTRY,metadata_snapshot_version=3,record=profile))
    for name,item in inputs.items():
        table='synthetic.'+native.POST+'_'+core.metadata.STEMS[name]
        item['table']=table
        schema=json.dumps(dict(type='struct',fields=item['payload']['schema']))
        evidence['products'][name]=dict(origin=table,detail=dict(format='delta',id=item['table_id']),detail_after=dict(id=item['table_id']),
            requested_version=None if name==bridge.CONTROL else 7,history=[dict(version=7)],schema_json=schema)
    return SimpleNamespace(_enterprise_evidence=evidence), inputs, raw


class Frame:
    def __init__(self,schema):self.schema=SimpleNamespace(json=lambda:json.dumps(dict(type='struct',fields=schema)))
    def collect(self):raise AssertionError('No collect in preparation')


class Spark:
    def __init__(self,inputs):self.inputs=inputs;self.actions=[];self.read=self
    def format(self,value):assert value=='delta';return self
    def option(self,name,value):assert name=='versionAsOf';self.version=value;return self
    def table(self,name):
        item=next(i for i in self.inputs.values() if i['table']==name)
        if item['version']!=self.version:raise RuntimeError('Synthetic missing version')
        self.actions.append(('schema',name,self.version));return Frame(item['payload']['schema'])
    def sql(self,query):
        assert query.startswith('DESCRIBE DETAIL '),query
        name=query.removeprefix('DESCRIBE DETAIL ').replace('`','')
        item=next(i for i in self.inputs.values() if i['table']==name)
        self.actions.append(('detail',name));return SimpleNamespace(first=lambda:dict(format='delta',id=item['table_id']))


def selection(meta,inputs,raw):
    return core.prepare_selection(meta,mode=core.MODE,post=native.POST,run=native.RUN,
        versions={name:7 for name in bridge.INPUTS},retained_sha256=bridge.digest(raw),selected_by='SYNTHETIC_OPERATOR',spark=Spark(inputs))


def result(shared=True,limitation=False):
    meta,inputs,raw=setup(shared,limitation)
    s=selection(meta,inputs,raw)
    # Deterministic synthetic retention; this is clearly a synthetic record.
    s['selection_provenance']['selected_at_utc']='2026-09-27T01:00:00+00:00'
    return core.reconstruct(s,inputs,raw),raw


def model_review_result(mixed):
    """Source-shaped original remainder; no projected treatment/customer fields.

    A null retained resolution signal follows Cell 14's original otherwise
    branch. Expected original products/counts are supplied explicitly here;
    the production reconstruction and all reconciliation checks still run.
    """
    meta,inputs,raw=setup(False);doc=bridge.strict_json(raw)
    p=inputs[bridge.ENRICHMENT]['payload'];cols=[f['name'] for f in p['schema']]
    target=2 if mixed else 3
    if mixed:
        key=p['data'][target][cols.index('asset_key')]
        serial=p['data'][target][cols.index('canonical_serial')]
        p['data'][target]=deepcopy(p['data'][3])
        p['data'][target][cols.index('asset_key')]=key
        p['data'][target][cols.index('canonical_serial')]=serial
    p['data'][target][cols.index('final_resolution_signal')]=None
    r=dict(zip(cols,p['data'][target]))
    # Preserve a conflicting-looking upstream customer context to prove it is
    # not silently substituted for the absent model-review case recipient.
    p['data'][target][cols.index('case_primary_customer')]='POST_DT'
    source_member=native.records(doc,2,2)[0]
    review_cols=bridge.check.contracts(native.sources())['mde_case_model_review_df']
    review={c:source_member.get(c) for c in review_cols}
    review.update({c:r[c] for c in review_cols if c in r})
    review.update(case_rule_key='MDE_CASE_MODEL_REVIEW_REQUIRED',ca_mde_applicability_state='APPLICABLE')
    doc['cells'][2]['outputs'].append(native.table([review],review_cols))
    if mixed:
        qs=native.records(doc,0)
        qs=[row for row in qs if row['mde_qualification_cohort']!='CURRENT_IDENTITY_TEMPORAL_REVIEW']
        qs[-1].update(assets=2,share_of_mde_missing_pct=Decimal('50.00'),resolved_mde_entities=2,stale_mde_entities=2)
        native.put(doc,0,1,qs)
        gs=native.records(doc,1)[1:];gs[0].update(assets=2,share_of_current_unresolved_pct=Decimal('100.00'));native.put(doc,1,1,gs)
        for label,value in [('Qualification cohorts observed',3),('Current identity/evidence review required',0),('Potential current control-review candidates',2)]:native.cert(doc,0,label,value)
        for label,value in [('MDE identity resolution still required',0),('Current applicable enterprise MDE control-review candidates',2)]:native.cert(doc,1,label,value)
    else:
        for output in (1,2):native.payload(doc,2,output)['data']=[]
    counts=[4,1,1,0 if mixed else 1,2 if mixed else 1,1 if mixed else 0,1 if mixed else 0,0,0]
    acc=native.records(doc,2,3)
    for row,n in zip(acc,counts):row['count']=n
    native.put(doc,2,3,acc)
    updates={'Current identity/temporal review outside control-case set':counts[3],
             'Owner-qualified enterprise residual candidates':counts[4],
             'Residual case members formed':counts[5],'Deterministic enterprise assurance cases':counts[6],
             'Enterprise service-owner cases':counts[6],'Residual members requiring model review':1,
             'Case formation readiness':'PARTIAL_MODEL_REVIEW_REQUIRED',
             'Residual member reconciliation':str(counts[4])+' / '+str(counts[4])}
    for label,value in updates.items():native.cert(doc,2,label,value)
    raw=bridge.encoded(doc)
    return core.reconstruct(selection(meta,inputs,raw),inputs,raw),raw,review


class PopulationTests(unittest.TestCase):
    def test_nonempty_model_review_only_and_mixed_end_to_end(self):
        from openpyxl import load_workbook
        for mixed in (False,True):
            with self.subTest(mixed=mixed):
                b,raw,review=model_review_result(mixed)
                self.assertNotIn('recommended_treatment',review)
                self.assertNotIn('primary_customer',review)
                self.assertNotIn('case_id',review)
                with fixtures.synthetic_directory() as d:
                    dest=d/'model-review.json'
                    receipt=core.retain_bundle(b,dest,approved_directory=d,original_raw=raw)
                    reopened=core.load_bundle(dest,expected_sha256=receipt['sha256'],original_raw=raw)
                    book=d/'model-review.xlsx'
                    population.save_workbook(reopened,book,approved_directory=d,original_raw=raw,synthetic=True)
                    wb=load_workbook(book)
                    pop=wb['MDE Population'];headers=[c.value for c in pop[8]]
                    keys=[pop.cell(i,1).value for i in range(9,pop.max_row+1)]
                    self.assertEqual(set(keys),{r['asset_key'] for r in b['qualification']})
                    self.assertEqual(len(keys),len(set(keys)))
                    row=9+keys.index(review['asset_key'])
                    self.assertIn('Original case model left this condition unresolved',pop.cell(row,headers.index('Supported next question')+1).value)
                    self.assertIsNone(pop.cell(row,headers.index('Original recipient')+1).value)
                    self.assertIsNone(pop.cell(row,headers.index('Original case')+1).value)
                    self.assertEqual(pop.cell(row,headers.index('qualification.case_primary_customer')+1).value,'POST_DT')
                    self.assertEqual(pop.cell(row,1).hyperlink.location,"'Model review'!A9")
                    ws=wb['Model review'];cols=[c.value for c in ws[8] if c.value is not None]
                    self.assertEqual(cols,list(review))
                    self.assertEqual([ws.cell(9,j+1).value for j in range(len(cols))],[population.evidence_value(review[c]) for c in cols])
                    self.assertEqual(ws.cell(9,cols.index('asset_key')+1).hyperlink.location,"'MDE Population'!A"+str(row))
                    self.assertEqual(len(b['original_products'][bridge.check.PRODUCTS[14][0]]),int(mixed))
                    self.assertEqual(len(b['original_products'][bridge.check.PRODUCTS[14][1]]),int(mixed))
                    if mixed:
                        case=b['original_products'][bridge.check.PRODUCTS[14][0]][0]
                        member=b['original_products'][bridge.check.PRODUCTS[14][1]][0]
                        formed=9+keys.index(member['asset_key'])
                        self.assertEqual(pop.cell(formed,7).value,case['primary_customer'])
                        self.assertEqual(pop.cell(formed,8).value,case['case_id'])
                        self.assertEqual(pop.cell(formed,8).hyperlink.location,"'Case review'!A9")

    def test_strict_route_unchanged_bytes_and_missing_witness_rejection(self):
        import types
        baseline=subprocess.check_output(['git','show','HEAD:notebooks/rcaip_mde_durable_bridge.py'],cwd=ROOT)
        old=types.ModuleType('baseline_bridge');exec(compile(baseline,'baseline_bridge','exec'),old.__dict__)
        for shared in (False,True):
            args=fixtures.setup(shared)
            self.assertEqual(bridge.encoded(old.reproduce(*args)),bridge.encoded(bridge.reproduce(*args)))
            args[1].clear()
            for module in (old,bridge):
                with self.assertRaisesRegex(bridge.InputContractError,'unavailable'):module.reproduce(*args)

    def test_explicit_route_separation_and_no_witnesses(self):
        b,raw=result()
        self.assertEqual(b['format'],core.FORMAT);self.assertEqual(b['limitation'],core.LIMITATION)
        self.assertNotIn('trusted_references',b);self.assertNotIn('evidence',b)
        with self.assertRaisesRegex(bridge.InputContractError,'Derived bundle contract'):bridge.validate_bundle(b,{},raw)
        with self.assertRaises(bridge.InputContractError):bridge.reproduce(b['manifest'],{}, {},b['inputs'],raw)
        b['selection']['mode']='AUTO'
        with self.assertRaisesRegex(bridge.InputContractError,'explicitly'):core.reconstruct(b['selection'],b['inputs'],raw)

    def test_existing_metadata_preparation_and_unverified_control_schema(self):
        meta,inputs,raw=setup();spark=Spark(inputs)
        s=core.prepare_selection(meta,mode=core.MODE,post=native.POST,run=native.RUN,versions={n:7 for n in bridge.INPUTS},
                                 retained_sha256=bridge.digest(raw),selected_by='SYNTHETIC_OPERATOR',spark=spark)
        self.assertEqual(sum(a[0]=='schema' for a in spark.actions),3)
        self.assertEqual(s['inputs'][bridge.CONTROL]['version'],7)
        self.assertEqual(meta._enterprise_evidence['products'][bridge.CONTROL]['requested_version'],None)
        inputs[bridge.CONTROL]['version']=8
        with self.assertRaisesRegex(RuntimeError,'missing version'):selection(meta,inputs,raw)

    def test_scope_identity_schema_version_and_explicit_requests(self):
        changes=[lambda s:s.update(run='foreign'),lambda s:s.update(post='foreign'),
                 lambda s:s['inputs'][bridge.SEARCH].update(table='synthetic.foreign'),
                 lambda s:s['inputs'][bridge.ENRICHMENT].update(table_id='other'),
                 lambda s:s['inputs'][bridge.CONTROL].update(version=8),
                 lambda s:s['inputs'][bridge.CONTROL].update(version='latest'),
                 lambda s:s['inputs'][bridge.SEARCH].update(schema_sha256='0'*64),
                 lambda s:s['selection_provenance'].update(meaning='attested')]
        for change in changes:
            meta,inputs,raw=setup();s=selection(meta,inputs,raw);change(s)
            with self.subTest(change=change),self.assertRaises(bridge.InputContractError):core.reconstruct(s,inputs,raw)
        meta,inputs,raw=setup();meta._enterprise_evidence['products'][bridge.SEARCH]['detail']['id']='replaced'
        with self.assertRaises(bridge.InputContractError):selection(meta,inputs,raw)

    def test_preparation_invalid_binding_stops_before_reads_and_schema_drift_stops(self):
        meta,inputs,raw=setup();spark=Spark(inputs)
        meta._enterprise_evidence['run_registry']['record']['results_namespace']='unsafe`name'
        with self.assertRaisesRegex(bridge.InputContractError,'namespace'):
            core.prepare_selection(meta,mode=core.MODE,post=native.POST,run=native.RUN,versions={n:7 for n in bridge.INPUTS},
                retained_sha256=bridge.digest(raw),selected_by='SYNTHETIC_OPERATOR',spark=spark)
        self.assertEqual(spark.actions,[])
        meta,inputs,raw=setup();inputs[bridge.ENRICHMENT]['payload']['schema'][0]['type']='double'
        with self.assertRaisesRegex(bridge.InputContractError,'schema changed'):selection(meta,inputs,raw)

    def test_exact_membership_context_and_retained_contradictions(self):
        for mode in ('missing','duplicate','foreign','scope','null_context','original_case','source'):
            with self.subTest(mode=mode):
                meta,inputs,raw=setup();s=selection(meta,inputs,raw)
                p=inputs[bridge.ENRICHMENT]['payload'];cols=[f['name'] for f in p['schema']]
                if mode=='missing':p['data'].pop()
                if mode=='duplicate':p['data'].append(deepcopy(p['data'][0]))
                if mode=='foreign':p['data'][0][cols.index('asset_key')]='OTHER'
                if mode=='scope':p['data'][0][cols.index('rcaip_run_id')]='OTHER'
                if mode=='null_context':
                    p=inputs[bridge.CONTROL]['payload'];cols=[f['name'] for f in p['schema']]
                    p['data'][0][cols.index('applicability_state')]=None
                if mode in ('original_case','source'):
                    doc=bridge.strict_json(raw)
                    if mode=='source':doc['cells'][0]['source']+='\nchanged = True\n'
                    else:native.set_value(doc,2,0,'recommended_treatment','invented',output=1)
                    raw=bridge.encoded(doc);s['retained_sha256']=bridge.digest(raw)
                with self.assertRaises(bridge.InputContractError):core.reconstruct(s,inputs,raw)

    def test_external_hash_semantics_and_nonoverwriting_retention(self):
        b,raw=result()
        with fixtures.synthetic_directory() as d:
            p=d/'reconstruction.json';receipt=core.retain_bundle(b,p,approved_directory=d,original_raw=raw)
            self.assertEqual(receipt['limitation'],core.LIMITATION)
            self.assertEqual(core.load_bundle(p,expected_sha256=receipt['sha256'],original_raw=raw),b)
            with self.assertRaisesRegex(bridge.InputContractError,'destination'):core.retain_bundle(b,p,approved_directory=d,original_raw=raw)
            with self.assertRaisesRegex(bridge.InputContractError,'hash/size'):core.load_bundle(p,expected_sha256='0'*64,original_raw=raw)
            bad=deepcopy(b);bad['qualification'][0]['mde_next_gate']='FORGED';p.write_bytes(bridge.encoded(bad))
            with self.assertRaisesRegex(bridge.InputContractError,'semantic'):core.load_bundle(p,expected_sha256=bridge.digest(p.read_bytes()),original_raw=raw)

    def test_complete_workbook_native_values_tables_navigation(self):
        from openpyxl import load_workbook
        b,raw=result()
        with fixtures.synthetic_directory() as d:
            p=d/'bundle.json';receipt=core.retain_bundle(b,p,approved_directory=d,original_raw=raw)
            loaded=core.load_bundle(p,expected_sha256=receipt['sha256'],original_raw=raw)
            dest=d/'workbook.xlsx';out=population.save_workbook(loaded,dest,approved_directory=d,original_raw=raw,synthetic=True)
            self.assertEqual(out['limitation'],core.LIMITATION)
            wb=load_workbook(dest);self.assertEqual(tuple(wb.sheetnames),population.SHEETS)
            self.assertEqual(wb['MDE Population'].max_row,8+len(b['qualification']))
            self.assertIn('UNESTABLISHED',str([c.value for row in wb['RCSO Summary'] for c in row]))
            for sheet,product in zip(('Case evidence','Members','Shared support','Model review'),bridge.check.PRODUCTS[14]):
                expected=b['original_products'][product]
                if not expected:continue
                cols=[c.value for c in wb[sheet][8] if c.value is not None]
                actual=[[wb[sheet].cell(i+9,j+1).value for j in range(len(cols))] for i in range(len(expected))]
                wanted=[[population.evidence_value(r[c]) for c in cols] for r in expected]
                self.assertEqual(actual,wanted)
            links=[c.hyperlink for ws in wb for row in ws for c in row if c.hyperlink]
            self.assertGreater(len(links),15)
            self.assertTrue(all(c.data_type not in ('f','e') for ws in wb for row in ws for c in row))

    def test_exact_numeric_temporal_literal_and_schema_retention(self):
        from openpyxl import load_workbook
        meta,inputs,raw=setup()
        p=inputs[bridge.ENRICHMENT]['payload']
        values=[Decimal('-0.0'),Decimal('0.12345678901234566'),9007199254740993,Decimal('123.4500'),
                '0000-01-01T00:00:00','+12000-01-01T00:00:00','2025-10-26T00:30:00Z','2025-10-26T01:30:00Z',
                '=SUM(A1:A2)','#N/A','literal\\u0001\x01','',None]
        types=['double','double','long','decimal(20,4)']+['string']*9
        for i,(v,t) in enumerate(zip(values,types)):
            p['schema'].append(dict(name='synthetic_evidence_'+str(i),type=t))
            for row in p['data']:row.append(v)
        inputs_meta=meta._enterprise_evidence['products'][bridge.ENRICHMENT]
        inputs_meta['schema_json']=json.dumps(dict(type='struct',fields=p['schema']))
        s=selection(meta,inputs,raw);b=core.reconstruct(s,inputs,raw)
        with fixtures.synthetic_directory() as d:
            dest=d/'evidence.xlsx';population.save_workbook(b,dest,approved_directory=d,original_raw=raw,synthetic=True)
            wb=load_workbook(dest);cols=[c.value for c in wb['MDE Population'][8]]
            for i,value in enumerate(values):
                col=cols.index('qualification.synthetic_evidence_'+str(i))+1
                self.assertEqual(wb['MDE Population'].cell(9,col).value,population.evidence_value(value))

    def test_varied_population_and_source_empty_guards(self):
        for shared,limitation in ((False,False),(True,False),(False,True)):
            b,raw=result(shared,limitation)
            wb=population.build_workbook(b,original_raw=raw,synthetic=True)
            self.assertEqual(wb['MDE Population'].max_row,8+len(b['qualification']))
        meta,inputs,raw=setup();s=selection(meta,inputs,raw)
        for item in inputs.values():item['payload']['data']=[]
        with self.assertRaisesRegex(bridge.InputContractError,'Cell 10 requires nonempty'):core.reconstruct(s,inputs,raw)

    def test_zero_original_cases_retained_without_new_case_formation(self):
        meta,inputs,raw=setup(False);doc=bridge.strict_json(raw)
        p=inputs[bridge.ENRICHMENT]['payload'];cols=[f['name'] for f in p['schema']]
        key=p['data'][3][cols.index('asset_key')];serial=p['data'][3][cols.index('canonical_serial')]
        p['data'][3]=deepcopy(p['data'][2]);p['data'][3][cols.index('asset_key')]=key;p['data'][3][cols.index('canonical_serial')]=serial
        qsummary=native.records(doc,0)
        qsummary=[r for r in qsummary if r['mde_qualification_cohort']!='CURRENT_STALE_OR_HISTORICAL_MDE_ONLY']
        qsummary[-1].update(assets=2,share_of_mde_missing_pct=Decimal('50.00'),resolved_mde_entities=2,historical_mde_entities=2)
        native.put(doc,0,1,qsummary)
        gsummary=native.records(doc,1)[:1];gsummary[0].update(assets=2,share_of_current_unresolved_pct=Decimal('100.00'));native.put(doc,1,1,gsummary)
        for label,value in [('Qualification cohorts observed',3),('Current identity/evidence review required',2),('Potential current control-review candidates',0)]:native.cert(doc,0,label,value)
        for label,value in [('MDE identity resolution still required',2),('Current applicable enterprise MDE control-review candidates',0)]:native.cert(doc,1,label,value)
        for output in (1,2):native.payload(doc,2,output)['data']=[]
        acc=native.records(doc,2,3)
        for r,n in zip(acc,[4,1,1,2,0,0,0,0,0]):r['count']=n
        native.put(doc,2,3,acc)
        for label,value in [('Current identity/temporal review outside control-case set',2),('Owner-qualified enterprise residual candidates',0),
                            ('Residual case members formed',0),('Deterministic enterprise assurance cases',0),('Enterprise service-owner cases',0),('Residual member reconciliation','0 / 0')]:native.cert(doc,2,label,value)
        raw=bridge.encoded(doc);s=selection(meta,inputs,raw);b=core.reconstruct(s,inputs,raw)
        self.assertEqual(b['original_products'][bridge.check.PRODUCTS[14][0]],[])
        with fixtures.synthetic_directory() as d:
            population.save_workbook(b,d/'zero-cases.xlsx',approved_directory=d,original_raw=raw,synthetic=True)

    def test_reordered_members_keep_navigation_and_complete_population(self):
        meta,inputs,raw=setup()
        for item in inputs.values():item['payload']['data'].reverse()
        b=core.reconstruct(selection(meta,inputs,raw),inputs,raw)
        with fixtures.synthetic_directory() as d:
            population.save_workbook(b,d/'permuted.xlsx',approved_directory=d,original_raw=raw,synthetic=True)

    def test_public_loading_operations_are_explicit_and_safe(self):
        modules=entry._modules(ROOT/'notebooks');self.assertIn('rcaip_mde_population',modules)
        meta,inputs,raw=setup();s=entry.prepare(meta,module_folder=ROOT/'notebooks',spark=Spark(inputs),mode=core.MODE,
            post=native.POST,run=native.RUN,versions={n:7 for n in bridge.INPUTS},retained_sha256=bridge.digest(raw),selected_by='SYNTHETIC_OPERATOR')
        self.assertEqual(s['mode'],core.MODE)
        for exc,kind in [(RuntimeError('PRIVATE'),entry.PopulationOperationError),(MemoryError('PRIVATE'),MemoryError),(KeyboardInterrupt('PRIVATE'),KeyboardInterrupt)]:
            with self.assertRaises(kind) as ctx:entry._safe('SYNTHETIC',lambda:(_ for _ in ()).throw(exc))
            self.assertNotIn('PRIVATE',str(ctx.exception))
        # Importing the public module invokes none of these enterprise operations.
        import ast
        tree=ast.parse((ROOT/'notebooks/RCAIP_MDE_Population_Review.py').read_text())
        self.assertFalse(any(isinstance(n,ast.Expr) and isinstance(n.value,ast.Call) for n in tree.body))


def generate(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    bundle,raw=result()
    path=directory/'RCAIP_MDE_SYNTHETIC_RECONSTRUCTION.json'
    receipt=core.retain_bundle(bundle,path,approved_directory=directory,original_raw=raw)
    loaded=core.load_bundle(path,expected_sha256=receipt['sha256'],original_raw=raw)
    workbook=directory/'RCAIP_MDE_Population_SYNTHETIC_REVIEW.xlsx'
    published=population.save_workbook(loaded,workbook,approved_directory=directory,original_raw=raw,synthetic=True)
    original=directory/'SYNTHETIC_MDE_ORIGINAL.ipynb'
    with original.open('xb') as stream:stream.write(raw)
    print(json.dumps(dict(bundle=receipt,workbook=published),indent=2))


if __name__=='__main__':
    if len(sys.argv)==3 and sys.argv[1]=='--generate':generate(sys.argv[2])
    else:unittest.main()
