"""Local synthetic acceptance; dataframe substitute, not PySpark or full report."""
import ast
from collections import Counter, OrderedDict
from datetime import datetime, timezone
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import unittest
import tempfile
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from ip001_synthetic_frame import D, F, extract

ROOT=Path(__file__).resolve().parents[1]
P=ROOT/'notebooks/RCAIP_IP_001_Open_Port_Intelligence v2.0.py'
R=ROOT/'notebooks/RCAIP_REPORT_001_DT_Validation_Package_V2.py'
ns={'F':F}
extract(P,{'IP001_HOST_GRAIN','printer_signal_pattern','ip001_same_rows','ip001_validate_rules',
    'ip001_validate_observations','ip001_enrich_ports','ip001_mark_action_members','ip001_attach_members','ip001_printer_ports'},ns)
kb=ast.parse((ROOT/'notebooks/RCAIP_KB_001_Knowledge_Base.py').read_text())
rules=next(n.value for n in kb.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='port_rules' for t in n.targets))
RULES=D([dict(zip(('port','port_label','review_priority','risk_theme','standard_recommendation','csa_relevance'),x)) for x in ast.literal_eval(rules)])
class StringType: pass
class ArrayType: pass
class OtherType: pass
T=type('Types',(),dict(BinaryType=OtherType,TimestampType=OtherType,DateType=OtherType))
D.schema=property(lambda d:type('Schema',(),{'fields':[type('Field',(),{'name':c,'dataType':ArrayType() if c=='port_members' else StringType()}) for c in d.columns]})())
rns={'F':F,'T':T,'COMPLEX_TYPES':(ArrayType,), 'pd':pd,'Workbook':Workbook,'get_column_letter':get_column_letter,
     'Alignment':Alignment,'Border':Border,'Font':Font,'PatternFill':PatternFill,'Side':Side,'datetime':datetime,'timezone':timezone}
tree=extract(R,{'PORT_SOURCE_KEYS','PORT_SHEET_KEYS','PORT_EVIDENCE_LIMIT','PORT_EVIDENCE_WARNING',
 'resolve_port_table','load_port_table','port_source_or_status','port_cell','port_member_counts','certify_port_frames',
 'verify_port_sheet','link_port_members','excel_safe_spark_dataframe','status_fill','safe_table_name','sanitize_excel_value'},rns)
for n in tree.body:
 if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and (t.id.startswith('COLOR_') or t.id.endswith(('_FILL','_FONT')) or t.id=='MEDIUM_BLUE') for t in n.targets):
  exec(compile(ast.Module(body=[n],type_ignores=[]),str(R),'exec'),rns)
sp=importlib.util.spec_from_file_location('audit',ROOT/'notebooks/rcaip_ip001_audit.py'); audit=importlib.util.module_from_spec(sp);sys.modules['audit']=audit;sp.loader.exec_module(audit)

def observation(port=23,**kw):
 row=dict(post_name='SYNTHETIC',target_cidr='192.0.2.0/24',post_slug='synthetic',rcaip_run_id='synthetic_run',scan_id='scan_a',host_key='HOST_A',ip_address='192.0.2.1',asset_domain='Synthetic',scan_scope='Synthetic',scan_location='Synthetic',scan_source_file='synthetic.xml',protocol='tcp',port=port,hostname=None,port_state='open',service='telnet',product=None,version=None)
 return dict(row,**kw)
def enrich(rows,rules=RULES):return ns['ip001_enrich_ports'](D(rows),rules)
def group(detail):
 env=dict(ns,open_port_review=detail)
 tr=ast.parse(P.read_text(encoding='utf-8'))
 for name in ('priority_rank_expression','device_port_rollup','device_port_prioritized','post_action_table'):
  assignments=[n for n in tr.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id==name for t in n.targets)]
  for node in assignments:exec(compile(ast.Module(body=[node],type_ignores=[]),str(P),'exec'),env)
 return env['post_action_table']
def summary(detail):
 env=dict(ns,open_port_review=detail)
 for node in ast.parse(P.read_text(encoding='utf-8')).body:
  if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='port_summary' for t in node.targets):
   exec(compile(ast.Module(body=[node],type_ignores=[]),str(P),'exec'),env)
 return env['port_summary']
def safe_frame(d):return rns['excel_safe_spark_dataframe'](d,[]).toPandas()
def artifact(body='<port protocol="tcp" portid="23"><state state="open"/><service name="telnet"/></port>',**kw):
 raw=('<nmaprun start="1000"><scaninfo protocol="tcp" services="21,23,80"/><host '+kw.get('host','')+'><address addr="192.0.2.1" addrtype="ipv4"/><ports>'+body+'</ports></host><runstats><finished time="1001" exit="'+kw.get('exit','success')+'"/></runstats></nmaprun>').encode()
 return dict(bytes=raw,sha256=hashlib.sha256(raw).hexdigest(),scan_id='scan_a',perspective='INTERNAL_OPENNET',time_basis='UNIX_EPOCH_UTC',
             scope_lineage=dict(post_slug='synthetic',rcaip_run_id='synthetic_run',scan_id='scan_a',sha256=hashlib.sha256(raw).hexdigest(),perspective='INTERNAL_OPENNET',reference='SYNTHETIC approved manifest'))
def binding(): return audit.Binding('synthetic.results',0,'synthetic-id',tuple((k,'int' if k=='port' else 'string') for k in observation()),'0'*64,'2026-10-14 10:00:00.100000')
def run_a(items,products=None,**kw):return audit.audit(items,products or {},post='synthetic',run='synthetic_run',**kw)

class Retention(unittest.TestCase):
 def test_unmapped_required_and_known_classification(self):
  d=enrich([observation(),observation(65023)],D([r for r in RULES.rows if r['port']!=23]))
  self.assertEqual(len(d.rows),2);self.assertTrue(d.rows[0]['required_port_review']);self.assertIsNone(d.rows[0]['review_priority'])
  self.assertEqual({r['port_rule_status'] for r in d.rows},{'UNMAPPED_RULE'})
  self.assertEqual(enrich([observation()]).rows[0]['review_priority'],'High')
 def test_exact_and_conflicting_duplicates(self):
  rule=next(r for r in RULES.rows if r['port']==23)
  self.assertEqual(len(enrich([observation()],D(RULES.rows+[rule])).rows),1)
  with self.assertRaisesRegex(RuntimeError,'conflicting'):enrich([observation()],D(RULES.rows+[dict(rule,risk_theme='different')]))
 def test_null_keys(self):
  for k in ns['IP001_HOST_GRAIN']+['protocol','port']:
   with self.subTest(k=k),self.assertRaisesRegex(RuntimeError,'key'):enrich([observation(**{k:None})])
 def test_same_observation_printer_and_null_names(self):
  for name in [None,'','NEUTRAL']:
   rows=[observation(p,hostname=name,service='jetdirect' if p==9100 else 'service') for p in (21,23,80,9100)]
   self.assertEqual(len(ns['ip001_printer_ports'](enrich(rows)).rows),4)
   self.assertEqual(len(ns['ip001_printer_ports'](enrich(rows[:3])).rows),0)
 def test_no_cross_scope_spread(self):
  rows=[observation(9100,service='jetdirect')]
  for k in ['scan_id','post_slug','rcaip_run_id','host_key','scan_source_file','ip_address']:
   rows.append(observation(**{k:'OTHER'}))
  self.assertEqual(len(ns['ip001_printer_ports'](enrich(rows)).rows),1)
 def test_protocol_and_source_repetition(self):
  rows=[observation(),observation(),observation(protocol='udp')]
  d=enrich(rows);g=group(d)
  self.assertEqual(len(g.rows[0]['port_members']),3)
  self.assertEqual(Counter(json.loads(x)['protocol'] for x in g.rows[0]['port_members']),Counter(tcp=2,udp=1))
 def test_equal_totals_wrong_members(self):
  with self.assertRaisesRegex(RuntimeError,'membership'):
   ns['ip001_same_rows'](D([observation(21),observation(23)]),D([observation(21),observation(21)]),list(observation()),'synthetic')
 def test_provenance_multiplication(self):
  d=enrich([observation()])
  with self.assertRaisesRegex(RuntimeError,'membership'):ns['ip001_same_rows'](d,D(d.rows*2),d.columns,'provenance')
 def test_actual_action_selection_required_priority_override(self):
  assignment=next(n for n in ast.parse(P.read_text()).body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='post_action_table' for t in n.targets))
  cols=next(n for n in ast.walk(assignment) if isinstance(n,ast.Call) and isinstance(n.func,ast.Attribute) and n.func.attr=='select').args
  row={n.value:'' for n in cols};row.update(final_priority='Review',required_port_review=1)
  env=dict(F=F,device_port_prioritized=D([row,dict(row,required_port_review=0)]))
  exec(compile(ast.Module(body=[assignment],type_ignores=[]),str(P),'exec'),env)
  self.assertEqual(len(env['post_action_table'].rows),1)
 def test_report_accounting_and_lossless_cell(self):
  d=enrich([observation(),observation()]);g=group(d);d=ns['ip001_mark_action_members'](d,g)
  frames={'Open_Port_Detail':safe_frame(d),'Open_Port_Summary':safe_frame(summary(d)),'Open_Port_Actions':safe_frame(g)}
  rns['certify_port_frames'](frames)
  frames['Open_Port_Actions']=safe_frame(group(enrich([observation()])))
  with self.assertRaisesRegex(RuntimeError,'Required'):rns['certify_port_frames'](frames)
  self.assertEqual(rns['port_cell'](0.12345678901234566),'0.12345678901234566')
  self.assertEqual(rns['port_cell']('a\x01\\b'),'a\\u0001\\\\b')
  self.assertTrue(rns['port_cell']('=1+1').startswith("'"))
  with self.assertRaisesRegex(RuntimeError,'bound'):rns['port_cell']('a'*32768)
 def test_explicit_report_binding(self):
  rns['RCAIP_PORT_REPORT_BINDINGS']={}
  self.assertEqual(rns['resolve_port_table']('open_port_detail',['A'])['port_status'],'UNAVAILABLE')
  rns['RCAIP_PORT_REPORT_BINDINGS']={'open_port_detail':None}
  self.assertEqual(rns['resolve_port_table']('open_port_detail',['A'])['port_status'],'NOT_PRODUCED')
  rns['get_rcaip_table']=lambda *a,**k:'synthetic.results'
  b=dict(logical_name='A',table='synthetic.results',version=0,generation='2026-10-14 10:00:00.100000',table_id='id',execution_id='e',producer_sha256='0'*64,schema=(),write_token='token')
  rns['RCAIP_PORT_REPORT_BINDINGS']={'open_port_detail':b}
  self.assertEqual(rns['resolve_port_table']('open_port_detail',['A'])['port_binding'],b)
  b['version']='latest'
  with self.assertRaisesRegex(RuntimeError,'binding'):rns['resolve_port_table']('open_port_detail',['A'])
 def test_actual_optional_read_error_loading_path(self):
  loop=next(n for n in tree.body if isinstance(n,ast.For) and isinstance(n.target,ast.Tuple) and isinstance(n.target.elts[0],ast.Name) and n.target.elts[0].id=='source_key')
  def unavailable(config):raise TimeoutError('SENSITIVE details')
  env=dict(rns, SOURCE_CONTRACT={'open_port_detail':dict(table_name='synthetic.results',available=True,required=False,shared=False)},
           REPORT_SOURCE_DFS={'open_port_detail':'stale'},source_validation_records=[],load_port_table=unavailable)
  exec(compile(ast.Module(body=[loop],type_ignores=[]),str(R),'exec'),env)
  self.assertNotIn('open_port_detail',env['REPORT_SOURCE_DFS'])
  self.assertEqual(env['SOURCE_CONTRACT']['open_port_detail']['port_status'],'UNAVAILABLE_READ_ERROR')
  self.assertNotIn('SENSITIVE',str(env['source_validation_records']))
 def test_empty_members_and_readiness_are_not_compliance(self):
  d=enrich([observation()]);g=group(d);d=ns['ip001_mark_action_members'](d,g)
  frames={'Open_Port_Detail':safe_frame(d).iloc[:0], 'Open_Port_Summary':safe_frame(summary(d)).iloc[:0], 'Open_Port_Actions':safe_frame(g).iloc[:0]}
  rns['certify_port_frames'](frames)
  self.assertIn('NOT ESTABLISHED',rns['PORT_EVIDENCE_WARNING'])
  self.assertIn('No compliance conclusion',rns['PORT_EVIDENCE_WARNING'])
 def test_report_wrong_scope_or_generation(self):
  class Reader:
   def option(self,k,v):self.version=v;return self
   def table(self,t):return self.frame
  reader=Reader();rns.update(spark=type('Spark',(),{'read':reader,'conf':type('Conf',(),{'get':lambda *a:'UTC'})()})(),RCAIP_POST_SLUG='synthetic',RCAIP_RUN_ID='synthetic_run')
  row=observation(generated_at_utc='g')
  for field in ('post_slug','rcaip_run_id','generated_at_utc'):
   reader.frame=D([dict(row,**{field:'wrong'})])
   with self.assertRaisesRegex(RuntimeError,'mismatch'):rns['load_port_table']({'port_binding':dict(version=0,table='synthetic.results',generation='g')})
  reader.frame=D([row]);self.assertEqual(rns['load_port_table']({'port_binding':dict(version=0,table='synthetic.results',generation='g')}).count(),1)
  reader.frame=D([{k:v for k,v in row.items() if k!='rcaip_run_id'}])
  with self.assertRaisesRegex(RuntimeError,'authority'):rns['load_port_table']({'port_binding':dict(version=0,table='synthetic.results',generation='g')})

class IndependentAudit(unittest.TestCase):
 def test_not_produced_and_read_error(self):
  s=run_a([artifact()],{'general':'READ_ERROR','report_detail':None}).status()
  self.assertEqual(s['general'],'UNAVAILABLE_READ_ERROR');self.assertEqual(s['report_detail'],'NOT_PRODUCED')
  self.assertEqual(s['q1143_scope'],'UNESTABLISHED');self.assertEqual(s['overall_compliance'],'NOT_ASSESSED')
  self.assertNotIn('192.0.2.1',repr(s));self.assertNotIn('192.0.2.1',repr(run_a([artifact()])))
 def test_membership_and_scope(self):
  row=observation(generated_at_utc='2026-10-14 10:00:00.100000')
  self.assertEqual(run_a([artifact()],{'general':dict(binding=binding(),rows=[row])}).status()['general'],'AVAILABLE_POSITIVE')
  self.assertEqual(run_a([artifact()],{'general':dict(binding=binding(),rows=[dict(row,port=21)])}).status()['general'],'MEMBERSHIP_MISMATCH')
  with self.assertRaisesRegex(audit.AuditError,'SCOPE'):run_a([artifact()],{'general':dict(binding=binding(),rows=[dict(row,post_slug='other')])})
 def test_states_and_partial_positive(self):
  sample=[dict(scan_id='scan_a',ip_address='192.0.2.1',approved=True)]
  for state in ('open','closed','filtered','unfiltered','open|filtered','closed|filtered'):
   a=artifact('<port protocol="tcp" portid="23"><state state="'+state+'"/></port>',host='timedout="true"')
   s=run_a([a],sample=sample).status()
   self.assertEqual(s['observed_state_counts'][state],1);self.assertEqual(s['coverage_counts']['INCOMPLETE'],3)
   self.assertEqual(s['positive_observations'],int(state=='open'))
 def test_discovery_unscanned_extraports(self):
  sample=[dict(scan_id='scan_a',ip_address='192.0.2.1',approved=True)]
  for xml,expected in [('', 'UNESTABLISHED'),('<extraports state="closed" count="3"/>','EXTRAPORTS_LIMITATION')]:
   self.assertEqual(run_a([artifact(xml)],sample=sample).status()['coverage_counts'][expected],3)
  self.assertEqual(run_a([artifact('',exit='error')],sample=sample).status()['coverage_counts']['INCOMPLETE'],3)
 def test_bounds_hash_and_version(self):
  with self.assertRaisesRegex(audit.AuditError,'BYTE_BOUND'):run_a([artifact()],max_bytes=1)
  a=artifact();a['sha256']='0'*64
  with self.assertRaisesRegex(audit.AuditError,'HASH'):run_a([a])
  b=binding();object.__setattr__(b,'version',None)
  with self.assertRaisesRegex(audit.AuditError,'VERSION'):b.validate()
  with self.assertRaisesRegex(audit.AuditError,'DUPLICATE'):run_a([artifact(),artifact()])

 def test_pinned_read_adapter(self):
  import types
  class Row(dict):
   def asDict(self,recursive=False):return dict(self)
  class ReadFrame:
   def __init__(self,rows):self.rows=rows;self.schema=types.SimpleNamespace(fields=[types.SimpleNamespace(name=k,dataType=types.SimpleNamespace(simpleString=lambda k=k:'int' if k=='port' else 'string')) for k in observation()])
   def limit(self,n):self.bound=n;return ReadFrame(self.rows[:n])
   def collect(self):return [Row(x) for x in self.rows]
  class Reader:
   def option(self,k,v):self.request=(k,v);return self
   def table(self,t):self.table_name=t;return self.frame
  reader=Reader();reader.frame=ReadFrame([observation(generated_at_utc='2026-10-14 10:00:00.100000')])
  identity=types.SimpleNamespace(select=lambda *a:types.SimpleNamespace(collect=lambda:[{'id':'synthetic-id'}]))
  spark=types.SimpleNamespace(read=reader,sql=lambda query:identity)
  old={k:sys.modules.get(k) for k in ('pyspark','pyspark.sql')}
  sys.modules['pyspark']=types.ModuleType('pyspark');sys.modules['pyspark.sql']=types.SimpleNamespace(functions=F)
  try:
   rows=audit.read_delta_rows(spark,binding(),'synthetic','synthetic_run',1)
   self.assertEqual(len(rows),1);self.assertEqual(reader.request,('versionAsOf',0))
   reader.frame=ReadFrame(reader.frame.rows*2)
   with self.assertRaisesRegex(audit.AuditError,'BOUND'):audit.read_delta_rows(spark,binding(),'synthetic','synthetic_run',1)
   spark.sql=lambda q:types.SimpleNamespace(select=lambda *a:types.SimpleNamespace(collect=lambda:[{'id':'changed'}]))
   with self.assertRaisesRegex(audit.AuditError,'IDENTITY'):audit.read_delta_rows(spark,binding(),'synthetic','synthetic_run')
   def timeout(q):raise TimeoutError('sensitive operational content')
   spark.sql=timeout
   with self.assertRaisesRegex(audit.AuditError,'^INPUT_READ_UNAVAILABLE$'):audit.read_delta_rows(spark,binding(),'synthetic','synthetic_run')
  finally:
   for k,v in old.items():
    if v is None:sys.modules.pop(k,None)
    else:sys.modules[k]=v
 def test_independent_printer_minimum_and_unknown_state(self):
  a=artifact('<port protocol="tcp" portid="23"><state state="open"/><service name="telnet"/></port><port protocol="tcp" portid="9100"><state state="open"/><service name="jetdirect"/></port>')
  row=observation(9100,service='jetdirect',generated_at_utc='2026-10-14 10:00:00.100000')
  self.assertEqual(run_a([a],{'q1143':dict(binding=binding(),rows=[row])}).status()['q1143_missing'],1)
  self.assertEqual(run_a([artifact()]).status()['report_detail'],'UNAVAILABLE_BINDING')
  with self.assertRaisesRegex(audit.AuditError,'STATE'):run_a([artifact('<port protocol="tcp" portid="23"><state state="SENSITIVE"/></port>')])

class WorkbookPath(unittest.TestCase):
 def test_actual_materializer_and_renderer_150_plus_repetition(self):
  d=enrich([observation(host_key='SYNTHETIC-'+str(i),service='telnet') for i in range(150)]+[observation(host_key='SYNTHETIC-0')])
  g=group(d)
  d=ns['ip001_mark_action_members'](d,g)
  summary_node=next(n for n in ast.parse(P.read_text()).body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='port_summary' for t in n.targets))
  summary_env=dict(F=F,open_port_review=d)
  exec(compile(ast.Module(body=[summary_node],type_ignores=[]),str(P),'exec'),summary_env)
  keys=['asset_domain','scan_scope','scan_location','protocol','port','port_label','review_priority','risk_theme','csa_relevance']
  summary=ns['ip001_attach_members'](summary_env['port_summary'],d,keys)
  # Actual producer summary expression and member helper, no asset bootstrap.
  frames={'open_port_summary':summary,'open_port_actions':g,'open_port_detail':d}
  env=dict(rns,REPORT_SOURCE_DFS=frames,SOURCE_CONTRACT={k:dict(port_binding={'generation':'g'}) for k in frames},
           RCAIP_POST_NAME='SYNTHETIC ONLY',RCAIP_COUNTRY_NAME='SYNTHETIC',RCAIP_RUN_ID='synthetic_run')
  rns.update(env)
  mapping=next(n.value for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='RCAIP_REPORT_WORKBOOK_SHEETS' for t in n.targets))
  # Pull the actual three configured worksheet expressions, not a parallel renderer.
  items=mapping.args[0].elts
  entries=[n for n in items if ast.literal_eval(n.elts[0]) in env['PORT_SHEET_KEYS']]
  for entry in entries:
   name=ast.literal_eval(entry.elts[0]);config=eval(compile(ast.Expression(entry.elts[1]),str(R),'eval'),env)
   env.setdefault('RCAIP_REPORT_WORKBOOK_SHEETS',OrderedDict())[name]=config
  env.update(workbook_pandas={},port_raw_frames={},manifest_records=[])
  materializer=next(n for n in tree.body if isinstance(n,ast.For) and isinstance(n.target,ast.Tuple) and isinstance(n.target.elts[0],ast.Name) and n.target.elts[0].id=='sheet_order')
  exec(compile(ast.Module(body=[materializer],type_ignores=[]),str(R),'exec'),env)
  env['certify_port_frames'](env['port_raw_frames'])
  env['link_port_members'](env['workbook_pandas'],env['port_raw_frames'])
  env['workbook']=Workbook();env['workbook'].remove(env['workbook'].active)
  renderer=next(n for n in tree.body if isinstance(n,ast.For) and isinstance(n.target,ast.Tuple) and isinstance(n.target.elts[0],ast.Name) and n.target.elts[0].id=='sheet_index')
  exec(compile(ast.Module(body=[renderer],type_ignores=[]),str(R),'exec'),env)
  out=ROOT/'outputs/ip001-review';out.mkdir(exist_ok=True)
  path=out/'RCAIP_IP001_SYNTHETIC_PORT_REVIEW.xlsx';env['workbook'].save(path)
  book=load_workbook(path)
  self.assertEqual(book.sheetnames,list(env['PORT_SHEET_KEYS']))
  for name in book.sheetnames:
   env['verify_port_sheet'](book[name],env['workbook_pandas'][name])
   self.assertIn('No compliance conclusion',book[name]['A2'].value)
   self.assertEqual(book[name].auto_filter.ref.split(':')[0],'A5')
  self.assertEqual(book['Open_Port_Detail'].max_row,156)
  book.close()

if __name__=='__main__':unittest.main(verbosity=2)
