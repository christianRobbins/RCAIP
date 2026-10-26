"""Synthetic tests for the sole Python presentation route; no Spark import/execution.

Run: python tests/test_cell24_python.py
Generate: python tests/test_cell24_python.py --generate <synthetic-output.xlsx>
Fixture is a preserved export of the original JavaScript syntheticRuntime().
"""
import copy
from datetime import date, datetime, timezone
from decimal import Decimal
import importlib.util
import json
import os
from pathlib import Path
import sys
from types import SimpleNamespace
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('cell24', ROOT / 'notebooks/rcaip_cell24_presentation.py')
c = importlib.util.module_from_spec(spec)
spec.loader.exec_module(c)
EXPECTED_POST = 'synthetic_demo'
EXPECTED_RUN = 'synthetic_cell24_review_001'


def fixture():
    return json.loads((ROOT / 'tests/fixtures/cell24_synthetic.json').read_text())


def rows(r, kind):
    return r[c.PRODUCTS[kind]]['rows']


def valid(r):
    return c.validate_cell24(r, EXPECTED_POST, EXPECTED_RUN)


class FakeType:
    def __init__(self, name, element=None):
        self.name = name
        self.elementType = element

    def typeName(self):
        return self.name


class FakeRow(dict):
    def asDict(self, recursive=False):
        return dict(self)


class FakeFrame:
    """Only the exact adapter surface exists; no source/table APIs."""
    def __init__(self, product, session):
        self.product = product
        self.sparkSession = session
        self.calls = []
        self.expressions = []
        self.bound = None
        self.on_collect = None
        fields = []
        for name in product['columns']:
            value = next((r[name] for r in product['rows'] if r[name] is not None), None)
            kind = ('array' if isinstance(value, list) else 'boolean' if type(value) is bool else
                    'long' if type(value) is int else 'double' if type(value) is float else
                    'decimal' if isinstance(value, Decimal) else 'timestamp' if isinstance(value, datetime) else
                    'date' if isinstance(value, date) else 'string')
            fields.append(SimpleNamespace(name=name, dataType=FakeType(kind, FakeType('string'))))
        self.schema = SimpleNamespace(fields=fields)

    def selectExpr(self, *expressions):
        self.calls.append('selectExpr'); self.expressions = expressions
        return self

    def limit(self, n):
        self.calls.append(('limit', n)); self.bound = n
        return self

    def collect(self):
        self.calls.append('collect')
        if self.on_collect:
            self.on_collect()
        out = copy.deepcopy(self.product['rows'][:self.bound])
        for row in out:
            for field, expression in zip(self.schema.fields, self.expressions):
                if expression.startswith('CAST(') and row[field.name] is not None:
                    value = row[field.name]
                    row[field.name] = value.isoformat() if isinstance(value, (datetime, date)) else str(value)
        return [FakeRow(r) for r in out]


def session_fixture(r=None):
    r = fixture() if r is None else r
    session = SimpleNamespace(conf=SimpleNamespace(get=lambda key: 'UTC'))
    return {**{k: r[k] for k in (c.READY, c.STATUS, c.POST, c.RUN)},
            **{name: FakeFrame(r[name], session) for name in c.PRODUCTS.values()}}


class ConsumerTest(unittest.TestCase):
    def test_original_fixture_all_fields_unchanged(self):
        r = fixture(); before = copy.deepcopy(r); result = valid(r)
        for kind in c.PRODUCTS:
            self.assertEqual(result[kind], r[c.PRODUCTS[kind]])
        self.assertEqual(r, before)
        self.assertEqual(len(result['cases']['rows']), 2)
        self.assertEqual(len(result['members']['rows']), 3)

    def test_exact_navigator_aggregate_regression(self):
        r = fixture()
        rows(r, 'members')[1]['asset_id'] = rows(r, 'members')[0]['asset_id']
        rows(r, 'cases')[0]['affected_asset_count'] = 1
        rows(r, 'cases')[0]['affected_asset_keys'] = [rows(r, 'members')[0]['asset_id']]
        with self.assertRaisesRegex(c.InputContractError, 'Aggregate affected_asset_count != member-row count'):
            valid(r)

    def test_same_asset_in_different_cases_is_allowed_without_deduplication(self):
        r = fixture(); shared = rows(r, 'members')[0]['asset_id']
        rows(r, 'members')[2]['asset_id'] = shared
        rows(r, 'cases')[1]['affected_asset_keys'] = [shared]
        result = valid(r)
        self.assertEqual(len(result['members']['rows']), 3)
        self.assertEqual(sum(x['affected_asset_count'] for x in result['cases']['rows']), 3)

    def test_original_fail_closed_requirements(self):
        mutations = []
        for name in c.PRODUCTS.values():
            mutations.append((f'missing {name}', lambda r, name=name: r.pop(name)))
        for field, value in [(c.READY, False), (c.READY, 'true'), (c.STATUS, 'NOT_CLOSED'), (c.POST, 'foreign'), (c.RUN, 'other')]:
            mutations.append((field + str(value), lambda r, field=field, value=value: r.update({field: value})))
        for kind in ('cases', 'members'):
            for field, value in [('post_profile_key', 'foreign'), ('post_profile_key', None), ('assessment_run_id', 'other')]:
                mutations.append((kind + field + str(value), lambda r, kind=kind, field=field, value=value: rows(r, kind)[0].update({field: value})))
            mutations.append((kind + ' schema', lambda r, kind=kind: r[c.PRODUCTS[kind]]['columns'].pop()))
        case_bad = {'case_instance_id': '', 'affected_member_count': 99, 'affected_asset_keys': ['foreign'],
                    'residual_concern': None, 'action_owner': 'INVENTED', 'case_class': 'CYBERSECURITY_CONTROL_CONDITION', 'source_record_count': 0}
        member_bad = {'case_instance_id': 'ORPHAN', 'asset_id': None, 'primary_customer': 'POST_DT',
                      'recommended_next_action': 'No action', 'evidence_confidence': 'HIGH'}
        for kind, bad in [('cases', case_bad), ('members', member_bad)]:
            for field, value in bad.items():
                mutations.append((kind + field, lambda r, kind=kind, field=field, value=value: rows(r, kind)[0].update({field: value})))
            mutations.append((kind + ' duplicate', lambda r, kind=kind: rows(r, kind).append(copy.deepcopy(rows(r, kind)[0]))))
            mutations.append((kind + ' partial', lambda r, kind=kind: rows(r, kind).pop()))
        for field, value in [('measurement_unit', 'ASSURANCE_CASE'), ('active_workload', 'YES'), ('count', -1), ('count', 0.5), ('stage_order', '1')]:
            mutations.append(('closure' + field + str(value), lambda r, field=field, value=value: rows(r, 'closure')[0].update({field: value})))
        mutations += [('missing closure', lambda r: rows(r, 'closure').pop()),
                      ('duplicate stage', lambda r: rows(r, 'closure').__setitem__(2, copy.deepcopy(rows(r, 'closure')[1]))),
                      ('multi-case member', lambda r: rows(r, 'members')[2].update(member_id=rows(r, 'members')[0]['member_id']))]
        for index in (1, 9, 12, 13):
            mutations.append((f'closure count {index}', lambda r, index=index: rows(r, 'closure')[index].update(count=rows(r, 'closure')[index]['count'] + 1)))
        for name, mutate in mutations:
            with self.subTest(name=name):
                r = fixture(); mutate(r)
                with self.assertRaises(c.InputContractError):
                    valid(r)
        self.assertEqual(len(mutations), 44)

    def test_empty_residual_and_post_normalization(self):
        r = fixture(); r[c.POST] = ' SYNTHETIC_DEMO '
        rows(r, 'cases')[0]['post_profile_key'] = ' SYNTHETIC_DEMO '
        valid(r)
        rows(r, 'cases').clear(); rows(r, 'members').clear()
        rows(r, 'closure')[7]['count'] = 4
        for row in rows(r, 'closure')[8:]: row['count'] = 0
        self.assertEqual(valid(r)['cases']['rows'], [])

    def test_numeric_temporal_literal_transport(self):
        self.assertEqual(c.evidence_value(0.12345678901234566), '0.12345678901234566')
        self.assertEqual(c.evidence_value(-0.0), '-0.0')
        self.assertEqual(c.evidence_value(Decimal('1.230000')), '1.230000')
        self.assertEqual(c.evidence_value(10**30), str(10**30))
        self.assertEqual(c.evidence_value(12), 12)
        self.assertEqual(c.evidence_value(''), '""')
        self.assertIsNone(c.evidence_value(None))
        self.assertEqual(c.evidence_value('10000-01-01T00:00:00Z'), '"10000-01-01T00:00:00Z"')
        self.assertEqual(c.evidence_value('\\u0001\x01'), '\\\\u0001\\u0001')
        for value in ([-0.0], [1], [Decimal('1.2')], [[1]], float('inf'), float('nan'), Decimal('NaN')):
            with self.subTest(value=str(value)), self.assertRaises(c.InputContractError): c.evidence_value(value)

    def test_adapter_bound_and_lossless_projection(self):
        r = fixture(); p = r[c.PRODUCTS['members']]
        p['columns'] += ['test_decimal', 'test_timestamp', 'test_large_integer']
        for row in p['rows']:
            row.update(test_decimal=Decimal('1.230000'), test_timestamp='10000-01-01 00:00:00.123456', test_large_integer=10**20)
        ns = session_fixture(r)
        frame = ns[c.PRODUCTS['members']]
        next(f for f in frame.schema.fields if f.name == 'test_timestamp').dataType = FakeType('timestamp')
        result = c.collect_cell24_session(ns, max_rows=20)
        for name in c.PRODUCTS.values():
            self.assertEqual(ns[name].calls, ['selectExpr', ('limit', 21), 'collect'])
        self.assertIn('CAST(`test_timestamp` AS STRING) AS `test_timestamp`', frame.expressions)
        self.assertIn('CAST(`test_decimal` AS STRING) AS `test_decimal`', frame.expressions)
        self.assertEqual(result['members']['rows'][0]['test_timestamp'], '10000-01-01 00:00:00.123456')
        self.assertEqual(result['members']['rows'][0]['test_decimal'], '1.230000')
        self.assertEqual(result['timestamp_session_timezone'], 'UTC')

    def test_native_timestamp_rejects_unsupported_zone_before_any_collection(self):
        # These UTC instants share the example_post_a local clock at the autumn rollback.
        instants = ['2025-10-26 00:30:00', '2025-10-26 01:30:00']
        for kind in c.PRODUCTS:
            for zone in ('Europe/London', '+01:00', '-05:00', 'Europe/Paris', 'GMT', 'utc'):
                with self.subTest(kind=kind, zone=zone):
                    r = fixture(); product = r[c.PRODUCTS[kind]]
                    product['columns'].append('native_instant')
                    for i, row in enumerate(product['rows']):
                        row['native_instant'] = instants[i % 2]
                    ns = session_fixture(r)
                    frame = ns[c.PRODUCTS[kind]]
                    frame.schema.fields[-1].dataType = FakeType('timestamp')
                    frame.sparkSession.conf.get = lambda key, zone=zone: zone
                    with self.assertRaisesRegex(c.InputContractError, 'Native timestamp transport requires'):
                        c.collect_cell24_session(ns)
                    for name in c.PRODUCTS.values():
                        self.assertEqual(ns[name].calls, [])

    def test_native_timestamp_supported_utc_preserves_distinct_instants(self):
        # Synthetic Spark CAST results: UTC does not collapse the rollback pair.
        values = ['2025-10-26 00:30:00', '2025-10-26 01:30:00', '0000-01-01 00:00:00']
        for zone in ('UTC', 'Etc/UTC', 'Z', '+00:00'):
            with self.subTest(zone=zone):
                r = fixture(); product = r[c.PRODUCTS['members']]
                product['columns'].append('native_instant')
                for row, value in zip(product['rows'], values): row['native_instant'] = value
                ns = session_fixture(r); frame = ns[c.PRODUCTS['members']]
                frame.schema.fields[-1].dataType = FakeType('timestamp')
                frame.sparkSession.conf.get = lambda key, zone=zone: zone
                result = c.collect_cell24_session(ns)
                self.assertEqual([row['native_instant'] for row in result['members']['rows']], values)
                self.assertIn('CAST(`native_instant` AS STRING) AS `native_instant`', frame.expressions)

    def test_zone_free_temporal_types_do_not_require_utc(self):
        r = fixture(); product = r[c.PRODUCTS['members']]
        extra = {'local_clock': '2025-10-26 01:30:00', 'extended_date': '10000-01-01',
                 'temporal_text': '0000-01-01T00:00:00'}
        product['columns'].extend(extra)
        for row in product['rows']: row.update(extra)
        ns = session_fixture(r); frame = ns[c.PRODUCTS['members']]
        for field, kind in zip(frame.schema.fields[-3:], ('timestamp_ntz', 'date', 'string')):
            field.dataType = FakeType(kind)
        frame.sparkSession.conf.get = lambda key: 'Europe/London'
        result = c.collect_cell24_session(ns)
        for row in result['members']['rows']:
            self.assertEqual({key: row[key] for key in extra}, extra)
        self.assertIn('`temporal_text`', frame.expressions)

    def test_adapter_missing_mixed_session_and_limit(self):
        for name in (*c.PRODUCTS.values(), c.READY, c.STATUS, c.POST, c.RUN):
            ns = session_fixture(); del ns[name]
            with self.subTest(name=name), self.assertRaisesRegex(c.InputContractError, 'Missing producer-session object'):
                c.collect_cell24_session(ns)
        ns = session_fixture(); ns[c.PRODUCTS['closure']].sparkSession = object()
        with self.assertRaisesRegex(c.InputContractError, 'different Spark sessions'): c.collect_cell24_session(ns)
        with self.assertRaisesRegex(c.InputContractError, 'collection exceeds'): c.collect_cell24_session(session_fixture(), max_rows=1)

    def test_adapter_does_not_set_readiness_or_hide_failures(self):
        ns = session_fixture(); original = dict(ns)
        c.collect_cell24_session(ns)
        self.assertEqual(ns, original)
        for exc in (TimeoutError('synthetic read failure'), MemoryError('synthetic resource failure'), KeyboardInterrupt()):
            ns = session_fixture()
            def fail(exc=exc): raise exc
            ns[c.PRODUCTS['cases']].on_collect = fail
            expected = c.InputContractError if isinstance(exc, TimeoutError) else type(exc)
            with self.assertRaises(expected): c.collect_cell24_session(ns)
        ns = session_fixture()
        ns[c.PRODUCTS['members']].on_collect = lambda: ns.update({c.READY: False})
        with self.assertRaisesRegex(c.InputContractError, 'changed during collection'): c.collect_cell24_session(ns)

    def test_adapter_rejects_numeric_arrays_before_collection(self):
        ns = session_fixture()
        ns[c.PRODUCTS['cases']].schema.fields[-1].dataType = FakeType('array', FakeType('double'))
        with self.assertRaisesRegex(c.InputContractError, 'unsupported numeric/nested array'): c.collect_cell24_session(ns)
        self.assertEqual(ns[c.PRODUCTS['cases']].calls, [])

    def test_no_render_bypass_after_mutation(self):
        data = valid(fixture()); data['authority'][c.READY] = False
        with self.assertRaises(c.InputContractError): c.build_cell24_workbook(data, synthetic=True)


def generate(output):
    r = fixture()
    edge = {'edge_formula': '=1+1', 'edge_error': '#N/A', 'edge_decimal': Decimal('1.230000'),
            'edge_large': 12345678901234567890, 'edge_empty': '', 'edge_null': None,
            'edge_control': 'Literal \\u0001 / control \x01', 'edge_negative_zero': -0.0,
            'edge_timestamp': '10000-01-01T00:00:00Z', 'edge_quote': "'literal", 'edge_plus': '+1+1'}
    r[c.PRODUCTS['members']]['columns'] += list(edge)
    for row in rows(r, 'members'): row.update(edge)
    # Actual production adapter/renderer route exercised with synthetic frame substitutes.
    ns = session_fixture(r)
    workbook = c.workbook_from_cell24_session(ns, synthetic=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    reopened = load_workbook(output)
    data = c.collect_cell24_session(session_fixture(r))
    for kind, sheet in [('cases', 'Case evidence'), ('members', 'Members'), ('closure', 'Branch closure')]:
        p = data[kind]; ws = reopened[sheet]
        assert ws['A2'].value == c.SYNTHETIC_LABEL
        for i, row in enumerate(p['rows'], 8):
            for j, field in enumerate(p['columns'], 1):
                assert ws.cell(i, j).value == c.evidence_value(row[field]), (sheet, i, field)
                assert ws.cell(i, j).data_type not in ('f', 'e')
        table = next(iter(ws.tables.values()))
        assert table.ref == table.autoFilter.ref
        assert range_boundaries(table.ref) == (1, 7, len(p['columns']), 7 + len(p['rows']))
        assert [x.name for x in table.tableColumns] == p['columns']
    if os.environ.get('RCAIP_REVIEW_TEMP'):
        expected = {sheet: [data[kind]['columns'], *[
            [c.evidence_value(row[field]) for field in data[kind]['columns']]
            for row in data[kind]['rows']]]
            for kind, sheet in [('cases', 'Case evidence'), ('members', 'Members'), ('closure', 'Branch closure')]}
        folder = Path(os.environ['RCAIP_REVIEW_TEMP'])
        folder.mkdir(parents=True, exist_ok=True)
        (folder / 'expected.json').write_text(json.dumps(expected), encoding='utf-8')
    review_values = [cell.value for row in reopened['Case review'] for cell in row]
    for case in rows(r, 'cases'):
        for _, field in c.CARD_FIELDS: assert case[field] in review_values
        assert review_values.count(case['case_instance_id']) == 1
    links = 0
    for ws in reopened:
        for row in ws:
            for cell in row:
                if cell.hyperlink:
                    links += 1
                    sheet, addr = cell.hyperlink.location.split('!')
                    assert reopened[sheet.strip("'")][addr].value is not None
    assert links == 10, links
    with ZipFile(output) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml', '.rels')): ET.fromstring(z.read(name))
        assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
    print(f'PASS: generated/reopened {output}; exact evidence cells, narrative, tables/filters, 10 internal links, XML; no desktop Excel claim')


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--generate':
        generate(Path(sys.argv[2]))
    else:
        unittest.main(verbosity=2)
