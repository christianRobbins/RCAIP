"""Prepared synthetic real-Spark smoke gate. Import has no Spark actions.
Call run_memory only after separate approval. Delta requires a reviewed plan
and a separate confirmation. No session/config creation, cleanup or live inputs.
"""
import ast
import hashlib
import json
import re
import sys
import types
from pathlib import Path
from uuid import uuid4

SOURCES = {
    'producer': ('RCAIP_IP_001_Open_Port_Intelligence v2.0.py', '7fbd254de65d6bb050094da9f5f2aad0249584ca1e14cd71307f92196fe2746b'),
    'report': ('RCAIP_REPORT_001_DT_Validation_Package_V2.py', 'ec8cac3b130dabe8277b928c293a89b1aaab29532ff7dae624e9f4367c3c29b8'),
    'audit': ('rcaip_ip001_audit.py', '4896cfd6c1933bfe3773bd0837b63b1cf4b2b055a5663f0b1a0580ce38df373f'),
}
POST, RUN, GENERATION = 'synthetic', 'synthetic_smoke', '2026-10-14 10:00:00.100000'


def check(value, code):
    if not value:
        raise AssertionError(code)


def rejected(operation):
    try:
        operation()
    except (RuntimeError, ValueError):
        return
    raise AssertionError('EXPECTED_REJECTION_MISSING')


def selected(tree, names, env):
    found = set()
    for node in tree.body:
        ids = ({node.name} if isinstance(node, ast.FunctionDef) else
               {t.id for t in node.targets if isinstance(t, ast.Name)} if isinstance(node, ast.Assign) else set())
        if ids & set(names):
            exec(compile(ast.Module(body=[node], type_ignores=[]), '<reviewed-selection>', 'exec'), env)
            found |= ids & set(names)
    check(found == set(names), 'SOURCE_SELECTION_MISSING')


def loop(tree, target, env):
    matches = [n for n in tree.body if isinstance(n, ast.For) and isinstance(n.target, ast.Tuple)
               and getattr(n.target.elts[0], 'id', None) == target]
    check(len(matches) == 1, 'SOURCE_LOOP_AMBIGUOUS')
    exec(compile(ast.Module(body=matches, type_ignores=[]), '<reviewed-loop>', 'exec'), env)


def load(spark, source_dir):
    # Verify ALL sources before executing any selected function or module.
    texts = {}
    for role, (name, digest) in SOURCES.items():
        raw = (Path(source_dir) / name).read_bytes().replace(b'\r\n', b'\n')
        check(hashlib.sha256(raw).hexdigest() == digest, 'SOURCE_HASH_MISMATCH')
        texts[role] = raw
    from pyspark.sql import functions as F, types as T
    from pyspark.sql import SparkSession
    check(isinstance(spark, SparkSession), 'REAL_SPARK_SESSION_REQUIRED')
    check(spark.conf.get('spark.sql.session.timeZone') in ('UTC', 'Etc/UTC', '+00:00', 'Z'), 'UTC_REQUIRED')
    import pandas as pd
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from datetime import datetime, timezone, date
    from decimal import Decimal
    import math
    from collections import OrderedDict
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    env = dict(F=F, T=T, spark=spark, pd=pd, Workbook=openpyxl.Workbook, get_column_letter=get_column_letter,
               Alignment=Alignment, Border=Border, Font=Font, PatternFill=PatternFill, Side=Side,
               datetime=datetime, timezone=timezone, date=date, Decimal=Decimal, math=math, json=json,
               OrderedDict=OrderedDict, ILLEGAL_CHARACTERS_RE=ILLEGAL_CHARACTERS_RE,
               COMPLEX_TYPES=(T.ArrayType, T.MapType, T.StructType),
               RCAIP_POST_SLUG=POST, RCAIP_RUN_ID=RUN, RCAIP_POST_NAME='SYNTHETIC ONLY',
               RCAIP_COUNTRY_NAME='SYNTHETIC', RCAIP_RUN_MODE='SYNTHETIC', IP001_GENERATED_AT_UTC=GENERATION)
    pt, rt = ast.parse(texts['producer']), ast.parse(texts['report'])
    selected(pt, ['IP001_HOST_GRAIN', 'printer_signal_pattern', 'ip001_same_rows', 'ip001_validate_rules',
                  'ip001_validate_observations', 'ip001_enrich_ports', 'ip001_attach_members',
                  'ip001_mark_action_members', 'ip001_printer_ports', 'add_runtime_columns', 'ip001_write_with_metadata'], env)
    selected(rt, ['PORT_SOURCE_KEYS', 'PORT_SHEET_KEYS', 'PORT_EVIDENCE_LIMIT', 'PORT_EVIDENCE_WARNING',
                  'resolve_port_table', 'load_port_table', 'placeholder_schema', 'placeholder_dataframe', 'port_source_or_status',
                  'port_cell', 'port_member_counts', 'certify_port_frames', 'link_port_members',
                  'verify_port_sheet', 'excel_safe_spark_dataframe', 'sanitize_excel_value', 'status_fill', 'safe_table_name'], env)
    styles = [t.id for n in rt.body if isinstance(n, ast.Assign) for t in n.targets if isinstance(t, ast.Name)
              and (t.id.startswith('COLOR_') or t.id.endswith(('_FILL', '_FONT')) or t.id == 'MEDIUM_BLUE')]
    selected(rt, styles, env)
    module = types.ModuleType('_synthetic_reviewed_ip001_audit')
    sys.modules[module.__name__] = module  # dataclass requires its real module namespace
    exec(compile(texts['audit'], '<reviewed-audit>', 'exec'), module.__dict__)
    return env, pt, rt, module


def fixture(env):
    T, spark = env['T'], env['spark']
    strings = ['post_name', 'target_cidr', 'post_slug', 'rcaip_run_id', 'scan_id', 'host_key', 'ip_address',
               'asset_domain', 'scan_scope', 'scan_location', 'scan_source_file', 'protocol', 'hostname',
               'port_state', 'service', 'product', 'version']
    schema = T.StructType([T.StructField(c, T.StringType(), True) for c in strings] +
                         [T.StructField('port', T.IntegerType(), True), T.StructField('service_confidence', T.IntegerType(), True)])
    base = dict(post_name='SYNTHETIC', target_cidr='192.0.2.0/24', post_slug=POST, rcaip_run_id=RUN,
                scan_id='scan_a', host_key='PRINTER', ip_address='192.0.2.1', asset_domain='Synthetic',
                scan_scope='Synthetic', scan_location='Synthetic', scan_source_file='invented.xml',
                protocol='tcp', hostname=None, port_state='open', service='service', product=None, version=None,
                port=23, service_confidence=None)
    rows = [dict(base, port=p, service='jetdirect' if p == 9100 else 'service', service_confidence=10 if p == 21 else None)
            for p in (21, 23, 80, 9100)]
    rows += [dict(base, port=65023, host_key='UNMAPPED'), dict(base, protocol='udp', host_key='UDP'), dict(rows[0])]
    rows += [dict(base, host_key='HOST_'+str(i)) for i in range(101)]
    # Same host identity but separate observation/Post/run must not inherit signal.
    rows += [dict(base, **change) for change in ({'scan_id':'scan_b'}, {'post_slug':'other_synthetic'}, {'rcaip_run_id':'other_run'})]
    rules_schema = 'port int, port_label string, review_priority string, risk_theme string, standard_recommendation string, csa_relevance string'
    rules = spark.createDataFrame([(21,'Neutral','Low','Review','Review configuration','General'),
                                  (80,'Neutral','Low','Review','Review configuration','General'),
                                  (9100,'Neutral','Low','Review','Review configuration','General')], rules_schema)
    return spark.createDataFrame(rows, schema), rules


def products(env, pt, observations, rules):
    detail = env['ip001_enrich_ports'](observations, rules)
    env['open_port_review'] = detail
    for name in ['priority_rank_expression', 'device_port_rollup', 'device_port_prioritized', 'post_action_table', 'port_summary']:
        selected(pt, [name], env)
    detail = env['ip001_mark_action_members'](detail, env['post_action_table'])
    return dict(open_port_detail=env['add_runtime_columns'](detail),
                open_port_actions=env['add_runtime_columns'](env['post_action_table']),
                open_port_summary=env['add_runtime_columns'](env['port_summary']))


def workbook(env, rt, data, status, directory, label):
    import copy
    from openpyxl import load_workbook
    env.update(REPORT_SOURCE_DFS=data, SOURCE_CONTRACT={k: {'port_binding': {'generation': GENERATION}}
               if k in data else {'port_status': status} for k in env['PORT_SOURCE_KEYS']},
               workbook_pandas={}, port_raw_frames={}, manifest_records=[], RCAIP_REPORT_WORKBOOK_SHEETS=env['OrderedDict']())
    mapping = next(n.value for n in rt.body if isinstance(n, ast.Assign) and any(getattr(t,'id',None)=='RCAIP_REPORT_WORKBOOK_SHEETS' for t in n.targets))
    for entry in mapping.args[0].elts:
        name = ast.literal_eval(entry.elts[0])
        if name in env['PORT_SHEET_KEYS']:
            env['RCAIP_REPORT_WORKBOOK_SHEETS'][name] = eval(compile(ast.Expression(entry.elts[1]), '<reviewed-sheet>', 'eval'), env)
    loop(rt, 'sheet_order', env)
    env['certify_port_frames'](env['port_raw_frames'])
    if data and data['open_port_detail'].limit(1).count():
        frames = copy.deepcopy(env['port_raw_frames'])
        frames['Open_Port_Detail'].loc[0,'port'] = 65534
        rejected(lambda: env['certify_port_frames'](frames))
        frames = copy.deepcopy(env['port_raw_frames'])
        frames['Open_Port_Actions'].loc[0,'hostname'] = 'incorrect-parent'
        rejected(lambda: env['certify_port_frames'](frames))
    env['link_port_members'](env['workbook_pandas'], env['port_raw_frames'])
    env['workbook'] = env['Workbook'](); env['workbook'].remove(env['workbook'].active)
    loop(rt, 'sheet_index', env)
    path = Path(directory)/('SYNTHETIC_'+label+'.xlsx')
    check(not path.exists(), 'WORKBOOK_ALREADY_EXISTS')
    env['workbook'].save(path)
    book = load_workbook(path)
    check(book.sheetnames == list(env['PORT_SHEET_KEYS']), 'SHEET_SET')
    for name in book.sheetnames:
        env['verify_port_sheet'](book[name], env['workbook_pandas'][name])
        check(book[name].auto_filter.ref.startswith('A5:'), 'FILTER_RANGE')
        for row in book[name]:
            for cell in row:
                if cell.hyperlink:
                    sheet, anchor = cell.hyperlink.target[1:].split('!')
                    check(book[sheet.strip("'")][anchor].value is not None, 'LINK_TARGET')
    book.close()
    return str(path)


def synthetic_xml(rows):
    """Encode ONLY supplied invented source observations, preserving multiplicity.
    Group by scan/address, not asset identity; duplicate port elements are retained.
    Callers supply original fixture rows independently of the product under test.
    """
    import xml.etree.ElementTree as ET
    scans = {}
    for row in rows:
        check(row['post_slug'] == POST and row['rcaip_run_id'] == RUN, 'SYNTHETIC_XML_SCOPE')
        scans.setdefault(row['scan_id'], {}).setdefault(row['ip_address'], []).append(row)
    if not scans:
        scans['synthetic_empty_scan'] = {}
    artifacts = []
    for scan, hosts in scans.items():
        root = ET.Element('nmaprun', start='1000')
        protocols = {}
        for members in hosts.values():
            for row in members:
                protocols.setdefault(row['protocol'], set()).add(row['port'])
        for protocol, ports in protocols.items():
            ET.SubElement(root, 'scaninfo', protocol=protocol, services=','.join(map(str, sorted(ports))))
        for address, members in hosts.items():
            host = ET.SubElement(root, 'host')
            ET.SubElement(host, 'address', addr=address, addrtype='ipv4')
            ports = ET.SubElement(host, 'ports')
            for row in members:
                port = ET.SubElement(ports, 'port', protocol=row['protocol'], portid=str(row['port']))
                ET.SubElement(port, 'state', state=row['port_state'])
                service = {k:row[k] for k in ('product','version') if row.get(k) is not None}
                if row.get('service') is not None:
                    service['name'] = row['service']
                if service:
                    ET.SubElement(port, 'service', **service)
        ET.SubElement(ET.SubElement(root, 'runstats'), 'finished', time='1001', exit='success')
        raw = ET.tostring(root, encoding='utf-8')
        digest = hashlib.sha256(raw).hexdigest()
        artifacts.append(dict(bytes=raw, sha256=digest, scan_id=scan,
            perspective='INTERNAL_OPENNET', time_basis='UNIX_EPOCH_UTC',
            scope_lineage=dict(post_slug=POST, rcaip_run_id=RUN, scan_id=scan,
                sha256=digest, perspective='INTERNAL_OPENNET', reference='SYNTHETIC SMOKE ONLY')))
    return artifacts


def assert_audit(audit, artifacts, binding, rows, expected, missing=0, extra=0):
    result = audit.audit(artifacts, {'general': {'binding':binding, 'rows':rows},
        'actions':None, 'report_detail':'READ_ERROR'}, post=POST, run=RUN)
    summary = result.status()
    check(summary['general'] == expected, 'AUDIT_STAGE_UNEXPECTED')
    check((summary['general_missing'], summary['general_extra']) == (missing, extra), 'AUDIT_DISCREPANCY_UNEXPECTED')
    check(summary['actions'] == 'NOT_PRODUCED' and summary['report_detail'] == 'UNAVAILABLE_READ_ERROR', 'AUDIT_AVAILABILITY_COLLAPSED')
    check(summary['overall_compliance'] == 'NOT_ASSESSED' and summary['q1143_scope'] == 'UNESTABLISHED', 'AUDIT_APPLICABILITY_INVENTED')
    return summary


def audit_controls(audit, artifacts, binding, rows):
    """A mismatch is an expected negative control ONLY in the explicit mutation."""
    summaries = [assert_audit(audit, artifacts, binding, rows,
                 'AVAILABLE_POSITIVE' if rows else 'AVAILABLE_ZERO_POSITIVES')]
    summaries.append(assert_audit(audit, synthetic_xml([]), binding, [], 'AVAILABLE_ZERO_POSITIVES'))
    if rows:
        changed = [dict(row) for row in rows]
        changed[0]['port'] = 65534  # absent from the finite invented fixture
        summaries.append(assert_audit(audit, artifacts, binding, changed, 'MEMBERSHIP_MISMATCH', 1, 1))
    return summaries


def run_memory(spark, source_dir, *, approved=False, allow_local_workbooks=False):
    """No table reads/writes. FAIL retains exception only on returned private object."""
    result = {'memory': 'NOT_RUN', 'delta': 'NOT_RUN', 'stage': 'approval', 'workbooks': [], 'audit_checks': []}
    if not approved:
        return result
    try:
        result['stage'] = 'runtime_and_source'
        check(allow_local_workbooks, 'LOCAL_WORKBOOK_PERMISSION_REQUIRED')
        env, pt, rt, audit = load(spark, source_dir)
        F = env['F']
        result['stage'] = 'transformations'
        obs, rules = fixture(env)
        rejected(lambda: env['ip001_validate_rules'](rules.unionByName(rules.withColumn('port_label',F.lit('conflict')))))
        rejected(lambda: env['ip001_validate_observations'](obs.withColumn('host_key', F.lit(None).cast('string'))))
        data = products(env, pt, obs, rules)
        detail = data['open_port_detail']
        check(detail.count() == 111, 'FULL_MULTIPLICITY')
        check(env['ip001_printer_ports'](detail).count() == 5, 'CROSS_SCOPE_PRINTER_SPREAD')
        check(detail.filter((F.col('port')==23) & F.col('review_priority').isNotNull()).count()==0, 'INVENTED_PRIORITY')
        check(detail.filter(F.col('port')==65023).count()==1, 'UNMAPPED_LOSS')
        check(detail.filter(F.col('protocol')=='udp').count()==1, 'PROTOCOL_LOSS')
        check(detail.filter(F.col('service_confidence').isNull()).count()==109, 'NULL_INTEGER_LOSS')
        rejected(lambda: env['ip001_same_rows'](obs, obs.withColumn('port',F.lit(12345)), obs.columns,'mutation'))
        # Report fixture must have a single scope; cross-scope association above
        # is tested separately, never silently filtered by the production report.
        scoped = obs.filter((F.col('post_slug')==POST) & (F.col('rcaip_run_id')==RUN))
        data = products(env, pt, scoped, rules)
        result['stage'] = 'timestamp_audit'
        for captured in ('2026-10-14 10:00:00.100000','2026-10-14 10:00:00.000000'):
            timestamp_frame = scoped.limit(1).withColumn('generated_at_utc',F.lit(captured).cast('timestamp').cast('string'))
            rows = timestamp_frame.collect()
            row = rows[0].asDict(); original = row['generated_at_utc']
            binding = audit.Binding('synthetic.results',0,'synthetic-id',
                tuple((field.name, field.dataType.simpleString()) for field in timestamp_frame.schema.fields),
                '0'*64,captured)
            artifacts = synthetic_xml([row])
            result['audit_checks'].extend(audit_controls(audit, artifacts, binding, [row]))
            check(row['generated_at_utc']==original, 'EVIDENCE_REWRITTEN')
            row['generated_at_utc']='2026-10-14 10:00:01'
            rejected(lambda: assert_audit(audit, artifacts, binding, [row], 'AVAILABLE_POSITIVE'))
        result['stage'] = 'report_reopen'
        import tempfile
        directory = tempfile.mkdtemp(prefix='rcaip_synthetic_smoke_')  # no automatic cleanup
        for label, frames, state in [('positive',data,None), ('empty',{k:v.limit(0) for k,v in data.items()},None),
                                     ('not_produced',{},'NOT_PRODUCED'), ('read_unavailable',{},'UNAVAILABLE_READ_ERROR')]:
            result['workbooks'].append(workbook(env,rt,frames,state,directory,label))
        result.update(memory='PASS',stage='complete')
    except Exception as exc:
        result.update(memory='FAIL', _exception=exc)
    return result


def delta_plan(scratch_schema):
    """Pure planning; does not inspect/create/switch schemas or touch Spark."""
    check(bool(re.fullmatch(r'[A-Za-z_]\w*(?:\.[A-Za-z_]\w*)?', scratch_schema)), 'EXPLICIT_SCRATCH_SCHEMA_REQUIRED')
    prefix = scratch_schema+'.rcaip_synthetic_smoke_'+uuid4().hex
    return {'tables': {k:prefix+'_'+k for k in ('summary','actions','detail','empty')},
            'operations':['existence check','create new Delta tables with unique userMetadata',
                          'append different synthetic commit','bounded history and detail metadata',
                          'pinned versionAsOf reads','report/audit handoff'], 'cleanup':'NONE'}


def run_delta(spark, source_dir, plan, *, confirm_synthetic_writes=False):
    result = {'memory':'NOT_RUN','delta':'NOT_RUN','stage':'approval','created_tables':[],'audit_checks':[]}
    if not confirm_synthetic_writes:
        return result
    try:
        result['stage']='plan_and_source'
        tables = plan['tables']
        check(set(tables)=={'summary','actions','detail','empty'}, 'PLAN_TABLE_SET')
        check(all(re.fullmatch(r'[A-Za-z_]\w*(?:\.[A-Za-z_]\w*)?\.rcaip_synthetic_smoke_[0-9a-f]{32}_'+k,v) for k,v in tables.items()), 'PLAN_TARGET_INVALID')
        check(len({v.rsplit('_',1)[0] for v in tables.values()})==1,'PLAN_TARGET_MIXED')
        env, pt, rt, audit = load(spark,source_dir)
        # Read only the explicitly planned names; no live registry/KB lookup.
        check(all(not spark.catalog.tableExists(v) for v in tables.values()),'PREEXISTING_TARGET')
        obs,rules=fixture(env);F=env['F']
        obs=obs.filter((F.col('post_slug')==POST)&(F.col('rcaip_run_id')==RUN))
        data=products(env,pt,obs,rules)
        original_xml = synthetic_xml([row.asDict(recursive=True) for row in obs.collect()])
        execution=dict(state='RUNNING',execution_id=uuid4().hex,post_slug=POST,rcaip_run_id=RUN,
                       generation=GENERATION,sources={'producer':{'sha256':SOURCES['producer'][1]}},
                       actual_xml=[{'synthetic_only':True}],outputs={})
        # Test-only capture fixture, never a historical/production witness.
        roles={'summary':'open_port_summary','actions':'open_port_actions','detail':'open_port_detail','empty':'open_port_detail'}
        for role,key in roles.items():
            result['stage']='delta_'+role
            logical=audit.PORT_PRODUCTS[key] if role!='empty' else 'SYNTHETIC_EMPTY'
            frame=data[key] if role!='empty' else data[key].limit(0)
            table=tables[role];token=audit.output_write_token(execution,logical)
            # Same actual writer metadata option as reviewed save_post_table;
            # errorifexists protects concurrent/preexisting objects from overwrite.
            env['ip001_write_with_metadata'](spark,frame,table,token,'errorifexists')
            result['created_tables'].append(table)
            identity=audit.table_identity(spark,table)
            audit.capture_output(spark,execution,logical,table,token,identity)
            original_record = execution['outputs'].pop(logical)
            env['ip001_write_with_metadata'](spark,frame,table,'SYNTHETIC_OTHER_COMMIT','append')
            audit.capture_output(spark,execution,logical,table,token,identity)
            record=execution['outputs'][logical]
            check(record['version']==original_record['version'],'CAPTURE_CHANGED_ORIGINAL_TOKEN_VERSION')
            loaded=env['load_port_table']({'port_binding':record})
            check(loaded.count()==frame.count(),'PINNED_MEMBERSHIP')
            native=frame
            if 'post_slug' not in native.columns:
                native=native.withColumn('post_slug',F.col('runtime_post_slug'))
            env['ip001_same_rows'](native,loaded,native.columns,'Delta report handoff')
            if role in ('detail','empty'):
                binding=audit.Binding(table,record['version'],identity,record['schema'],record['producer_sha256'],GENERATION)
                rows=audit.read_delta_rows(spark,binding,POST,RUN)
                artifacts = synthetic_xml([]) if role == 'empty' else original_xml
                result['audit_checks'].extend(audit_controls(audit, artifacts, binding, rows))
            rejected(lambda: env['load_port_table']({'port_binding':dict(record,generation='2026-10-14 10:00:01')})) if role!='empty' else None
            rejected(lambda: env['load_port_table']({'port_binding':dict(record,table_id='wrong-id')}))
        execution['state']='CERTIFIED'
        bindings=audit.port_report_bindings(execution)
        env.update(RCAIP_PORT_REPORT_BINDINGS=bindings, get_rcaip_table=lambda name, **kwargs: next(r['table'] for r in bindings.values() if r['logical_name']==name))
        for key,record in bindings.items():
            env['resolve_port_table'](key,[record['logical_name']])
        result.update(delta='PASS',stage='complete')
    except Exception as exc:
        result.update(delta='FAIL',_exception=exc)
    return result


def status(result):
    """Shareable fixed statuses only. Keep paths, object names and exception local."""
    return {k:result[k] for k in ('memory','delta','stage')}
