"""Correction regressions: synthetic transports, actual extracted report code.
No real Spark/Delta, producer top-level execution or enterprise data.
"""
import ast
import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import types
import unittest
from collections import OrderedDict
from openpyxl import Workbook, load_workbook
import test_ip001_retention as t


def materialize(data, status=None):
    env=dict(t.rns,REPORT_SOURCE_DFS=data,SOURCE_CONTRACT={k:{'port_binding':{'generation':'g'}} for k in data},
             RCAIP_POST_NAME='SYNTHETIC ONLY',RCAIP_COUNTRY_NAME='SYNTHETIC',RCAIP_RUN_ID='synthetic_run',
             workbook_pandas={},port_raw_frames={},manifest_records=[],RCAIP_REPORT_WORKBOOK_SHEETS=OrderedDict())
    for key in env['PORT_SOURCE_KEYS']:
        env['SOURCE_CONTRACT'].setdefault(key,{'port_status':status or 'NOT_PRODUCED'})
    env['placeholder_dataframe']=lambda message:t.D([{'source_status':status or 'NOT_PRODUCED','message':message}])
    t.extract(t.R,{'port_source_or_status'},env)
    mapping=next(n.value for n in t.tree.body if isinstance(n,ast.Assign) and any(isinstance(k,ast.Name) and k.id=='RCAIP_REPORT_WORKBOOK_SHEETS' for k in n.targets))
    for entry in mapping.args[0].elts:
        name=ast.literal_eval(entry.elts[0])
        if name in env['PORT_SHEET_KEYS']:
            env['RCAIP_REPORT_WORKBOOK_SHEETS'][name]=eval(compile(ast.Expression(entry.elts[1]),str(t.R),'eval'),env)
    node=next(n for n in t.tree.body if isinstance(n,ast.For) and isinstance(n.target,ast.Tuple) and getattr(n.target.elts[0],'id',None)=='sheet_order')
    exec(compile(ast.Module(body=[node],type_ignores=[]),str(t.R),'exec'),env)
    env['certify_port_frames'](env['port_raw_frames']);env['link_port_members'](env['workbook_pandas'],env['port_raw_frames'])
    return env


def render(env,path):
    env['workbook']=Workbook();env['workbook'].remove(env['workbook'].active)
    node=next(n for n in t.tree.body if isinstance(n,ast.For) and isinstance(n.target,ast.Tuple) and getattr(n.target.elts[0],'id',None)=='sheet_index')
    exec(compile(ast.Module(body=[node],type_ignores=[]),str(t.R),'exec'),env)
    env['workbook'].save(path)
    book=load_workbook(path)
    for name in book.sheetnames:
        env['verify_port_sheet'](book[name],env['workbook_pandas'][name])
        assert 'No compliance conclusion' in book[name]['A2'].value
        assert book[name].auto_filter.ref.startswith('A5:')
        for row in book[name]:
            for cell in row:
                if cell.hyperlink:
                    sheet,anchor=cell.hyperlink.target[1:].split('!')
                    assert book[sheet.strip("'")][anchor].value is not None
    book.close()


def products(rows, rules=t.RULES):
    detail=t.enrich(rows,rules);actions=t.group(detail)
    detail=t.ns['ip001_mark_action_members'](detail,actions)
    return dict(open_port_detail=detail,open_port_actions=actions,open_port_summary=t.summary(detail))


class CorrectedProduct(unittest.TestCase):
    def test_actual_positive_and_nullable_paths(self):
        rows=[t.observation(p,service='jetdirect' if p==9100 else 'service',service_confidence=10 if p==21 else None) for p in (21,23,80,9100)]
        rows += [t.observation(65023,host_key='UNMAPPED',ip_address='192.0.2.2',service_confidence=None),
                 t.observation(23,host_key='OVERRIDE',ip_address='192.0.2.3',service_confidence=None)]
        rows += [t.observation(23,host_key=f'HOST-{i}',service_confidence=None) for i in range(101)]
        rows += [dict(rows[0])]
        rules=t.D([r for r in t.RULES.rows if r['port']!=23])
        data=products(rows,rules)
        self.assertEqual(len(data['open_port_detail'].rows),108)
        self.assertEqual(len(t.ns['ip001_printer_ports'](data['open_port_detail']).rows),5)
        override=next(r for r in data['open_port_actions'].rows if r['host_key']=='OVERRIDE')
        self.assertEqual(override['final_priority'],'Review');self.assertEqual(override['required_port_review'],1)
        self.assertIn('post_response',override);self.assertIn('recommended_action',override)
        env=materialize(data)
        self.assertIsNone(env['workbook_pandas']['Open_Port_Detail']['service_confidence'].iloc[1])
        render(env,t.ROOT/'outputs/ip001-review/RCAIP_IP001_SYNTHETIC_PORT_REVIEW.xlsx')
    def test_empty_not_produced_read_unavailable_reopen(self):
        positive=products([t.observation()])
        empty={k:t.D([],v.columns) for k,v in positive.items()}
        folder=t.ROOT/'outputs/ip001-review'
        for index,(data,status) in enumerate([(empty,None),({},'NOT_PRODUCED'),({},'UNAVAILABLE_READ_ERROR')]):
            env=materialize(data,status)
            for frame in env['workbook_pandas'].values():
                self.assertEqual(len(frame),1);self.assertNotIn('port_members',frame.columns)
            render(env,Path(folder)/f'RCAIP_IP001_SYNTHETIC_STATUS_{index}.xlsx')
    def test_native_nonfinite_still_rejected(self):
        for number in (float('nan'),float('inf'),-float('inf')):
            with self.assertRaises(RuntimeError):materialize(products([t.observation(service_confidence=number)]))
    def test_metrics_keys_and_alias_conflict(self):
        env=materialize(products([t.observation()]))
        for field,value in [('review_port_instances',8),('affected_hosts',7),('protocol','udp')]:
            frames=copy.deepcopy(env['port_raw_frames']);frames['Open_Port_Summary'].loc[0,field]=value
            with self.assertRaises(RuntimeError):t.rns['certify_port_frames'](frames)
        for field,value in [('review_port_count',9),('open_review_ports','80'),('hostname','WRONG'),('target_cidr','WRONG')]:
            frames=copy.deepcopy(env['port_raw_frames']);frames['Open_Port_Actions'].loc[0,field]=value
            with self.assertRaises(RuntimeError):t.rns['certify_port_frames'](frames)
        d=t.D([dict(t.observation(),runtime_post_slug='conflict',generated_at_utc='g')])
        reader=types.SimpleNamespace(option=lambda *a:types.SimpleNamespace(table=lambda *a:d))
        scope=dict(F=t.F,spark=types.SimpleNamespace(read=reader,conf=types.SimpleNamespace(get=lambda *a:'UTC')),RCAIP_POST_SLUG='synthetic',RCAIP_RUN_ID='synthetic_run')
        t.extract(t.R,{'load_port_table'},scope)
        with self.assertRaisesRegex(RuntimeError,'aliases'):scope['load_port_table']({'port_binding':dict(version=0,table='synthetic.results',generation='g')})
    def test_unsafe_xml_and_unestablished_lineage_preserve_positive(self):
        for prefix in [b'<!DOCTYPE nmaprun SYSTEM "file:///private">',b'<!DOCTYPE nmaprun []>',b'<!DOCTYPE nmaprun><!DOCTYPE nmaprun>']:
            a=t.artifact();raw=prefix+a['bytes'];a.update(bytes=raw,sha256=hashlib.sha256(raw).hexdigest())
            with self.assertRaises(t.audit.AuditError):t.run_a([a])
        a=t.artifact();a['scope_lineage']['rcaip_run_id']='wrong'
        s=t.run_a([a],sample=[dict(scan_id='scan_a',ip_address='192.0.2.1',approved=True)]).status()
        self.assertEqual(s['positive_observations'],1);self.assertNotIn('EXPLICIT_OBSERVATION',s['coverage_counts'])
        for malformed in ([{}],['bad']):
            with self.assertRaises(t.audit.AuditError) as caught:t.run_a(malformed)
            self.assertEqual(str(caught.exception),'AUDIT_INPUT_OR_TRANSPORT_FAILURE')


class Transport:
    def __init__(self):
        self.identity='table-id';self.history=[];self.requests=[];self.frame=t.D([dict(t.observation(),runtime_post_slug='synthetic',generated_at_utc='g')])
        self.conf=types.SimpleNamespace(get=lambda *a:'UTC');self.read=self
    def sql(self,query):
        self.requests.append(query)
        return t.D([{'id':self.identity}]) if query.startswith('DESCRIBE DETAIL') else t.D(self.history)
    def option(self,k,v):self.requests.append((k,v));return self
    def table(self,table):self.requests.append(table);return self.frame


class Prospective(unittest.TestCase):
    def setUp(self):
        self.old={k:sys.modules.get(k) for k in ('pyspark','pyspark.sql')}
        sys.modules['pyspark']=types.ModuleType('pyspark');sys.modules['pyspark.sql']=types.SimpleNamespace(functions=t.F)
        # Explicit schema support in the substitute, not an actual Spark datatype.
        t.StringType.simpleString=lambda self:'string';t.ArrayType.simpleString=lambda self:'array<string>'
    def tearDown(self):
        for key,value in self.old.items():
            if value is None:sys.modules.pop(key,None)
            else:sys.modules[key]=value
    def plan(self,spark):
        sources={role:dict(path=t.P,sha256=hashlib.sha256(t.P.read_bytes()).hexdigest()) for role in ('producer','kb_builder','report')}
        return t.audit.prepare_execution(spark,post='synthetic',run='synthetic_run',kb_table='synthetic.kb',kb_version=0,kb_table_id='table-id',source_files=sources,xml_identities=[dict(source_file_name='synthetic.xml',source_file_sha256='a'*64)])
    def test_explicit_input_and_exact_output_not_latest(self):
        spark=Transport();plan=self.plan(spark)
        self.assertEqual(plan['kb']['version'],0);self.assertTrue(all(v is None for v in t.audit.port_report_bindings(plan).values()))
        self.assertEqual(t.audit.verified_producer_source(plan),t.P.read_bytes())
        t.audit.pinned_kb(spark,plan,'synthetic.kb');self.assertIn(('versionAsOf',0),spark.requests)
        plan['state']='RUNNING';t.audit.bind_xml_inputs(plan,[dict(post_slug='synthetic',rcaip_run_id='synthetic_run',source_file_name='synthetic.xml',source_file_sha256='a'*64)])
        spark.frame=spark.frame.withColumn('generated_at_utc',t.F.lit(plan['generation']))
        for logical in t.audit.PORT_PRODUCTS.values():
            token=t.audit.output_write_token(plan,logical)
            spark.history=[{'version':12,'userMetadata':'another writer'},{'version':11,'userMetadata':token}]
            t.audit.capture_output(spark,plan,logical,'synthetic.output',token,'table-id')
        plan['state']='CERTIFIED';bindings=t.audit.port_report_bindings(plan)
        self.assertTrue(all(v['version']==11 for v in bindings.values()))
        self.assertTrue(all(v['execution_id']==plan['execution_id'] for v in bindings.values()))
        report=dict(F=t.F,spark=spark,RCAIP_PORT_REPORT_BINDINGS=bindings,
                    get_rcaip_table=lambda *a,**k:'synthetic.output',RCAIP_POST_SLUG='synthetic',RCAIP_RUN_ID='synthetic_run')
        t.extract(t.R,{'resolve_port_table','load_port_table'},report)
        for key,logical in t.audit.PORT_PRODUCTS.items():
            config=report['resolve_port_table'](key,[logical])
            self.assertEqual(report['load_port_table'](config).count(),1)

    def test_replacement_generation_and_missing_token_stop(self):
        spark=Transport();plan=self.plan(spark);plan['state']='RUNNING';plan['actual_xml']=[{}]
        token=t.audit.output_write_token(plan,'IP001_PORT_SUMMARY')
        spark.identity='replacement'
        with self.assertRaisesRegex(t.audit.AuditError,'BINDING_CHANGED'):t.audit.pinned_kb(spark,plan,'synthetic.kb')
        with self.assertRaisesRegex(t.audit.AuditError,'REPLACED'):t.audit.capture_output(spark,plan,'IP001_PORT_SUMMARY','synthetic.output',token,'table-id')
        spark.identity='table-id';spark.history=[]
        with self.assertRaisesRegex(t.audit.AuditError,'NOT_FOUND'):t.audit.capture_output(spark,plan,'IP001_PORT_SUMMARY','synthetic.output',token,'table-id')
        spark.history=[dict(version=11,userMetadata=token)]
        with self.assertRaisesRegex(t.audit.AuditError,'GENERATION'):t.audit.capture_output(spark,plan,'IP001_PORT_SUMMARY','synthetic.output',token,'table-id')
        with self.assertRaisesRegex(t.audit.AuditError,'NOT_CERTIFIED'):t.audit.port_report_bindings(plan)
    def test_changed_source_and_mixed_source_rejected(self):
        spark=Transport();plan=self.plan(spark);plan['sources']['producer']['sha256']='b'*64
        with self.assertRaisesRegex(t.audit.AuditError,'CHANGED'):t.audit.verified_producer_source(plan)
        plan['state']='RUNNING';plan['actual_xml']=[{}]
        token=t.audit.output_write_token(plan,'IP001_PORT_SUMMARY');bad=json.loads(token);bad['producer_sha256']='c'*64
        with self.assertRaisesRegex(t.audit.AuditError,'TOKEN'):t.audit.capture_output(spark,plan,'IP001_PORT_SUMMARY','synthetic.output',json.dumps(bad),'table-id')


if __name__=='__main__':unittest.main(verbosity=2)
