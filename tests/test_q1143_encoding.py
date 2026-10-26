"""Synthetic workbook regression; no Spark or publication execution."""
import io
import json
import re
import unittest
from openpyxl import load_workbook
import test_q1143_review as original


def decode_string(cell):
    """Independent implementation of the documented string transport grammar."""
    if cell is None:
        return None
    text = re.sub(r'\\\\|\\u([0-9a-f]{4})',
                  lambda m: chr(int(m[1], 16)) if m[1] else '\\', cell)
    return json.loads(text) if text.startswith('"') else text


class EncodingReview(unittest.TestCase):
    def test_actual_workbook_preserves_collision_pairs_and_literal_values(self):
        values = ['', '""', '2026-10-15', '"2026-10-15"',
                  None, 'literal\\u0000\x00', '=1+1', '"quoted\\text"']
        data, execution = original.fixture()
        for row, value in zip(data['ports'], values, strict=True):
            row['scan_end_time_utc'] = value
        workbook = load_workbook(io.BytesIO(original.q.build_workbook(
            original.model(data, execution), synthetic=True)))
        sheet = workbook['Port Evidence']
        col = [c.value for c in sheet[4]].index('scan_end_time_utc') + 1
        cells = [sheet.cell(row, col) for row in range(5, sheet.max_row + 1)]
        exported = [c.value for c in cells]
        self.assertCountEqual([decode_string(v) for v in exported], values)
        self.assertEqual(len(set(exported)), len(values))
        self.assertTrue(all(c.data_type == 's' for c in cells if c.value is not None))
        workbook.close()

    def test_table_filters_and_empty_header_filters(self):
        for data in (None, {'ports': [], 'devices': []}):
            fixture, execution = original.fixture()
            workbook = load_workbook(io.BytesIO(original.q.build_workbook(
                original.model(fixture if data is None else data, execution), synthetic=True)))
            for sheet in workbook:
                if sheet.tables:
                    self.assertIsNone(sheet.auto_filter.ref)
                    for table in sheet.tables.values():
                        self.assertEqual(table.autoFilter.ref, table.ref)
                else:
                    self.assertIsNotNone(sheet.auto_filter.ref)
            workbook.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
