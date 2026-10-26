"""Independent read-only checks of the consumer-generated SYNTHETIC workbook.

Run with RCAIP_SYNTHETIC_OUTPUT and RCAIP_REVIEW_TEMP set by the fixture runner.
Never pass State workbooks here. This is not desktop Excel acceptance.
"""
import json
import os
from pathlib import Path
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


class SyntheticWorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.path = Path(os.environ['RCAIP_SYNTHETIC_OUTPUT'])
        if cls.path.name != 'RCAIP_Cell24_SYNTHETIC_REVIEW.xlsx':
            raise ValueError('This test reads only the named synthetic review artifact')
        cls.expected = json.loads((Path(os.environ['RCAIP_REVIEW_TEMP']) / 'expected.json').read_text())
        cls.book = load_workbook(cls.path, data_only=False)

    def test_labels_and_separate_grains(self):
        self.assertEqual(self.book.sheetnames, ['Case review', 'Case evidence', 'Members', 'Branch closure'])
        for sheet in self.book:
            self.assertEqual(sheet['A2'].value, 'SYNTHETIC TEST DATA / NOT A example_post_a ASSESSMENT')
        self.assertEqual(self.book['Case review']['B5'].value, 2)
        self.assertEqual(self.book['Case review']['B6'].value, 3)
        self.assertEqual(len(self.expected['Branch closure']) - 1, 14)

    def test_every_source_cell_round_trips(self):
        for name, values in self.expected.items():
            for r, row in enumerate(values, 7):
                for c, value in enumerate(row, 1):
                    with self.subTest(sheet=name, row=r, col=c):
                        self.assertEqual(self.book[name].cell(r, c).value, value)

    def test_reader_sees_complete_analytical_chain(self):
        fields = ['condition_title', 'executive_summary', 'why_it_matters', 'evidence_summary',
                  'benign_explanations_considered', 'residual_concern', 'recommended_action',
                  'primary_customer', 'action_owner', 'priority_band', 'priority_basis',
                  'confidence_band', 'confidence_basis']
        headers, *cases = self.expected['Case evidence']
        displayed = [row[1].value for row in self.book['Case review'].iter_rows(min_row=8, max_col=2)]
        for case in cases:
            for field in fields:
                self.assertIn(case[headers.index(field)], displayed)
        for case_id in (c[headers.index('case_instance_id')] for c in cases):
            self.assertEqual(displayed.count(case_id), 1)

    def test_no_formulas_or_excel_errors_in_evidence(self):
        for sheet in self.book:
            for row in sheet:
                for cell in row:
                    self.assertNotIn(cell.data_type, ('f', 'e'), (sheet.title, cell.coordinate))

    def test_native_tables_and_filters(self):
        names = set()
        for sheet_name, values in self.expected.items():
            sheet = self.book[sheet_name]
            self.assertEqual(len(sheet.tables), 1)
            table = next(iter(sheet.tables.values()))
            self.assertNotIn(table.name, names)
            names.add(table.name)
            self.assertEqual(range_boundaries(table.ref), (1, 7, len(values[0]), 6 + len(values)))
            self.assertEqual(table.autoFilter.ref, table.ref)
            self.assertEqual([c.name for c in table.tableColumns], values[0])
            self.assertEqual(len(values[0]), len(set(values[0])))

    def test_xml_is_well_formed_and_no_external_links(self):
        with ZipFile(self.path) as archive:
            self.assertIsNone(archive.testzip())
            for name in archive.namelist():
                self.assertFalse(name.startswith('xl/externalLinks/'))
                if name.endswith(('.xml', '.rels')):
                    ET.fromstring(archive.read(name))
            tables = [n for n in archive.namelist() if n.startswith('xl/tables/') and n.endswith('.xml')]
            self.assertEqual(len(tables), 3)


if __name__ == '__main__':
    unittest.main(verbosity=2)
