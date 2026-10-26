"""Synthetic source-shaped retained outputs. Never execute captured DIAG source."""
import copy
import ast
import hashlib
import io
import json
import re
from pathlib import Path
import sys
import tempfile
import textwrap
import unittest
from decimal import Decimal

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'notebooks'))
import rcaip_ip_population as ip
import test_cell24_retained as old
from openpyxl import load_workbook

POST = 'synthetic_demo'
RUN = 'synthetic_cell24_review_001'
STAMP = '2026-09-26T12:00:00Z'


def fixture():
    book = old.synthetic_notebook()
    cell24 = book['cells'][0]
    payload = cell24['outputs'][1]['metadata'][ip.retained.MIME]
    member_cols = [f['name'] for f in payload['schema']]
    members = [dict(zip(member_cols, r)) for r in payload['data']]
    for row in members:
        row['member_id'] = 'TENABLE_UUID::' + row['candidate_tenable_uuid'].lower()
    payload['data'] = [[r[c] for c in member_cols] for r in members]
    source = (ROOT / 'notebooks/RCAIP_DIAG_001_Asset_Visibility_Funnel.py').read_text(encoding='utf-8-sig')
    sources = {n: next(c for c in source.split('# COMMAND ----------') if f'# Cell {n} ' in c) for n in ip.PRODUCTS}
    cols = {p: ip.source_columns(sources[n], p) for n, ps in ip.PRODUCTS.items() for p in ps}
    raw = []
    for i in range(7):
        raw.append({'asset_key': f'SYNTHETIC-START-{i}', 'candidate_ip': f'198.51.100.{i + 1}',
                    'candidate_tenable_uuid': f'synthetic-uuid-{min(i, 2) if i < 5 else i}',
                    'candidate_tenable_last_observed': '2026-08-01T00:00:00Z',
                    'candidate_tenable_age_days': 3, 'candidate_tenable_acr_score': Decimal('0.12345678901234566'),
                    'device_family': 'SYNTHETIC DEVICE', 'network_scope': 'SYNTHETIC NETWORK',
                    'identity_basis': 'IP', 'identity_strength': 'WEAK', 'source_mask_active': 'TENABLE'})
    raw += [{**r, 'asset_key': r['asset_id']} for r in members]
    groups = {}
    for r in raw:
        r['tenable_identity_group_key'] = 'TENABLE_UUID::' + r['candidate_tenable_uuid'].lower()
        groups.setdefault(r['tenable_identity_group_key'], []).append(r)
        r.update(canonical_ip=r['candidate_ip'], tenable_ip=r['candidate_ip'], tenable_host_uuid=r['candidate_tenable_uuid'],
                 tenable_last_observed_timestamp=r['candidate_tenable_last_observed'], tenable_days_since_observed=r['candidate_tenable_age_days'],
                 tenable_acr_score=r['candidate_tenable_acr_score'],
                 tenable_key_drivers=r.get('candidate_tenable_key_drivers'),
                 tenable_organization_name=r.get('candidate_tenable_organization_name'),
                 tenable_event_time=r.get('candidate_tenable_event_time'),
                 tenable_observation_age_bucket=r.get('candidate_tenable_age_bucket'))
    products = {p: [] for p in cols}
    def add(n, index, values):
        p = ip.PRODUCTS[n][index]; products[p].append({c: values.get(c) for c in cols[p]})
    for i, r in enumerate(raw):
        group = r['tenable_identity_group_key']
        r.update(tenable_uuid_reconciliation_class='UUID_LINKS_EXISTING_DURABLE_ASSET' if i < 2 else 'UUID_COLLAPSES_MULTIPLE_IP_ASSET_KEYS' if i < 5 else 'UNIQUE_TENABLE_UUID_NO_EXISTING_ASSET_LINK',
                 stronger_non_ip_asset_keys=[f'SYNTHETIC-TARGET-{i}'] if i < 2 else [], candidate_asset_key_count=len(groups[group]))
        r['source_local_uuid_class'] = r['tenable_uuid_reconciliation_class']
        r['tenable_dns_name'] = '=literal\\source\x01' if i == 0 else '#N/A' if i == 1 else '' if i == 2 else None
        add(16, 0, r); add(17, 0, r)
        if i < 5: continue
        add(20, 1, r)
        r.update(member_next_gate='DEVICE_CLASS_AND_INVENTORY_VALIDATION' if i == 5 else 'TARGETED_IDENTITY_ENUMERATION',
                 member_qualification_class='TENABLE_INFRASTRUCTURE_DEVICE_CONTEXT_CANDIDATE' if i == 5 else 'FRESH_GENERAL_PURPOSE_COMPUTE_UNCORROBORATED',
                 workload_promotion_state='NOT_READY_FOR_WORKLOAD_PROMOTION' if i == 5 else 'RESIDUAL_IDENTITY_REVIEW_CANDIDATE')
        add(21, 0, r)
        if i == 5: continue
        r.update(temporal_next_gate='CURRENTNESS_VALIDATION_BEFORE_IDENTITY_ENUMERATION' if i == 6 else 'FOCUSED_IDENTITY_DEVICE_ENUMERATION',
                 temporal_collection_class='LATER_NMAP_COLLECTION_AFTER_TENABLE_LAST_OBSERVED' if i == 6 else 'TENABLE_OBSERVED_AFTER_ALL_RELEVANT_NMAP_COLLECTIONS',
                 workload_promotion_state='NOT_READY_FOR_WORKLOAD_PROMOTION' if i == 6 else 'RESIDUAL_CURRENT_IDENTITY_REVIEW_CANDIDATE',
                 temporal_interpretation='Validate currentness before identity enumeration; non-observation does not prove absence.' if i == 6 else 'Recorded identity observation follows the relevant collection; identity review remains appropriate.')
        add(22, 0, r); add(22, 1 if i == 6 else 2, r)
        if i > 6: add(23, 0, r)
    for i, group in enumerate(list(groups)[:3]):
        rs = groups[group]
        add(20, 0, {'analysis_subject_key': group, 'cohort_id': f'SYNTHETIC-COHORT-{i}',
                    'cohort_status': 'RECONCILIATION_PATTERN_SUPPORTED',
                    'explanation_family': 'SOURCE_LOCAL_IDENTITY_MODEL_RECONCILIATION' if i < 2 else 'ENTERPRISE_IDENTITY_AND_LIFECYCLE_RECONCILIATION',
                    'explanation_class': 'TENABLE_UUID_ALREADY_REPRESENTED_BY_STRONGER_ASSET' if i < 2 else 'CURRENT_TARGET_INVENTORY_IP_VALIDATION',
                    'provisional_treatment_lane': 'RCAIP_ENGINEERING' if i < 2 else 'ENTERPRISE_DATA_RECONCILIATION',
                    'candidate_tenable_uuid': rs[0]['candidate_tenable_uuid'], 'candidate_asset_count': len(rs),
                    'candidate_asset_keys': [r['asset_key'] for r in rs], 'candidate_ips': [r['candidate_ip'] for r in rs],
                    'target_asset_count': 1, 'target_asset_keys': [f'SYNTHETIC-TARGET-{i}'], 'target_names': ['Synthetic stronger asset'],
                    'evidence_source_family_sets': 'TENABLE_HOST_UUID' if i < 2 else 'CMDB|TENABLE',
                    'recommended_treatment': 'Review the retained identity relationship before consolidation; no automatic merge.'})
    checks = {16: {'Active unresolved-IP candidates': 10, 'Distinct candidate asset keys': 10},
              17: {'Active unresolved-IP candidate rows': 10, 'Tenable identity groups after UUID grouping': 8},
              20: {'Starting unresolved-IP asset rows': 10, 'Tenable source-local identity groups': 8, 'No-link candidate asset rows advancing': 5},
              21: {'No-link asset rows represented': 5, 'True no-link Tenable identity groups entering Cell 21': 5},
              22: {'Targeted candidate asset rows': 4},
              23: {'Cell 22 residual current identity rows': 3, 'Cell 22 residual distinct identity groups': 3}}
    checks[17].update({'UUID groups linking an existing stronger Asset Plane identity': 2, 'Candidate rows in existing-asset UUID link groups': 2})
    checks[20].update({'Source-local UUID groups linked to stronger Asset Plane identities': 2,
                      'Remaining groups with stronger exact-IP relationships': 1, 'Remaining groups with no stronger exact-IP relationship': 5,
                      'Linked identity groups assigned to explanatory cohorts': 3, 'Linked/no-link overlap': 0, 'Duplicate explanatory cohort assignments': 0})
    checks[21].update({'Groups not ready for workload promotion': 1, 'Residual targeted-enumeration candidates': 4})
    checks[22].update({'Currentness-validation candidates before identity enumeration': 1,
                      'Residual current identity-enumeration candidates': 3, 'Temporal evidence-limitation / timing-review candidates': 0})
    checks[23].update({'Residual identities missing Cell 21 context': 0, 'Invalid Cell 22 temporal-gate rows': 0, 'Identity groups assigned to multiple customers': 0})
    for n, ps in ip.PRODUCTS.items():
        outputs = []
        for p in ps:
            fields = []
            for c in cols[p]:
                # Source-shaped null columns keep the type demonstrated by the
                # same field in the other synthetic producer outputs.
                v = next((r[c] for records in products.values() for r in records if c in r and r[c] is not None), None)
                kind = {'type': 'array', 'elementType': 'string', 'containsNull': False} if isinstance(v, list) else 'boolean' if type(v) is bool else 'long' if type(v) is int else 'double' if isinstance(v, Decimal) else 'string'
                fields.append({'name': c, 'type': json.dumps(kind)})
            outputs.append({'output_type': 'display_data', 'metadata': {ip.retained.MIME: {
                'schema': fields, 'overflow': False, 'data': [[r[c] for c in cols[p]] for r in products[p]]}}})
        marker = f'RCAIP DIAG 001 - CELL {n}' + (' v1.1' if n == 22 else '') + ' CERTIFICATION'
        saved = marker + f'\nStatus: PASS\nPost: {POST}\nRun ID: {RUN}\n'
        saved += ''.join(f'{label}: {v}\n' for label, v in checks[n].items())
        outputs.append({'output_type': 'stream', 'name': 'stdout', 'text': saved})
        book['cells'].append({'cell_type': 'code', 'source': sources[n], 'execution_count': n, 'metadata': {'synthetic': True}, 'outputs': outputs})
    return book


def table(book, number, product=0):
    return book['cells'][list(ip.PRODUCTS).index(number) + 1]['outputs'][product]['metadata'][ip.retained.MIME]


def change(book, number, product, row, col, value):
    t = table(book, number, product); columns = [f['name'] for f in t['schema']]
    t['data'][row][columns.index(col)] = value


def signature(ws):
    return ([(c.coordinate, c.value, c.data_type, str(c._style), c.number_format,
              c.hyperlink.target if c.hyperlink else None, c.comment.text if c.comment else None)
             for row in ws for c in row], str(ws.merged_cells), ws.freeze_panes,
            {k: str(v) for k, v in ws.tables.items()})


class PopulationTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'SYNTHETIC.ipynb'

    def load(self, book, **kw):
        old.write_fixture(self.path, book)
        return ip.load_population(self.path, expected_post=POST, expected_run=RUN, **kw)

    def test_complete_population_and_native_arrays(self):
        data = self.load(fixture())
        self.assertEqual(len(data['population']), 10)
        self.assertEqual([data['counts'][f] for f in ip.FAMILIES], [(2, 2), (3, 1), (1, 1), (1, 1), (3, 3), (0, 0)])
        linked = next(r for r in data['population'] if r['outcome_family'] == ip.FAMILIES[1])
        self.assertEqual(len(linked['evidence'][ip.PRODUCTS[20][0]]['candidate_ips']), 3)
        self.assertEqual(len(linked['stronger_asset_keys']), 1)
        self.assertTrue(all(r['remaining_question'] for r in data['population']))

    def test_missing_duplicate_foreign_and_ambiguous_members(self):
        mutations = [
            ('Starting/UUID membership', lambda b: table(b, 17)['data'].pop()),
            ('duplicate/blank identity', lambda b: table(b, 16)['data'].append(copy.deepcopy(table(b, 16)['data'][0]))),
            ('membership mismatch', lambda b: change(b, 20, 0, 0, 'candidate_asset_keys', ['foreign'])),
            ('No-link group association', lambda b: change(b, 20, 1, 0, 'tenable_identity_group_key', 'foreign')),
            ('conflicting', lambda b: change(b, 21, 0, 0, 'network_scope', 'foreign')),
            ('Temporal input membership', lambda b: table(b, 22)['data'].pop()),
            ('subproduct membership', lambda b: table(b, 22, 1)['data'].clear()),
            ('Residual/customer/Cell 24', lambda b: table(b, 23)['data'].pop()),
            ('conflicting', lambda b: change(b, 23, 0, 0, 'primary_customer', 'MONITORING')),
        ]
        for reason, mutate in mutations:
            with self.subTest(reason=reason):
                b = fixture(); mutate(b)
                with self.assertRaisesRegex(ip.consumer.InputContractError, reason): self.load(b)

    def test_source_scope_completeness(self):
        for n in ip.PRODUCTS:
            index = list(ip.PRODUCTS).index(n) + 1
            for kind in ('run', 'post', 'fail', 'source', 'duplicate', 'missing', 'overflow', 'schema', 'count'):
                with self.subTest(cell=n, kind=kind):
                    b = fixture(); c = b['cells'][index]
                    reason = ''
                    if kind == 'run': c['outputs'][-1]['text'] = c['outputs'][-1]['text'].replace(RUN, 'foreign'); reason = 'foreign run'
                    if kind == 'post': c['outputs'][-1]['text'] = c['outputs'][-1]['text'].replace(POST, 'foreign'); reason = 'foreign Post'
                    if kind == 'fail': c['outputs'][-1]['text'] = c['outputs'][-1]['text'].replace('Status: PASS', 'Status: FAIL'); reason = 'certification'
                    if kind == 'source': c['source'] += '\nchanged_source = True\n'; reason = 'changed executable'
                    if kind == 'duplicate': b['cells'].append(copy.deepcopy(c)); reason = 'ambiguous producer'
                    if kind == 'missing': c['outputs'].pop(0); reason = 'missing required'
                    if kind == 'overflow': table(b, n)['overflow'] = True; reason = 'Overflow'
                    if kind == 'schema': table(b, n)['schema'][0]['name'] = 'wrong'; reason = 'missing required'
                    if kind == 'count':
                        lines = c['outputs'][-1]['text'].splitlines(); lines[-1] = lines[-1].split(':')[0] + ': 999'; c['outputs'][-1]['text'] = '\n'.join(lines); reason = 'saved count mismatch'
                    with self.assertRaisesRegex(ip.consumer.InputContractError, reason): self.load(b)

    def test_selection_is_not_routing(self):
        b = fixture(); original = self.load(b)
        key = original['cell24']['members']['rows'][0]['asset_id']
        selected = self.load(b, proposed_asset_keys=[key])
        self.assertEqual(selected['cell24'], original['cell24'])
        self.assertEqual(sum(r['proposed_response'] for r in selected['population']), 1)
        for keys in ([key, key], ['foreign'], ['SYNTHETIC-START-0']):
            with self.assertRaisesRegex(ip.consumer.InputContractError, 'Proposed response'): self.load(b, proposed_asset_keys=keys)

    def test_inherited_aliases_and_exact_age_counterexample(self):
        # Independent derivation from the reviewed source selects, never execution.
        source = (ROOT / 'notebooks/RCAIP_DIAG_001_Asset_Visibility_Funnel.py').read_text(encoding='utf-8-sig')
        def aliases(number, name):
            cell = next(c for c in source.split('# COMMAND ----------') if f'# Cell {number} ' in c)
            assignment = next(n for n in ast.walk(ast.parse(cell)) if isinstance(n, ast.Assign)
                and any(isinstance(t, ast.Name) and t.id == name for t in n.targets))
            return {n.func.value.args[0].value: n.args[0].value for n in ast.walk(assignment)
                    if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.func.attr == 'alias'
                    and isinstance(n.func.value, ast.Call) and isinstance(n.func.value.func, ast.Attribute)
                    and n.func.value.func.attr == 'col'}
        left, right = aliases(16, 'tenable_context_df'), aliases(20, 'candidate_tenable_df')
        b = fixture()
        cols16 = {f['name'] for f in table(b, 16)['schema']}
        cols20 = {f['name'] for f in table(b, 20, 1)['schema']}
        expected = {(left[k], right[k]) for k in left.keys() & right.keys() if left[k] in cols16 and right[k] in cols20}
        self.assertEqual(set(ip.INHERITED_ALIASES), expected)
        for before, after in expected:
            with self.subTest(alias=before):
                b = fixture()
                value = 800 if before == 'tenable_days_since_observed' else Decimal('800.000') if before == 'tenable_acr_score' else ' SYNTHETIC-UUID-5 ' if before == 'tenable_host_uuid' else 'conflicting retained text'
                for n in (16, 17): change(b, n, 0, 5, before, value)
                with self.assertRaisesRegex(ip.consumer.InputContractError, 'Cell 16/20 inherited evidence: conflicting ' + before):
                    self.load(b)
        # Signed zero and exact precision are also compared without tolerance.
        b = fixture()
        for n in (16, 17): change(b, n, 0, 8, 'tenable_acr_score', Decimal('0.0'))
        with self.assertRaisesRegex(ip.consumer.InputContractError, 'conflicting tenable_acr_score'): self.load(b)

    def test_invalid_saved_decisions_stop_before_mapping(self):
        fields = ('member_qualification_class', 'member_next_gate', 'workload_promotion_state')
        bad = [('ARBITRARY_CLASS', 'ARBITRARY_GATE', 'ARBITRARY_STATE'),
               ('FRESH_GENERAL_PURPOSE_COMPUTE_UNCORROBORATED', 'TARGETED_IDENTITY_ENUMERATION', 'UNSUPPORTED_PROMOTION'),
               ('TENABLE_INFRASTRUCTURE_DEVICE_CONTEXT_CANDIDATE', 'TARGETED_IDENTITY_ENUMERATION', 'NOT_READY_FOR_WORKLOAD_PROMOTION')]
        for values in bad:
            b = fixture()
            for field, value in zip(fields, values): change(b, 21, 0, 0, field, value)
            with self.assertRaisesRegex(ip.consumer.InputContractError, 'Cell 21: (unsupported|contradictory)'): self.load(b)
        for field in fields:
            b = fixture(); change(b, 21, 0, 0, field, '')
            with self.assertRaisesRegex(ip.consumer.InputContractError, 'Cell 21: missing decision field ' + field): self.load(b)
        for field in ('explanation_class', 'recommended_treatment', 'cohort_status'):
            b = fixture(); change(b, 20, 0, 0, field, '')
            with self.assertRaisesRegex(ip.consumer.InputContractError, 'Cell 20: missing decision field ' + field): self.load(b)
        for field, value in [('explanation_class', 'UNKNOWN'), ('cohort_status', 'MODEL_REVIEW_REQUIRED'),
                             ('provisional_treatment_lane', 'POST_DT')]:
            b = fixture(); change(b, 20, 0, 0, field, value)
            with self.assertRaisesRegex(ip.consumer.InputContractError, 'Cell 20: (unsupported|contradictory)'): self.load(b)

    def test_recorded_review_states_and_scoped_translations(self):
        # A source-supported review remains a review, never a benign outcome.
        b = fixture()
        change(b, 20, 0, 2, 'explanation_class', 'LINKED_RELATIONSHIP_MODEL_REVIEW_REQUIRED')
        change(b, 20, 0, 2, 'cohort_status', 'MODEL_REVIEW_REQUIRED')
        change(b, 20, 0, 2, 'provisional_treatment_lane', 'RCAIP_ENGINEERING_REVIEW')
        change(b, 21, 0, 0, 'member_qualification_class', 'OTHER_DEVICE_CONTEXT_REVIEW')
        change(b, 21, 0, 0, 'member_next_gate', 'MODEL_OR_DEVICE_CONTEXT_REVIEW')
        data = self.load(b)
        r = next(r for r in data['population'] if r['asset_key'] == 'SYNTHETIC-START-5')
        self.assertIn('review', r['remaining_question'].lower())
        self.assertEqual(r['original_next_step'], 'MODEL_OR_DEVICE_CONTEXT_REVIEW')
        r = next(r for r in data['population'] if r['asset_key'] == 'SYNTHETIC-START-2')
        self.assertIn('no resolved or benign outcome', r['reason'])
        for cell, field, value in [(21, 'member_next_gate', 'UNKNOWN'), (20, 'member_next_gate', 'TARGETED_IDENTITY_ENUMERATION')]:
            with self.assertRaisesRegex(ip.consumer.InputContractError, 'unsupported display value'): ip.display_code(cell, field, value)
        self.assertEqual(set(ip.DISPLAY[20, 'explanation_class']), set(ip.LINK_DECISIONS))
        self.assertEqual(set(ip.DISPLAY[21, 'member_qualification_class']), set(ip.QUALIFICATION_DECISIONS))
        self.assertEqual(set(ip.DISPLAY[21, 'member_next_gate']), {x[0] for x in ip.QUALIFICATION_DECISIONS.values()})
        self.assertEqual(set(ip.DISPLAY[22, 'temporal_next_gate']), {x[0] for x in ip.TEMPORAL_DECISIONS.values()})

    def test_selected_summary_zero_one_multiple_and_native_evidence(self):
        data = self.load(fixture())
        keys = [m['asset_id'] for m in data['cell24']['members']['rows']]
        for selection in ([], keys[:1], keys):
            with self.subTest(selected=len(selection)):
                wb = ip.build_population_workbook(self.path, expected_post=POST, expected_run=RUN,
                    proposed_asset_keys=selection, synthetic=True, generated_at=STAMP)
                out = Path(self.temp.name) / 'selection.xlsx'; wb.save(out)
                check = load_workbook(out); self.addCleanup(check.close)
                summary, detail = check['RCSO Summary'], check['IP Population Review']
                cards = [c.row for row in summary for c in row if c.value == ip.PROPOSED_LABEL]
                self.assertEqual(len(cards), len(selection))
                by_key = {r['asset_key']: r for r in data['population']}
                for first, key in zip(cards, sorted(selection)):
                    item = by_key[key]
                    self.assertEqual(summary.cell(first + 1, 2).value, ip.consumer.evidence_value(item['observed_ip']))
                    self.assertEqual(summary.cell(first + 3, 2).value, ip.consumer.evidence_value(item['observation_time']))
                    self.assertEqual(summary.cell(first + 4, 2).value, ip.TIME_BASIS)
                    self.assertEqual(summary.cell(first + 5, 2).value, item['remaining_question'])
                    self.assertEqual(summary.cell(first + 6, 2).value, item['customer_or_lane'])
                    for col in (1, 4):
                        location = summary.cell(first + 8, col).hyperlink.location
                        sheet, address = location.split('!')
                        target = check[sheet.strip("'")][address]
                        if col == 1: self.assertEqual(check['IP Population Review'].cell(target.row, 2).value, key)
                        else: self.assertEqual(check['Case review'].cell(target.row, 2).value, item['case_id'])
                headers = {c.value: c.column for c in detail[7]}
                for line, item in enumerate(data['population'], 8):
                    selected = item['asset_key'] in selection
                    self.assertEqual(detail.cell(line, headers['Communication review status']).value,
                        ip.PROPOSED_LABEL if selected else 'NOT SELECTED FOR THIS PROPOSAL')
                    self.assertEqual(detail.cell(line, 1).fill.fgColor.rgb == '00FFF2CC', selected)
                    for product in [p for ps in ip.PRODUCTS.values() for p in ps]:
                        expected = ip.consumer.evidence_value(ip.retained._json(item['evidence'][product])) if product in item['evidence'] else 'NOT IN THIS BRANCH'
                        self.assertEqual(detail.cell(line, headers[product + ' (native evidence)']).value, expected)
                original = ip.consumer.build_cell24_workbook(data['cell24'], synthetic=True, generated_at=STAMP)
                for ws in original: self.assertEqual(signature(ws), signature(wb[ws.title]))

    def test_lossless_and_workbook_preserves_original_sheets(self):
        b = fixture(); data = self.load(b)
        wb = ip.build_population_workbook(self.path, expected_post=POST, expected_run=RUN, synthetic=True, generated_at=STAMP)
        original = ip.consumer.build_cell24_workbook(data['cell24'], synthetic=True, generated_at=STAMP)
        for ws in original: self.assertEqual(signature(ws), signature(wb[ws.title]))
        out = Path(self.temp.name) / 'SYNTHETIC.xlsx'; wb.save(out)
        check = load_workbook(out); self.addCleanup(check.close)
        self.assertEqual(check.sheetnames, ip.SHEETS)
        self.assertEqual(check['IP Population Review'].max_row, 17)
        texts = '\n'.join(str(c.value) for ws in check for row in ws for c in row)
        for value in ('0.12345678901234566', '-0.0', '0000-01-01', '10000-01-01', 'literal', 'CURRENTNESS_VALIDATION'):
            self.assertIn(value, texts)
        for ws in check:
            for table_obj in ws.tables.values():
                self.assertEqual(table_obj.ref, table_obj.autoFilter.ref)
            for row in ws:
                for c in row: self.assertNotEqual(c.data_type, 'f')
        self.assertEqual(sum(bool(check['IP Population Review'].cell(i, 12).hyperlink) for i in range(8, 18)), 3)

    def test_input_order_independent(self):
        b = fixture(); first = self.load(b)['population']
        for n, ps in ip.PRODUCTS.items():
            for i in range(len(ps)): table(b, n, i)['data'].reverse()
        self.assertEqual(self.load(b)['population'], first)

    def test_empty_population(self):
        b = fixture()
        for cell in b['cells']:
            for output in cell['outputs']:
                payload = output.get('metadata', {}).get(ip.retained.MIME)
                if payload is not None:
                    columns = [f['name'] for f in payload['schema']]
                    if columns == list(ip.consumer.CLOSURE_COLUMNS):
                        for row in payload['data']: row[3] = 0
                    else: payload['data'] = []
                if output.get('output_type') == 'stream':
                    output['text'] = re.sub(r'(?m)^([^\n]+: )\d+( / \d+)?$',
                        lambda m: m[1] + ('0 / 0' if m[2] else '0'), output['text'])
        self.assertEqual(self.load(b)['population'], [])
        wb = ip.build_population_workbook(self.path, expected_post=POST, expected_run=RUN, synthetic=True)
        self.assertEqual(wb['RCSO Summary']['B19'].value, 0)
        self.assertEqual(len(wb['IP Population Review'].tables), 0)

    def test_entrypoint_uses_reviewed_pagination_without_execution(self):
        def statements(path):
            source = path.read_text(encoding='utf-8-sig')
            block = source.split('# PAGINATED_DBFS_READ_START')[1].split('# PAGINATED_DBFS_READ_END')[0]
            return ast.dump(ast.parse(textwrap.dedent(block)), include_attributes=False)
        self.assertEqual(statements(ROOT / 'notebooks/RCAIP_IP_Population_Review.py'),
                         statements(ROOT / 'notebooks/rcaip_cell24_retained_publication.py'))
        # Parse, never import or execute the enterprise caller.
        tree = ast.parse((ROOT / 'notebooks/RCAIP_IP_Population_Review.py').read_text())
        assignment = next(n for n in tree.body if isinstance(n, ast.Assign))
        for name, digest in ast.literal_eval(assignment.value).items():
            text = (ROOT / 'notebooks' / (name + '.py')).read_text(encoding='utf-8-sig')
            self.assertEqual(hashlib.sha256((text.rstrip('\n') + '\n').encode()).hexdigest(), digest)

    def test_scope_is_configuration_and_uuid_can_be_absent(self):
        b = fixture()
        # Existing NO_UUID key recipe, with no invented UUID evidence.
        key = 'SYNTHETIC-START-5'
        change(b, 16, 0, 5, 'tenable_host_uuid', None)
        change(b, 17, 0, 5, 'tenable_host_uuid', None)
        change(b, 17, 0, 5, 'tenable_uuid_reconciliation_class', 'NO_TENABLE_HOST_UUID')
        for n, product in ((20, 1), (21, 0)):
            change(b, n, product, 0, 'candidate_tenable_uuid', None)
            change(b, n, product, 0, 'tenable_identity_group_key', 'NO_UUID::' + key)
        data = self.load(b)
        self.assertEqual(next(r['identity_group'] for r in data['population'] if r['asset_key'] == key), 'NO_UUID::' + key)
        raw = old.fixture_json(b).replace(POST, 'other_synthetic_post').replace(RUN, 'other_synthetic_run')
        self.path.write_text(raw, encoding='utf-8')
        result = ip.load_population(self.path, expected_post='other_synthetic_post', expected_run='other_synthetic_run')
        self.assertEqual(result['counts'], data['counts'])

    def test_no_residual_and_temporal_other_remain_visible(self):
        b = fixture()
        # All four temporal groups defer; starting records are unchanged.
        t = table(b, 22); columns = [f['name'] for f in t['schema']]
        for r in t['data']:
            r[columns.index('temporal_next_gate')] = 'CURRENTNESS_VALIDATION_BEFORE_IDENTITY_ENUMERATION'
            r[columns.index('temporal_collection_class')] = 'LATER_NMAP_COLLECTION_AFTER_TENABLE_LAST_OBSERVED'
            r[columns.index('workload_promotion_state')] = 'NOT_READY_FOR_WORKLOAD_PROMOTION'
        c = table(b, 22, 1); ccols = [f['name'] for f in c['schema']]
        c['data'] = [[dict(zip(columns, r))[k] for k in ccols] for r in t['data']]
        table(b, 22, 2)['data'] = []; table(b, 23)['data'] = []
        for i in (0, 1): b['cells'][0]['outputs'][i]['metadata'][ip.retained.MIME]['data'] = []
        closure = b['cells'][0]['outputs'][2]['metadata'][ip.retained.MIME]
        for row in closure['data']:
            if row[0] == 8: row[3] = 4
            if row[0] >= 9: row[3] = 0
        cert = b['cells'][0]['outputs'][-1]['text']
        labels = ['Cell 22 residual current identity groups', 'Cell 23 RCSO-routed residual members',
                  'Cell 23 Post-routed residual members', 'Cell 23 customer-qualified residual rows',
                  'Final assurance-case members', 'Final deterministic assurance cases', 'Final RCSO cases', 'Final Post DT cases']
        for label in labels: cert = re.sub('^' + re.escape(label) + r': .*$', label + ': 0', cert, flags=re.M)
        cert = cert.replace('Cell 22 currentness-validation groups: 1', 'Cell 22 currentness-validation groups: 4')
        cert = cert.replace('Case-member reconciliation: 3 / 3', 'Case-member reconciliation: 0 / 0').replace('Customer-routing reconciliation: 3 / 3', 'Customer-routing reconciliation: 0 / 0')
        b['cells'][0]['outputs'][-1]['text'] = cert
        b['cells'][-1]['outputs'][-1]['text'] = b['cells'][-1]['outputs'][-1]['text'].replace(': 3', ': 0')
        temporal_cert = b['cells'][-2]['outputs'][-1]
        temporal_cert['text'] = temporal_cert['text'].replace('before identity enumeration: 1', 'before identity enumeration: 4').replace('identity-enumeration candidates: 3', 'identity-enumeration candidates: 0')
        data = self.load(b)
        self.assertEqual(data['counts'][ip.FAMILIES[4]], (0, 0))
        self.assertEqual(len(data['population']), 10)
        wb = ip.build_population_workbook(self.path, expected_post=POST, expected_run=RUN, synthetic=True)
        self.assertEqual(wb['Case review']['B5'].value, 0)
        # A temporal limitation stays explicit rather than disappearing.
        change(b, 22, 0, 0, 'temporal_next_gate', 'TEMPORAL_EVIDENCE_LIMITATION_REVIEW')
        change(b, 22, 0, 0, 'temporal_collection_class', 'TEMPORAL_RELATIONSHIP_REVIEW_REQUIRED')
        table(b, 22, 1)['data'].pop(0)
        closure['data'][7][3] = 3
        cert = b['cells'][0]['outputs'][-1]['text'].replace('currentness-validation groups: 4', 'currentness-validation groups: 3').replace('limitation / other groups: 0', 'limitation / other groups: 1')
        b['cells'][0]['outputs'][-1]['text'] = cert
        temporal_cert['text'] = temporal_cert['text'].replace('before identity enumeration: 4', 'before identity enumeration: 3').replace('timing-review candidates: 0', 'timing-review candidates: 1')
        self.assertEqual(self.load(b)['counts'][ip.FAMILIES[5]], (1, 1))


def generate(destination):
    destination = Path(destination); destination.mkdir(parents=True, exist_ok=True)
    path = destination / 'SYNTHETIC_IP_POPULATION.ipynb'
    old.write_fixture(path, fixture())
    wb = ip.build_population_workbook(path, expected_post=POST, expected_run=RUN,
        proposed_asset_keys=['SYNTHETIC-ASSET-1-1'], synthetic=True, generated_at=STAMP)
    output = destination / 'RCAIP_IP_Population_SYNTHETIC_REVIEW.xlsx'
    wb.save(output); wb.close()
    print(output)


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--generate': generate(sys.argv[2])
    else: unittest.main()
