"""Synthetic-only Q1143 tests; captured-reader/Workspace substitutes are labelled.
No Spark or Databricks execution; source-shaped fixture uses existing fake-frame
producer-expression fixture helpers exclusively to construct invented rows.
"""
import ast
from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import sys
from types import SimpleNamespace
import unittest
from openpyxl import load_workbook
import test_ip001_retention as t

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('q1143',ROOT/'notebooks/rcaip_q1143_review.py')
q=importlib.util.module_from_spec(spec);spec.loader.exec_module(q)
GEN='2026-10-15 10:00:00.100000'


def fixture():
    rows=[t.observation(p,service='jetdirect' if p==9100 else 'service',service_confidence=None,hostname=None) for p in (21,23,80,9100)]
    rows += [dict(rows[0]),t.observation(23,protocol='udp',service_confidence=10),
             t.observation(443,scan_id='scan_b',hostname='=SYNTHETIC()',service_confidence=None),
             t.observation(65023,host_key='PORT_ONLY',ip_address='192.0.2.2',service_confidence=None)]
    rules=t.D([dict(row,review_priority='Low') for row in t.RULES.rows if row['port']!=23])
    detail=t.enrich(rows,rules)
    # Actual device-prioritized expressions; deliberately no action/priority filter.
    env=dict(t.ns,open_port_review=detail)
    for name in ('priority_rank_expression','device_port_rollup','device_port_prioritized'):
        for node in ast.parse(t.P.read_bytes()).body:
            if isinstance(node,ast.Assign) and any(getattr(target,'id',None)==name for target in node.targets):
                exec(compile(ast.Module(body=[node],type_ignores=[]),str(t.P),'exec'),env)
    devices=[dict(row,printer_context_status='CANDIDATE_ONLY',post_response='',post_notes='') for row in env['device_port_prioritized'].rows if row['host_key']!='PORT_ONLY']
    deviceonly=deepcopy(devices[0]);deviceonly.update(host_key='DEVICE_ONLY',ip_address='192.0.2.3')
    members=[]
    for member in deviceonly['port_members']:
        value=json.loads(member);value.update(host_key='DEVICE_ONLY',ip_address='192.0.2.3');members.append(json.dumps(value))
    deviceonly['port_members']=members;devices.append(deviceonly)
    ports=[dict(row,printer_context_status='CANDIDATE_ONLY') for row in detail.rows]
    for row in ports+devices:row.update(generated_at_utc=GEN,runtime_post_slug='synthetic',scan_end_time_utc='0000-01-01 00:00:00')
    data={'ports':ports,'devices':devices}
    execution=dict(state='CERTIFIED',post_slug='synthetic',rcaip_run_id='synthetic_run',generation=GEN,execution_id='SYNTHETIC_EXECUTION',sources={'producer':{'sha256':q.PRODUCER_SHA}},outputs={})
    for kind,logical in q.PRODUCTS.items():
        types=lambda v:'boolean' if type(v) is bool else 'bigint' if type(v) is int else 'array<string>' if isinstance(v,list) else 'string'
        schema=[(k,types(v)) for k,v in data[kind][0].items()]
        record=dict(logical_name=logical,table='synthetic.q1143_'+kind,version=7 if kind=='ports' else 9,table_id='synthetic-id-'+kind,schema=schema,generation=GEN,execution_id=execution['execution_id'],producer_sha256=q.PRODUCER_SHA)
        record['write_token']=json.dumps(dict(logical_name=logical,execution_id=execution['execution_id'],producer_sha256=q.PRODUCER_SHA,generation=GEN,post_slug='synthetic',rcaip_run_id='synthetic_run',write_id='SYNTHETIC_WRITE_'+kind))
        execution['outputs'][logical]=record
    return data,execution


def model(data=None,execution=None):
    if data is None:data,execution=fixture()
    records=q.contracts(execution,'synthetic','synthetic_run')
    return q.validate(data,records,post='synthetic',run='synthetic_run',generation_key=t.audit.generation_key)


class Presentation(unittest.TestCase):
    def test_all_occurrences_groups_and_manual_defaults(self):
        m=model();self.assertEqual(len(m['data']['ports']),8);self.assertEqual(len(m['groups']),4)
        self.assertTrue(all(r['review_priority'] is None for r in m['data']['ports'] if r['port']==23))
        self.assertTrue(any(r['port']==21 and r['review_priority']=='Low' for r in m['data']['ports']))
        self.assertEqual(sum(bool(g['ports']) and not g['devices'] for g in m['groups']),1)
        self.assertEqual(sum(bool(g['devices']) and not g['ports'] for g in m['groups']),1)
        raw=q.build_workbook(m,synthetic=True);book=load_workbook(io.BytesIO(raw))
        self.assertEqual(book.sheetnames,q.SHEETS)
        self.assertEqual(book['Port Evidence'].max_row,12)
        for row in book['Assessor Checks'].iter_rows(min_row=5):
            self.assertIn(row[3].value,('NOT_YET_CONFIRMED','NOT_TESTED'))
            self.assertIsNone(row[4].value)
        text='\n'.join(str(c.value) for row in book['Printer Review'] for c in row)
        self.assertIn('tcp/21 open x2',text);self.assertIn('udp/23 open x1',text)
        self.assertIn('Not observed open; closure not verified',text)
        self.assertIn('Device input only',text);self.assertIn('Port input only',text)
        self.assertFalse(any(c.data_type=='f' for ws in book for row in ws for c in row))
        book.close()

    def test_exact_evidence_nulls_temporal_and_literal_values(self):
        data,execution=fixture();data['ports'][0]['service_confidence']=None
        data['devices'][0]['post_notes']='=FORMULA()\x01 literal\\u0001 #N/A'
        m=model(data,execution);book=load_workbook(io.BytesIO(q.build_workbook(m,synthetic=True)))
        for kind,name in [('ports','Port Evidence'),('devices','Device Evidence')]:
            native=[row for g in m['groups'] for row in g[kind]]
            columns=[x[0] for x in m['records'][kind]['schema']]
            for index,row in enumerate(native,5):
                for ci,key in enumerate(columns,2):self.assertEqual(book[name].cell(index,ci).value,q.evidence(row[key]))
        self.assertEqual(q.evidence(-0.0),'-0.0');self.assertEqual(q.evidence(Decimal('1.2300')),'1.2300')
        self.assertEqual(q.evidence(10**18),str(10**18))
        book.close()

    def test_missing_and_empty_inputs(self):
        data,execution=fixture();data={'ports':[],'devices':[]}
        book=load_workbook(io.BytesIO(q.build_workbook(model(data,execution),synthetic=True)))
        self.assertEqual(book['Printer Review'].max_row,4)
        self.assertEqual(book['Assessor Checks'].max_row,4);book.close()
        for kind in q.PRODUCTS:
            bad=deepcopy(execution);del bad['outputs'][q.PRODUCTS[kind]]
            with self.assertRaises(q.ContractError):q.contracts(bad,'synthetic','synthetic_run')
        with self.assertRaises(q.ContractError):q.validate({'ports':[]},{},post='synthetic',run='synthetic_run',generation_key=t.audit.generation_key)

    def test_binding_failures_before_reads(self):
        _,execution=fixture()
        for field,value in [('table_id',''),('version',-1),('generation','wrong'),('producer_sha256','0'*64),('execution_id','wrong'),('logical_name','wrong')]:
            with self.subTest(field=field):
                bad=deepcopy(execution);bad['outputs'][q.PRODUCTS['ports']][field]=value
                with self.assertRaises(q.ContractError):q.contracts(bad,'synthetic','synthetic_run')
        for change in [{'state':'RUNNING'},{'post_slug':'foreign'},{'rcaip_run_id':'foreign'}]:
            bad=deepcopy(execution);bad.update(change)
            with self.assertRaises(q.ContractError):q.contracts(bad,'synthetic','synthetic_run')
        bad=deepcopy(execution);bad['outputs'][q.PRODUCTS['ports']]['write_token']='{}'
        with self.assertRaises(q.ContractError):q.contracts(bad,'synthetic','synthetic_run')

    def test_row_mismatch_and_membership_discrepancy(self):
        for field,value in [('post_slug','foreign'),('runtime_post_slug','foreign'),('generated_at_utc','2026-10-15 11:00:00'),('host_key',None),('port',42)]:
            data,execution=fixture();data['ports'][0][field]=value
            with self.assertRaises((q.ContractError,t.audit.AuditError)):model(data,execution)
        data,execution=fixture();data['devices'][0]['port_members']=data['devices'][0]['port_members'][:-1]
        m=model(data,execution);self.assertTrue(any('member discrepancy' in g['reconciliation'] for g in m['groups']))
        self.assertEqual(len(m['data']['ports']),8)

    def test_reader_uses_only_captured_bindings(self):
        data,execution=fixture();calls=[]
        def read(spark,binding,post,run,max_rows):
            calls.append(binding);return deepcopy(data['ports' if binding.table.endswith('ports') else 'devices'])
        api=SimpleNamespace(__file__=t.audit.__file__,Binding=t.audit.Binding,generation_key=t.audit.generation_key,read_delta_rows=read,table_identity=lambda s,name:'synthetic-id-'+name.rsplit('_',1)[1])
        m=q.read_inputs(None,execution,api,expected_post='synthetic',expected_run='synthetic_run')
        self.assertEqual([b.version for b in calls],[7,9]);self.assertEqual(len(m['groups']),4)
        bad=deepcopy(execution);bad['state']='RUNNING';calls.clear()
        with self.assertRaises(q.ContractError):q.read_inputs(None,bad,api,expected_post='synthetic',expected_run='synthetic_run')
        self.assertEqual(calls,[])
        api.table_identity=lambda *a:'replaced'
        with self.assertRaises(q.ContractError):q.read_inputs(None,execution,api,expected_post='synthetic',expected_run='synthetic_run')

    def test_workspace_publication_nonoverwrite_and_byte_check(self):
        # Explicit SDK/reader substitutes: exercises actual publisher, no service.
        import types
        from unittest.mock import patch
        uploaded={};events=[]
        def upload(**kw):
            self.assertIs(kw['overwrite'],False)
            self.assertEqual(kw['format'],'AUTO')
            self.assertNotIn(kw['path'],uploaded)
            uploaded[kw['path']]=kw['content'].read();events.append(kw['path'])
        def download(**kw):return io.BytesIO(uploaded[kw['path']])
        client=SimpleNamespace(workspace=SimpleNamespace(upload=upload,download=download))
        modules={'databricks':types.ModuleType('databricks'),'databricks.sdk':types.ModuleType('databricks.sdk'),
                 'databricks.sdk.service':types.ModuleType('databricks.sdk.service'),
                 'databricks.sdk.service.workspace':types.ModuleType('databricks.sdk.service.workspace')}
        modules['databricks.sdk'].WorkspaceClient=lambda:client
        modules['databricks.sdk.service.workspace'].ImportFormat=SimpleNamespace(AUTO='AUTO')
        modules['databricks.sdk.service.workspace'].ExportFormat=SimpleNamespace(AUTO='AUTO')
        with patch.dict(sys.modules,modules),patch.object(q,'read_inputs',return_value=model()):
            first=q.publish(None,None,None,expected_post='synthetic',expected_run='synthetic_run',approved_report_dir='/Users/synthetic/reports',workspace_client=client)
            second=q.publish(None,None,None,expected_post='synthetic',expected_run='synthetic_run',approved_report_dir='/Users/synthetic/reports',workspace_client=client)
            self.assertNotEqual(first['workspace_path'],second['workspace_path'])
            self.assertTrue(first['published_bytes_verified'])
            client.workspace.download=lambda **kw:io.BytesIO(b'wrong')
            with self.assertRaisesRegex(q.ContractError,'PUBLISHED_BYTES_MISMATCH'):
                q.publish(None,None,None,expected_post='synthetic',expected_run='synthetic_run',approved_report_dir='/Users/synthetic/reports',workspace_client=client)

    def test_unsupported_lossy_values_stop(self):
        for value in (float('nan'),float('inf'),[-0.0],{'a':1},'x'*32768):
            with self.assertRaises(q.ContractError):q.evidence(value)


if __name__=='__main__':
    generate='--generate' in sys.argv
    if generate:sys.argv.remove('--generate')
    result=unittest.main(verbosity=2,exit=False).result
    if not result.wasSuccessful():raise SystemExit(1)
    if generate:
        folder=ROOT/'outputs/q1143-review';folder.mkdir(exist_ok=True)
        (folder/'RCAIP_Q1143_SYNTHETIC_REVIEW.xlsx').write_bytes(q.build_workbook(model(),synthetic=True))
