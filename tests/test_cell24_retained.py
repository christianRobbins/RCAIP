"""Synthetic saved notebook outputs only. Captured producer source is never run."""
import copy
import ast
from decimal import Decimal
import hashlib
import inspect
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'notebooks'))
import rcaip_cell24_retained as loader
from openpyxl import load_workbook


def synthetic_notebook(encoded=True):
    runtime = json.loads((ROOT / 'tests/fixtures/cell24_synthetic.json').read_text())
    source = (ROOT / 'notebooks/RCAIP_DIAG_001_Asset_Visibility_Funnel.py').read_text(encoding='utf-8-sig')
    source = next(s for s in source.split('# COMMAND ----------') if f'print("{loader.MARKER}")' in s)
    members = runtime[loader.consumer.PRODUCTS['members']]['rows']
    members[0]['source_provenance'] = '=literal\\backslash\x01'
    members[0]['candidate_tenable_last_observed'] = '0000-01-01 00:00:00.123456'
    members[1]['candidate_tenable_last_observed'] = '10000-01-01T00:00:00Z'
    members[0]['candidate_tenable_acr_score'] = Decimal('0.12345678901234566')
    members[1]['candidate_tenable_acr_score'] = Decimal('-0.0')
    outputs = []
    for name in loader.consumer.PRODUCTS.values():
        product = runtime[name]; fields = []
        for col in product['columns']:
            value = next((r[col] for r in product['rows'] if r[col] is not None), None)
            kind = ({'type': 'array', 'elementType': 'string', 'containsNull': False} if isinstance(value, list)
                    else 'boolean' if type(value) is bool else 'long' if type(value) is int
                    else 'double' if isinstance(value, (float, Decimal)) else 'string')
            fields.append({'name': col, 'type': json.dumps(kind) if encoded else kind})
        outputs.append({'output_type': 'display_data', 'data': {'text/html': 'IGNORED'},
                        'metadata': {loader.MIME: {'schema': fields, 'overflow': False,
                        'data': [[r[col] for col in product['columns']] for r in product['rows']]}}})
    # These are the saved producer's actual printed labels, with synthetic totals.
    certification = '''RCAIP DIAG 001 - CELL 24 CERTIFICATION
Status: PASS
Post: synthetic_demo
Run ID: synthetic_cell24_review_001
Starting active unresolved-IP analytical asset rows: 10
Tenable source-local identity groups: 8
Stronger source-local UUID relationship groups: 2
Stronger exact-IP relationship groups: 1
True no-link Tenable identity groups: 5
Cell 21 groups not ready for workload promotion: 1
Cell 22 temporal identity groups: 4
Cell 22 currentness-validation groups: 1
Cell 22 residual current identity groups: 3
Cell 22 temporal limitation / other groups: 0
Cell 23 RCSO-routed residual members: 2
Cell 23 Post-routed residual members: 1
Cell 23 enterprise-routed residual members: 0
Cell 23 customer-qualified residual rows: 3
Final assurance-case members: 3
Final deterministic assurance cases: 2
Final RCSO cases: 1
Final Post DT cases: 1
Final enterprise cases: 0
Final RCAIP engineering cases: 0
Final monitoring cases: 0
Members assigned to multiple cases: 0
Noncanonical case classes: 0
Cybersecurity-control cases created from unresolved identity alone: 0
Case-member reconciliation: 3 / 3
Identity-group reconciliation: 8 / 8
Analytical asset-row reconciliation: 10 / 10
No-link / Cell 21 reconciliation: 5 / 5
Temporal reconciliation: 4 / 4
Customer-routing reconciliation: 3 / 3
Unresolved-IP branch closure state: CLOSED_FOR_CURRENT_RUN
'''
    outputs.append({'output_type': 'stream', 'name': 'stdout', 'text': certification})
    cell = {'cell_type': 'code', 'source': source, 'outputs': outputs, 'execution_count': 24,
            'metadata': {'execution': {'iopub.execute_input': '2026-09-01T00:00:00Z'}}}
    return {'nbformat': 4, 'nbformat_minor': 5, 'metadata': {'synthetic': True}, 'cells': [cell]}


class Pairs(list):
    """Synthetic object pairs, including deliberate duplicate properties."""


def fixture_json(value, path=(), duplicate_outputs=True):
    if isinstance(value, (dict, Pairs)):
        pairs = list(value.items()) if isinstance(value, dict) else value
        if (isinstance(value, dict) and duplicate_outputs and len(path) == 4 and
                path[0] == 'cells' and type(path[1]) is int and
                path[2] == 'outputs' and type(path[3]) is int):
            pairs = [(k, v) for k, v in pairs for _ in range(2 if k == 'output_type' else 1)]
        return '{' + ','.join(json.dumps(k) + ':' + fixture_json(v, path + (k,), duplicate_outputs)
                              for k, v in pairs) + '}'
    if isinstance(value, list):
        return '[' + ','.join(fixture_json(v, path + (i,), duplicate_outputs)
                               for i, v in enumerate(value)) + ']'
    return loader._json(value)


def write_fixture(path, book, duplicate_outputs=True):
    path.write_text(fixture_json(book, duplicate_outputs=duplicate_outputs), encoding='utf-8')


def load(path):
    return loader.load_retained_cell24(path, expected_post='synthetic_demo', expected_run='synthetic_cell24_review_001')


class RetainedTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'SYNTHETIC.ipynb'

    def test_complete_and_lossless(self):
        book = synthetic_notebook(); write_fixture(self.path, book); data = load(self.path)
        row = data['members']['rows'][0]
        self.assertEqual(row['source_provenance'], '=literal\\backslash\x01')
        self.assertEqual(row['candidate_tenable_last_observed'], '0000-01-01 00:00:00.123456')
        self.assertEqual(row['candidate_tenable_acr_score'], Decimal('0.12345678901234566'))
        self.assertEqual(str(data['members']['rows'][1]['candidate_tenable_acr_score']), '-0.0')
        provenance = json.loads(data['historical_provenance'])
        self.assertEqual(provenance['input_sha256'], hashlib.sha256(self.path.read_bytes()).hexdigest())
        self.assertEqual(provenance['execution_count'], 24)
        self.assertFalse(hasattr(loader, loader.consumer.READY))
        self.assertEqual(provenance['json_compatibility'], {
            'redundant_output_type_count': 4,
            'structural_locations': [f'$.cells[0].outputs[{i}]' for i in range(4)]})

    def test_output_type_pairs_in_other_cells_and_selected_error(self):
        b = synthetic_notebook()
        b['cells'].append({'cell_type': 'code', 'source': 'pass', 'outputs': [
            {'output_type': kind} for kind in ('stream', 'display_data', 'execute_result', 'error')]})
        original = copy.deepcopy(b)
        write_fixture(self.path, b)
        raw = self.path.read_bytes(); data = load(self.path)
        self.assertEqual(b, original)
        self.assertEqual(self.path.read_bytes(), raw)
        provenance = json.loads(data['historical_provenance'])
        self.assertEqual(provenance['input_sha256'], hashlib.sha256(raw).hexdigest())
        self.assertEqual(provenance['json_compatibility']['redundant_output_type_count'], 8)
        self.assertEqual(provenance['json_compatibility']['structural_locations'][-1], '$.cells[1].outputs[3]')
        b['cells'][0]['outputs'][0]['output_type'] = 'error'
        write_fixture(self.path, b)
        with self.assertRaisesRegex(loader.consumer.InputContractError, 'Saved cell error'):
            load(self.path)

    def test_duplicate_exceptions_are_narrow(self):
        for values in [('stream', 'display_data'), ('stream',) * 3, (1, 1),
                       (None, None), ('unknown', 'unknown')]:
            with self.subTest(values=values):
                b = synthetic_notebook()
                b['cells'][0]['outputs'][0] = Pairs([('output_type', v) for v in values])
                write_fixture(self.path, b)
                with self.assertRaisesRegex(loader.consumer.InputContractError, 'Duplicate JSON key'):
                    load(self.path)
        duplicate = Pairs([('output_type', 'stream'), ('output_type', 'stream')])
        # Valid names at invalid structural paths, including object rather than array containers.
        misplaced = [duplicate, {'cells': {'0': {'outputs': [duplicate]}}},
                     {'cells': [{'outputs': {'0': duplicate}}]},
                     {'metadata': {'cells': [{'outputs': [duplicate]}]}},
                     {'cells': [{'outputs': [{'metadata': duplicate}]}]}]
        for b in misplaced:
            write_fixture(self.path, b)
            with self.assertRaisesRegex(loader.consumer.InputContractError, 'Duplicate JSON key'):
                load(self.path)

    def test_all_other_duplicate_properties_rejected(self):
        def dup(obj, key):
            return Pairs(list(obj.items()) + [(key, obj[key])])
        for key in ('data', 'schema', 'overflow'):
            b = synthetic_notebook(); output = b['cells'][0]['outputs'][0]
            output['metadata'][loader.MIME] = dup(output['metadata'][loader.MIME], key)
            write_fixture(self.path, b)
            with self.assertRaisesRegex(loader.consumer.InputContractError, 'Duplicate JSON key'): load(self.path)
        for key in ('source', 'metadata'):
            b = synthetic_notebook(); b['cells'][0] = dup(b['cells'][0], key)
            write_fixture(self.path, b)
            with self.assertRaisesRegex(loader.consumer.InputContractError, 'Duplicate JSON key'): load(self.path)
        for key in ('post_profile_key', 'assessment_run_id', 'certification'):
            b = synthetic_notebook()
            b['cells'][0]['metadata'] = Pairs([(key, 'same'), (key, 'same')])
            write_fixture(self.path, b)
            with self.assertRaisesRegex(loader.consumer.InputContractError, 'Duplicate JSON key'): load(self.path)
        b = synthetic_notebook(); output = b['cells'][0]['outputs'][-1]
        b['cells'][0]['outputs'][-1] = dup(output, 'text')
        write_fixture(self.path, b)
        with self.assertRaisesRegex(loader.consumer.InputContractError, 'Duplicate JSON key'): load(self.path)

    def test_nonduplicate_export_still_supported(self):
        b = synthetic_notebook(); write_fixture(self.path, b, duplicate_outputs=False)
        self.assertEqual(json.loads(load(self.path)['historical_provenance'])['json_compatibility'],
                         {'redundant_output_type_count': 0, 'structural_locations': []})

    def test_exact_minimal_variant_reaches_cell_selection(self):
        self.path.write_bytes(b'{"nbformat":4,"cells":[{"cell_type":"code","source":"pass","outputs":['
            b'{"output_type":"stream","name":"stdout","output_type":"stream","text":["SYNTHETIC ONLY"]}]}]}')
        with self.assertRaisesRegex(loader.consumer.InputContractError, 'Missing or ambiguous Cell 24 producer source'):
            load(self.path)

    def test_reject_invalid_saved_results(self):
        def tables(b): return b['cells'][0]['outputs']
        mutations = [
            lambda b: tables(b).pop(0),
            lambda b: tables(b)[0]['metadata'][loader.MIME].update(overflow=True),
            lambda b: tables(b)[0]['metadata'][loader.MIME].pop('overflow'),
            lambda b: tables(b)[-1].update(text=tables(b)[-1]['text'] + 'Status: FAIL\n'),
            lambda b: tables(b)[-1].update(text=tables(b)[-1]['text'].replace('Status: PASS', 'Status: FAIL')),
            lambda b: tables(b)[-1].update(text=tables(b)[-1]['text'].replace('cases: 2', 'cases: 9')),
            lambda b: tables(b)[-1].update(text=tables(b)[-1]['text'].replace('CLOSED_FOR_CURRENT_RUN', 'NOT_CLOSED')),
            lambda b: tables(b)[0]['metadata'][loader.MIME]['data'][0].__setitem__(26, 'foreign'),
            lambda b: b['cells'].append(copy.deepcopy(b['cells'][0])),
            lambda b: b['cells'][0].update(source=b['cells'][0]['source'] + '\nraise RuntimeError()'),
            lambda b: tables(b).append(copy.deepcopy(tables(b)[0])),
            lambda b: tables(b)[0]['metadata'][loader.MIME]['data'][0].pop(),
            lambda b: tables(b)[0]['metadata'][loader.MIME]['schema'][0].update(type='binary'),
            lambda b: tables(b).pop(),
            lambda b: tables(b)[0]['metadata'][loader.MIME].update(overflow='false'),
            lambda b: tables(b)[-1].update(text=tables(b)[-1]['text'] + 'Post: other\n'),
            lambda b: tables(b)[-1].update(text=tables(b)[-1]['text'].replace('3 / 3', '2 / 3')),
            lambda b: tables(b)[0]['metadata'][loader.MIME]['data'][0].__setitem__(27, 'foreign'),
            lambda b: tables(b)[0]['metadata'][loader.MIME]['data'][0].__setitem__(19, 99),
            lambda b: tables(b)[0]['metadata'].pop(loader.MIME),
        ]
        reasons = ['Three distinct', 'Missing/overflowed', 'Missing/overflowed',
                   'Missing/ambiguous certification field: Status', 'not PASS', 'Saved count mismatch',
                   'Saved branch not closed', 'foreign/missing Post', 'ambiguous Cell 24',
                   'Unexpected Cell 24 executable source', 'duplicate product schema',
                   'Malformed positional', 'Unsupported exported schema type', 'Missing/ambiguous saved certification',
                   'Missing/overflowed', 'Missing/ambiguous certification field: Post',
                   'Saved member reconciliation mismatch', 'foreign/missing run',
                   'Case member/relationship counts', 'Three distinct']
        self.assertEqual(len(mutations), len(reasons))
        for encoded in (False, True):
            for index, (mutate, reason) in enumerate(zip(mutations, reasons)):
                with self.subTest(encoded=encoded, index=index):
                    b = synthetic_notebook(encoded); mutate(b); write_fixture(self.path, b)
                    with self.assertRaisesRegex(loader.consumer.InputContractError, reason): load(self.path)

    def test_encoded_and_decoded_products_are_identical(self):
        results = []
        for encoded in (False, True):
            b = synthetic_notebook(encoded)
            # Ordinary evidence text resembling JSON must not be decoded.
            table = b['cells'][0]['outputs'][1]['metadata'][loader.MIME]
            col = loader.consumer.MEMBER_COLUMNS.index('source_provenance')
            table['data'][0][col] = '{"type":"array"}'
            write_fixture(self.path, b); results.append(load(self.path))
        for kind in loader.consumer.PRODUCTS:
            self.assertEqual(results[0][kind], results[1][kind])
        self.assertEqual(results[1]['members']['rows'][0]['source_provenance'], '{"type":"array"}')

    def test_descriptor_errors_reach_schema_boundary(self):
        descriptors = [('"string', 'Malformed JSON'),
                       ('{"type":"array","type":"string"}', 'Duplicate JSON key'),
                       (json.dumps(json.dumps('string')), 'Unsupported exported schema type'),
                       ('{"type":"array","elementType":"\\\"string\\\""}', 'Unsupported exported schema type'),
                       ('"binary"', 'Unsupported exported schema type'),
                       ('["string"]', 'Unsupported exported schema type')]
        for product in range(3):
            for descriptor, reason in descriptors:
                with self.subTest(product=product, descriptor=descriptor):
                    b = synthetic_notebook()
                    b['cells'][0]['outputs'][product]['metadata'][loader.MIME]['schema'][0]['type'] = descriptor
                    write_fixture(self.path, b)
                    with self.assertRaisesRegex(loader.consumer.InputContractError, reason): load(self.path)

    def test_explicit_fingerprint_representation(self):
        source = synthetic_notebook()['cells'][0]['source']
        self.assertEqual(loader.producer_fingerprint(source), loader.PRODUCER_AST_SHA256)
        tree = ast.parse(source.strip())
        if 'show_empty' in inspect.signature(ast.dump).parameters:
            omitted = hashlib.sha256(ast.dump(tree, show_empty=False).encode()).hexdigest()
            retained = hashlib.sha256(ast.dump(tree, show_empty=True).encode()).hexdigest()
            self.assertEqual(omitted, 'f02f2d35be59a0564f11bc5c4f0826ef6a508fd2dc70c172d9618793a8f90f85')
            self.assertEqual(retained, loader.PRODUCER_AST_SHA256)
            self.assertNotEqual(omitted, retained)
        else:
            self.assertEqual(hashlib.sha256(ast.dump(tree).encode()).hexdigest(), loader.PRODUCER_AST_SHA256)
        self.assertNotEqual(loader.producer_fingerprint(source + '\nraise RuntimeError()'), loader.PRODUCER_AST_SHA256)

    def test_same_cell_binding_and_no_fixed_position(self):
        b = synthetic_notebook(); removed = b['cells'][0]['outputs'].pop(0)
        b['cells'].insert(0, {'cell_type': 'code', 'source': 'pass', 'outputs': [removed]})
        write_fixture(self.path, b)
        with self.assertRaises(loader.consumer.InputContractError): load(self.path)
        b = synthetic_notebook(); b['cells'].insert(0, {'cell_type': 'markdown', 'source': 'Synthetic'})
        b['cells'][1]['outputs'].reverse(); write_fixture(self.path, b)
        self.assertEqual(len(load(self.path)['cases']['rows']), 2)

    def test_actual_loader_renderer_reopens(self):
        write_fixture(self.path, synthetic_notebook())
        output = Path(self.tmp.name) / 'SYNTHETIC.xlsx'
        verify_workbook(self.path, output)


def verify_workbook(notebook, output):
    data = load(notebook)
    wb = loader.workbook_from_retained_cell24(notebook, expected_post='synthetic_demo',
        expected_run='synthetic_cell24_review_001', synthetic=True)
    wb.save(output); reopened = load_workbook(output)
    assert reopened.sheetnames == ['Case review', 'Case evidence', 'Members', 'Branch closure']
    for ws in reopened:
        assert 'RETAINED NOTEBOOK RESULTS REPRODUCTION' in ws['A2'].value
        assert 'NOT LIVE AUTHORITY' in ws['A1'].value and ws['A1'].comment
        assert json.loads(ws['A1'].comment.text)['input_sha256'] == hashlib.sha256(notebook.read_bytes()).hexdigest()
    for kind, sheet in [('cases', 'Case evidence'), ('members', 'Members'), ('closure', 'Branch closure')]:
        rows = data[kind]['rows']
        if kind == 'closure': rows = sorted(rows, key=lambda r: r['stage_order'])
        for r, record in enumerate(rows, 8):
            for col, name in enumerate(data[kind]['columns'], 1):
                assert reopened[sheet].cell(r, col).value == loader.consumer.evidence_value(record[name])
                assert reopened[sheet].cell(r, col).data_type not in ('f', 'e')
        table = next(iter(reopened[sheet].tables.values()))
        assert table.autoFilter.ref == table.ref


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--generate':
        folder = Path(sys.argv[2]); folder.mkdir(parents=True, exist_ok=True)
        notebook = folder / 'Cell24_SYNTHETIC_RETAINED.ipynb'
        write_fixture(notebook, synthetic_notebook())
        verify_workbook(notebook, folder / 'RCAIP_Cell24_SYNTHETIC_RETAINED_REVIEW.xlsx')
        print('PASS synthetic retained notebook -> actual loader/consumer -> reopened workbook')
    else:
        unittest.main(verbosity=2)
